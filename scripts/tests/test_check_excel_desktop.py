"""The local Excel oracle refuses incomplete inputs and repaired workbooks."""

from __future__ import annotations

import importlib.util
import contextlib
import hashlib
import io
import json
import math
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest import mock

SCRIPT = Path(__file__).resolve().parents[1] / "check_excel_desktop.py"
SPEC = importlib.util.spec_from_file_location("check_excel_desktop", SCRIPT)
assert SPEC is not None and SPEC.loader is not None
ORACLE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(ORACLE)


class Collection:
    def __init__(self, items=()):
        self.items = list(items)
        self.Count = len(self.items)

    def Item(self, index):
        return self.items[index - 1]

    def __iter__(self):
        return iter(self.items)


def snapshot_book(name="book.xlsx", text="value"):
    font = SimpleNamespace(Name="Calibri", Size=11, Bold=False, Italic=False,
                           Underline=-4142, Strikethrough=False, Color=0,
                           OutlineFont=False, Shadow=False, Subscript=False, Superscript=False)
    interior = SimpleNamespace(Pattern=-4142, Color=16777215, PatternColor=0)
    borders = {index: SimpleNamespace(LineStyle=1, Weight=2, Color=0) for index in range(5, 11)}
    cell = SimpleNamespace(Value2=text, Formula=text, HasFormula=False, Address="$A$1", MergeCells=False,
                           Font=font, Interior=interior, Borders=SimpleNamespace(Item=borders.__getitem__),
                           NumberFormat="General", HorizontalAlignment=1, VerticalAlignment=1,
                           WrapText=False, ShrinkToFit=False, Orientation=0, IndentLevel=0)
    row = SimpleNamespace(Row=1, RowHeight=15, Hidden=False, OutlineLevel=1)
    column = SimpleNamespace(Column=1, ColumnWidth=18, Hidden=False, OutlineLevel=1)
    sheet = SimpleNamespace(Name="Data", Visible=-1, UsedRange=SimpleNamespace(CountLarge=1, Cells=[cell],
                            Rows=[row], Columns=[column]), Cells=SimpleNamespace(FormatConditions=Collection()),
                            ChartObjects=lambda: Collection(), Hyperlinks=Collection(), AutoFilterMode=False,
                            ListObjects=Collection(), Comments=Collection(), PivotTables=lambda: Collection())
    chart = SimpleNamespace(Name="Chart1", Visible=-1, ChartType=51, SeriesCollection=lambda: Collection())
    book = SimpleNamespace(Name=name, Worksheets=Collection([sheet]), Charts=Collection([chart]),
                           Sheets=Collection([sheet, chart]), ActiveSheet=sheet, Date1904=False, Names=Collection(),
                           Windows=Collection([SimpleNamespace(FreezePanes=False, SplitRow=0, SplitColumn=0)]))
    sheet.Activate = lambda: setattr(book, "ActiveSheet", sheet)
    chart.Activate = lambda: setattr(book, "ActiveSheet", chart)
    return book, sheet, cell


def mock_characters(cell, start, length=None):
    """Model either Python characters or UTF-16 positions explicitly, never COM."""
    text = cell.Value2
    if getattr(cell, "CharacterDisplayNormalization", False):
        text = text.replace("\r\n", "\n").replace("\r", "\n")
    if getattr(cell, "CharacterIndexing", "codepoint") == "utf16":
        encoded = text.encode("utf-16-le")
        units = [encoded[at:at + 2].decode("utf-16-le", errors="surrogatepass")
                 for at in range(0, len(encoded), 2)]
    else:
        units = list(text)
    fonts = getattr(cell, "CharacterFonts", [cell.Font] * len(units))
    if not hasattr(cell, "CharacterCalls"):
        cell.CharacterCalls = []
    cell.CharacterCalls.append((start, length))
    if length is None:
        return SimpleNamespace(Count=len(units), Text=text)
    selected = fonts[start - 1:start - 1 + length]
    font = SimpleNamespace(**{name: getattr(selected[0], name) if all(
        getattr(value, name) == getattr(selected[0], name) for value in selected) else None
        for name in ORACLE.FONT_PROPERTIES})
    content = "".join(units[start - 1:start - 1 + length])
    content = content.encode("utf-16-le", errors="surrogatepass").decode("utf-16-le", errors="surrogatepass")
    return SimpleNamespace(Count=length, Text=content, Font=font)


class DesktopOracle(unittest.TestCase):
    def setUp(self) -> None:
        # Keep every snapshot test independent of pywin32 and a running Excel.
        self.characters = mock.patch.object(ORACLE, "characters", side_effect=mock_characters, create=True).start()
        self.addCleanup(mock.patch.stopall)

    def test_snapshot_detects_rich_text_with_equal_aggregate_fonts(self) -> None:
        for property_name, value in (("Bold", True), ("Italic", True), ("Color", 255),
                                     ("Name", "Arial"), ("Size", 14), ("Underline", 2),
                                     ("Strikethrough", True), ("OutlineFont", True), ("Shadow", True),
                                     ("Subscript", True), ("Superscript", True)):
            before, _, left = snapshot_book(text="ab")
            after, _, right = snapshot_book(text="ab")
            plain = SimpleNamespace(**vars(left.Font))
            changed = SimpleNamespace(**{**vars(left.Font), property_name: value})
            # Mixed color can report zero; do not scan only aggregate-null fonts.
            setattr(left.Font, property_name, 0 if property_name == "Color" else None)
            setattr(right.Font, property_name, 0 if property_name == "Color" else None)
            left.CharacterFonts, right.CharacterFonts = [changed, plain], [plain, changed]
            with self.subTest(property=property_name), mock.patch.object(ORACLE, "special_cells", return_value=[]):
                difference = ORACLE.first_difference(ORACLE.snapshot(before), ORACLE.snapshot(after))
                self.assertIsNotNone(difference, "different authored runs cannot pass on aggregate Font equality")
                self.assertIn("rich_text", difference)

    def test_rich_text_normalizes_modeled_unicode_positions_without_losing_surrogates(self) -> None:
        before, _, left = snapshot_book(text="A\U0001f600e\u0301")
        after, _, right = snapshot_book(text="A\U0001f600e\u0301")
        bold = SimpleNamespace(**{**vars(left.Font), "Bold": True})
        left.CharacterFonts = [left.Font, bold, left.Font, left.Font]
        right.CharacterIndexing = "utf16"
        right.CharacterFonts = [right.Font, bold, bold, right.Font, right.Font]
        with mock.patch.object(ORACLE, "special_cells", return_value=[]):
            expected = ORACLE.snapshot(before)
            actual = ORACLE.snapshot(after)
            self.assertEqual(expected, actual)
            runs = actual["sheets"][0]["rich_text"][0][1]
            self.assertEqual([(run["start_utf16"], run["length_utf16"]) for run in runs], [(1, 1), (2, 2), (4, 2)])
            self.assertEqual(right.CharacterCalls[-1], (5, 1))
            # A font disagreement inside one surrogate pair is not a proven
            # scalar font. Refuse it instead of silently losing half a run.
            right.CharacterFonts[2] = right.Font
            with self.assertRaisesRegex(ValueError, "Font.Bold"):
                ORACLE.snapshot(after)

    def test_rich_text_proves_each_position_without_whole_text(self) -> None:
        # Whole-text access is not needed to prove font runs: each native
        # position must still reproduce the corresponding Value2 bytes.
        for text in ("SOH\x01CR\rX_x0041_", "\u20ac " + "x" * 300):
            for unavailable in (False, True):
                book, _, cell = snapshot_book(text=text)
                changed = SimpleNamespace(**{**vars(cell.Font), "Bold": True})
                cell.CharacterFonts = [cell.Font] * (len(text) - 1) + [changed]
                class Whole:
                    Count = len(text)
                    @property
                    def Text(self):
                        if unavailable:
                            raise ValueError("whole Characters.Text unavailable")
                        return "whole view differs"
                def positions(value, start, length=None):
                    actual = mock_characters(value, start, length)
                    return Whole() if length is None else actual
                with self.subTest(text=text[:12], unavailable=unavailable), \
                        mock.patch.object(ORACLE, "special_cells", return_value=[]), \
                        mock.patch.object(ORACLE, "characters", side_effect=positions):
                    runs = ORACLE.snapshot(book)["sheets"][0]["rich_text"][0][1]
                    self.assertEqual([(run["start_utf16"], run["length_utf16"]) for run in runs],
                                     [(1, len(text) - 1), (len(text), 1)])
                    self.assertTrue(runs[-1]["font"]["Bold"])
                    self.assertEqual(cell.CharacterCalls[-1], (len(text), 1))

    def test_rich_text_accepts_observed_cr_display_without_changing_value(self) -> None:
        # Excel 16.0 build20430: from-rust.xlsx Sheet1!B4 Value2 has CR at7,
        # while Characters(7,1).Text returns LF; Count and position stay intact.
        text = "SOH\x01CR\rX_x0041_"
        book, _, cell = snapshot_book(text=text)
        bold = SimpleNamespace(**{**vars(cell.Font), "Bold": True})
        cell.CharacterFonts = [cell.Font] * len(text)
        cell.CharacterFonts[6] = bold
        def displayed(value, start, length=None):
            part = mock_characters(value, start, length)
            part.Text = part.Text.replace("\r", "\n")
            return part
        with mock.patch.object(ORACLE, "special_cells", return_value=[]), \
                mock.patch.object(ORACLE, "characters", side_effect=displayed):
            observed = ORACLE.snapshot(book)
            self.assertEqual(observed["sheets"][0]["cells"], [["$A$1", text]])
            runs = observed["sheets"][0]["rich_text"][0][1]
            self.assertEqual([(run["start_utf16"], run["length_utf16"]) for run in runs], [(1,6),(7,1),(8,8)])
            self.assertTrue(runs[1]["font"]["Bold"])
            cell.Value2 = text.replace("\r", "\n")
            self.assertIsNotNone(ORACLE.first_difference(observed, ORACLE.snapshot(book)))
            cell.Value2 = text
            cell.CharacterFonts[6] = cell.Font
            self.assertNotEqual(runs, ORACLE.snapshot(book)["sheets"][0]["rich_text"][0][1])

    def test_rich_text_refuses_other_control_substitution_or_position_collapse(self) -> None:
        for raw, display in (("\n", "\r"), ("\r", ""), ("\r", " "), ("\r", "\r\n"), ("\x01", "\n")):
            book, _, _ = snapshot_book(text="A" + raw + "Z")
            def changed(value, start, length=None):
                part = mock_characters(value, start, length)
                if start == 2 and length == 1:
                    part.Text = display
                return part
            with self.subTest(raw=repr(raw), display=repr(display)), \
                    mock.patch.object(ORACLE, "special_cells", return_value=[]), \
                    mock.patch.object(ORACLE, "characters", side_effect=changed):
                with self.assertRaisesRegex(ValueError, r"Data!\$A\$1.*Characters\(2, 1\)"):
                    ORACLE.snapshot(book)

    def test_rich_text_maps_crlf_display_and_surrogate_spans_to_original_offsets(self) -> None:
        text = "A\U0001f600\r\nZ"
        observed = []
        for indexing in ("codepoint", "utf16"):
            book, _, cell = snapshot_book(text=text)
            cell.CharacterDisplayNormalization = True
            cell.CharacterIndexing = indexing
            bold = SimpleNamespace(**{**vars(cell.Font), "Bold": True})
            cell.CharacterFonts = ([cell.Font, cell.Font, bold, cell.Font] if indexing == "codepoint"
                                   else [cell.Font, cell.Font, cell.Font, bold, cell.Font])
            with mock.patch.object(ORACLE, "special_cells", return_value=[]):
                observed.append(ORACLE.snapshot(book))
            self.assertEqual(observed[-1]["sheets"][0]["cells"], [["$A$1", text]])
            runs = observed[-1]["sheets"][0]["rich_text"][0][1]
            self.assertEqual([(run["start_utf16"], run["length_utf16"]) for run in runs], [(1,3),(4,2),(6,1)])
            self.assertTrue(runs[1]["font"]["Bold"])
            if indexing == "utf16":
                self.assertEqual(cell.CharacterCalls, [(1,None),(1,1),(2,2),(4,1),(5,1)])
        self.assertEqual(observed[0], observed[1])

    def test_rich_text_still_refuses_unrecognized_display_collapse(self) -> None:
        book, _, cell = snapshot_book(text="A\n\nB")
        def collapsed(value, start, length=None):
            if length is None:
                return SimpleNamespace(Count=3)
            return mock_characters(value, start, length)
        with mock.patch.object(ORACLE, "special_cells", return_value=[]), \
                mock.patch.object(ORACLE, "characters", side_effect=collapsed):
            with self.assertRaisesRegex(ValueError, "Characters.Count"):
                ORACLE.snapshot(book)

    def test_rich_text_scan_is_linear_and_coalesces_uniform_fonts(self) -> None:
        for size in (64, 128):
            book, _, cell = snapshot_book(text="a" * size)
            with self.subTest(size=size), mock.patch.object(ORACLE, "special_cells", return_value=[]):
                runs = ORACLE.snapshot(book)["sheets"][0]["rich_text"][0][1]
                self.assertEqual(len(runs), 1)
                self.assertEqual(runs[0]["length_utf16"], size)
                self.assertEqual(len(cell.CharacterCalls), size + 1)

    def test_rich_text_refuses_unobservable_or_inconsistent_character_state_with_location(self) -> None:
        for failure in ("unavailable", "count", "text", "font", "indexing"):
            book, _, cell = snapshot_book(text="ab")

            def broken(value, start, length=None):
                if failure == "unavailable":
                    raise AttributeError("Characters unavailable")
                part = mock_characters(value, start, length)
                if length is None and failure == "count":
                    part.Count = 8
                elif length is not None and failure == "text":
                    part.Text = "other"
                elif length is not None and failure == "font":
                    part.Font = SimpleNamespace(**{**vars(cell.Font), "Bold": None})
                elif length is not None and failure == "indexing":
                    part.Text = "z"
                return part

            with self.subTest(failure=failure), mock.patch.object(ORACLE, "special_cells", return_value=[]), \
                    mock.patch.object(ORACLE, "characters", side_effect=broken):
                with self.assertRaisesRegex(ValueError, r"Data!\$A\$1.*unsupported rich text"):
                    ORACLE.snapshot(book)

    def test_rich_text_bounds_per_cell_and_workbook_work_before_scanning(self) -> None:
        book, sheet, cell = snapshot_book(text="a" * 32768)
        with mock.patch.object(ORACLE, "special_cells", return_value=[]):
            with self.assertRaisesRegex(ValueError, r"Data!\$A\$1.*32767.*bound"):
                ORACLE.snapshot(book)
            self.characters.assert_not_called()
        book, sheet, cell = snapshot_book(text="abc")
        second = SimpleNamespace(**{**vars(cell), "Address": "$A$2"})
        sheet.UsedRange.Cells.append(second)
        sheet.UsedRange.CountLarge = 2
        with mock.patch.object(ORACLE, "special_cells", return_value=[]), \
                mock.patch.object(ORACLE, "MAX_SNAPSHOT_CHARACTERS", 5, create=True):
            with self.assertRaisesRegex(ValueError, r"Data!\$A\$2.*5.*bound"):
                ORACLE.snapshot(book)
            self.assertFalse(hasattr(second, "CharacterCalls"))

    def test_rich_text_does_not_inspect_formula_results_numbers_or_empty_cells(self) -> None:
        for text, formula in (("answer", True), (12, False), (None, False), ("", False)):
            book, _, cell = snapshot_book(text=text)
            cell.HasFormula = formula
            with self.subTest(text=text), mock.patch.object(ORACLE, "special_cells", return_value=[]):
                self.assertEqual(ORACLE.snapshot(book)["sheets"][0]["rich_text"], [])
        self.characters.assert_not_called()

    def test_shared_repair_activity_is_indeterminate_and_never_copies_private_log_data(self) -> None:
        book = mock.Mock(RepairMode=False)
        app = mock.Mock()
        app.Workbooks.Open.return_value = book
        app.Workbooks.Count = 0
        private_log = mock.Mock(spec=Path)
        private_log.read_text.return_value = "private workbook and cell contents"
        with mock.patch.object(ORACLE, "repair_logs", return_value={}), \
                mock.patch.object(ORACLE, "changed_repair_logs", return_value=[private_log]):
            with self.assertRaisesRegex(RuntimeError, "indeterminate.*1 unattributed") as caught:
                with ORACLE.open_checked(app, Path("fixture.xlsx")):
                    self.fail("unattributed repair activity cannot verify a workbook")
        self.assertNotIn("private", str(caught.exception))
        private_log.read_text.assert_not_called()
        book.Close.assert_called_once_with(SaveChanges=False)

    def test_repair_mode_remains_direct_evidence_despite_shared_log_activity(self) -> None:
        book = mock.Mock(RepairMode=True)
        app = mock.Mock()
        app.Workbooks.Open.return_value = book
        app.Workbooks.Count = 0
        with mock.patch.object(ORACLE, "repair_logs", return_value={}), \
                mock.patch.object(ORACLE, "changed_repair_logs", return_value=[mock.Mock(spec=Path)]):
            with self.assertRaisesRegex(AssertionError, "RepairMode is true"):
                with ORACLE.open_checked(app, Path("repaired.xlsx")):
                    self.fail("a repaired workbook cannot enter the check")
        book.Close.assert_called_once_with(SaveChanges=False)

    def test_snapshot_detects_chart_and_comment_geometry_changes_inside_same_cells(self) -> None:
        def anchor():
            return SimpleNamespace(TopLeftCell=SimpleNamespace(Address="$A$1"),
                                   BottomRightCell=SimpleNamespace(Address="$B$2"),
                                   Left=1, Top=2, Width=20, Height=30, Placement=1)

        for surface in ("chart", "comment"):
            for property_name in ("Left", "Top", "Width", "Height", "Placement"):
                before, left, _ = snapshot_book()
                after, right, _ = snapshot_book()
                for sheet, changed in ((left, False), (right, True)):
                    shape = anchor()
                    if changed:
                        setattr(shape, property_name, getattr(shape, property_name) + 1)
                    if surface == "chart":
                        shape.Name = "EmbeddedChart"
                        shape.Chart = SimpleNamespace(SeriesCollection=lambda: Collection())
                        sheet.ChartObjects = lambda shape=shape: Collection([shape])
                    else:
                        sheet.Comments = Collection([SimpleNamespace(Parent=SimpleNamespace(Address="$A$1"),
                            Author="fixture", Text=lambda: "note", Shape=shape)])
                with self.subTest(surface=surface, property=property_name), \
                        mock.patch.object(ORACLE, "special_cells", return_value=[]):
                    self.assertNotEqual(ORACLE.snapshot(before), ORACLE.snapshot(after))

    def test_main_refuses_existing_fixture_staging_without_removing_it(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            staged = Path(directory) / "rich_excel.pending.xlsx"
            staged.write_bytes(b"previous unverified workbook")
            with mock.patch.object(sys, "argv", [str(SCRIPT), "fidelity", "--output-dir", directory]), \
                    mock.patch.object(sys, "platform", "win32"), mock.patch.object(ORACLE.subprocess, "Popen") as popen, \
                    mock.patch.object(ORACLE, "write_json", side_effect=AssertionError("validation reached output")):
                with self.assertRaisesRegex(SystemExit, "fresh"):
                    ORACLE.main()
                popen.assert_not_called()
            self.assertEqual(staged.read_bytes(), b"previous unverified workbook")

    def test_fixture_finalization_publishes_both_files_or_rolls_back(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            staged = output / "rich_excel.pending.xlsx"
            staged.write_bytes(b"verified workbook")
            result = {"fixture_candidate": str(staged)}
            with self.assertRaises(FileNotFoundError):
                ORACLE.finish_fixture(result, output, True)
            self.assertFalse((output / "rich_excel.xlsx").exists())
            staged.write_bytes(b"verified workbook")
            staged.with_suffix(".json").write_text('{"excel":{"version":"test"},"verified":true}', encoding="utf-8")
            ORACLE.finish_fixture(result, output, True)
            self.assertEqual(Path(result["fixture"]).read_bytes(), b"verified workbook")
            self.assertEqual(json.loads((output / "rich_excel.json").read_text(encoding="utf-8"))["excel"]["version"], "test")
            self.assertNotIn("fixture_candidate", result)

    def test_snapshot_restores_active_and_hidden_sheets_even_when_view_read_fails(self) -> None:
        book, sheet, _ = snapshot_book()
        active = book.Charts.Item(1)
        book.ActiveSheet = active
        sheet.Visible = 0
        with mock.patch.object(ORACLE, "special_cells", return_value=[]):
            ORACLE.snapshot(book)
            self.assertIs(book.ActiveSheet, active)
            self.assertEqual(sheet.Visible, 0)
            broken = mock.Mock()
            type(broken).FreezePanes = mock.PropertyMock(side_effect=ValueError("view failure"))
            book.Windows = Collection([broken])
            with self.assertRaisesRegex(ValueError, "view failure"):
                ORACLE.snapshot(book)
            self.assertIs(book.ActiveSheet, active)
            self.assertEqual(sheet.Visible, 0)

    def test_shift_coverage_requires_charts_comments_and_all_other_reference_surfaces(self) -> None:
        sheet = {"name": "Data", "formulas": ["$A$1"], "conditional_formats": [{}], "validations": [{}],
                 "hyperlinks": [{}], "autofilter": "$A$1:$B$2", "tables": [{}], "charts": [{"series": ["=Data!A1"]}],
                 "comments": [{}], "merges": ["$C$1:$D$1"]}
        snapshot = {"sheets": [sheet], "names": [{}]}
        ORACLE.require_shift_coverage(snapshot, "Data")
        for key in ("formulas", "conditional_formats", "validations", "hyperlinks", "autofilter", "tables", "charts", "comments", "merges"):
            with self.subTest(feature=key), self.assertRaisesRegex(ValueError, "coverage"):
                ORACLE.require_shift_coverage({**snapshot, "sheets": [{**sheet, key: []}]}, "Data")
        with self.assertRaisesRegex(ValueError, "defined names"):
            ORACLE.require_shift_coverage({**snapshot, "names": []}, "Data")

    def test_main_reports_unavailable_startup_identity_without_guessing_an_excel_process(self) -> None:
        child = mock.Mock(returncode=1)
        child.poll.return_value = 1
        with tempfile.TemporaryDirectory() as directory, \
                contextlib.redirect_stdout(io.StringIO()), \
                mock.patch.object(sys, "argv", [str(SCRIPT), "probe", "--output-dir", directory]), \
                mock.patch.object(sys, "platform", "win32"), mock.patch.object(ORACLE.subprocess, "Popen", return_value=child), \
                mock.patch.object(ORACLE, "visible_modal", return_value=None), mock.patch.object(ORACLE, "stop_owned_excel") as stop:
            self.assertEqual(ORACLE.main(), 1)
            result = json.loads((Path(directory) / "results.json").read_text(encoding="utf-8"))
            self.assertIn("identity unavailable", result["cleanup_limitation"])
            self.assertIn("not terminated", result["cleanup_limitation"])
            stop.assert_called_once()

    def test_fidelity_does_not_publish_before_the_worker_finishes(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            args = SimpleNamespace(exchange=Path(directory), rich=Path(directory) / "rich.xlsx", workbook=[],
                                   shift_original=None, shifted=None, sheet="Data", insert_row=3, insert_column=1,
                                   output_dir=Path(directory))
            book = mock.Mock()
            book.SaveAs.side_effect = lambda path, **kwargs: Path(path).write_bytes(b"mock workbook")
            result = {"excel": {"version": "test"}, "failures": [], "workbooks": {}}
            with mock.patch.object(ORACLE, "fidelity_inputs", return_value=[]), mock.patch.object(ORACLE, "inspect_workbooks"), \
                    mock.patch.object(ORACLE, "open_checked", side_effect=lambda *args: contextlib.nullcontext(book)), \
                    mock.patch.object(ORACLE, "snapshot", return_value={}), \
                    mock.patch.object(ORACLE, "require_shift_coverage", create=True):
                ORACLE.run_fidelity(mock.Mock(), args, result)
            self.assertFalse((Path(directory) / "rich_excel.xlsx").exists())
            self.assertTrue(Path(result["fixture_candidate"]).is_file())

    def test_shift_compares_persisted_native_geometry_and_records_save_normalization(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            args = SimpleNamespace(exchange=root, rich=root / "rich.xlsx", workbook=[],
                                   shift_original=None, shifted=None, sheet="Data", insert_row=3, insert_column=1,
                                   output_dir=root)
            original, saved, shifted, shifted_saved, rich = (mock.Mock() for _ in range(5))
            for book in (original, saved, shifted, shifted_saved, rich):
                book.SaveAs.side_effect = lambda path, **kwargs: Path(path).write_bytes(b"mock workbook")
            def opened(_app, path):
                books = {"from-openpyxl-anchored-fidelity.xlsx": original, "shift_excel.xlsx": saved,
                         "from-rust-shifted.xlsx": shifted, "shift_rust_excel.xlsx": shifted_saved}
                return contextlib.nullcontext(books.get(path.name, rich))
            def observed(book):
                if book is original or book is shifted:
                    return {"geometry": 1.8}
                if book is saved or book is shifted_saved:
                    return {"geometry": 1.2}
                return {}
            result = {"excel": {"version": "test"}, "failures": [], "workbooks": {}}
            with mock.patch.object(ORACLE, "fidelity_inputs", return_value=[]), mock.patch.object(ORACLE, "inspect_workbooks"), \
                    mock.patch.object(ORACLE, "open_checked", side_effect=opened), \
                    mock.patch.object(ORACLE, "snapshot", side_effect=observed), \
                    mock.patch.object(ORACLE, "require_shift_coverage"):
                ORACLE.run_fidelity(mock.Mock(), args, result)
            self.assertEqual(result["failures"], [])
            self.assertEqual(result["shift"]["expected_before_save"], {"geometry": 1.8})
            self.assertEqual(result["shift"]["expected"], {"geometry": 1.2})
            self.assertEqual(result["shift"]["actual_before_save"], {"geometry": 1.8})
            self.assertEqual(result["shift"]["actual"], {"geometry": 1.2})
            original.SaveAs.assert_called_once_with(str((root / "shift_excel.xlsx").resolve()),
                                                   FileFormat=51, AddToMru=False, Local=False)
            shifted.SaveAs.assert_called_once_with(str((root / "shift_rust_excel.xlsx").resolve()),
                                                  FileFormat=51, AddToMru=False, Local=False)

    def test_shift_refuses_differences_on_each_side_of_the_save_cycle(self) -> None:
        for failed_phase in ("before-save", "after-save"):
            with self.subTest(phase=failed_phase), tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                args = SimpleNamespace(exchange=root, shift_original=None, shifted=None, sheet="Data",
                                       insert_row=3, insert_column=1, output_dir=root)
                books = {name: mock.Mock() for name in (
                    "from-openpyxl-anchored-fidelity.xlsx", "shift_excel.xlsx",
                    "from-rust-shifted.xlsx", "shift_rust_excel.xlsx")}
                for book in books.values():
                    book.SaveAs.side_effect = lambda path, **kwargs: Path(path).write_bytes(b"mock workbook")
                snapshots = {
                    id(books["from-openpyxl-anchored-fidelity.xlsx"]): {"position": 1.8},
                    id(books["shift_excel.xlsx"]): {"position": 1.2},
                    id(books["from-rust-shifted.xlsx"]): {"position": 9 if failed_phase == "before-save" else 1.8},
                    id(books["shift_rust_excel.xlsx"]): {"position": 9 if failed_phase == "after-save" else 1.2},
                }
                result = {"failures": []}
                with mock.patch.object(ORACLE, "open_checked", side_effect=lambda _, path:
                                       contextlib.nullcontext(books[path.name])), \
                        mock.patch.object(ORACLE, "snapshot", side_effect=lambda book: snapshots[id(book)]), \
                        mock.patch.object(ORACLE, "require_shift_coverage"):
                    ORACLE.run_shift(mock.Mock(), args, result)
                self.assertEqual(len(result["failures"]), 1)
                self.assertIn(f"fidelity/shift/{failed_phase}: $.position", result["failures"][0])

    def test_snapshot_keeps_literal_workbook_names(self) -> None:
        before, _, _ = snapshot_book("source.xlsx", "[source.xlsx]")
        after, _, _ = snapshot_book("output.xlsx", "[output.xlsx]")
        with mock.patch.object(ORACLE, "special_cells", return_value=[]):
            self.assertNotEqual(ORACLE.snapshot(before), ORACLE.snapshot(after))

    def test_snapshot_detects_styles_geometry_visibility_and_chartsheet_loss(self) -> None:
        changes = {
            "font": lambda book, sheet, cell: setattr(cell.Font, "Color", 255),
            "fill": lambda book, sheet, cell: setattr(cell.Interior, "Color", 255),
            "border": lambda book, sheet, cell: setattr(cell.Borders.Item(7), "LineStyle", -4142),
            "row height": lambda book, sheet, cell: setattr(sheet.UsedRange.Rows[0], "RowHeight", 30),
            "column width": lambda book, sheet, cell: setattr(sheet.UsedRange.Columns[0], "ColumnWidth", 8),
            "row hidden": lambda book, sheet, cell: setattr(sheet.UsedRange.Rows[0], "Hidden", True),
            "column hidden": lambda book, sheet, cell: setattr(sheet.UsedRange.Columns[0], "Hidden", True),
            "sheet hidden": lambda book, sheet, cell: setattr(sheet, "Visible", 0),
            "chartsheet": lambda book, sheet, cell: setattr(book, "Charts", Collection()),
        }
        with mock.patch.object(ORACLE, "special_cells", return_value=[]):
            for feature, change in changes.items():
                before, _, _ = snapshot_book()
                after, sheet, cell = snapshot_book()
                change(after, sheet, cell)
                with self.subTest(feature=feature):
                    self.assertNotEqual(ORACLE.snapshot(before), ORACLE.snapshot(after))

    def test_main_preserves_unicode_result_when_console_cannot_encode_failure(self) -> None:
        child = mock.Mock(returncode=1)
        child.poll.return_value = 1
        raw = io.BytesIO()
        stream = io.TextIOWrapper(raw, encoding="cp1252", errors="strict")
        message = "formats/CJK: expected '\u4e2d\u6587', Excel '\u6c49\u5b57'"
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            def start(command, **_):
                state = Path(command[command.index("--state") + 1])
                state.write_text(json.dumps({"owned": False, "pid": 12}), encoding="utf-8")
                ORACLE.write_json(output/"results.json", {"schema_version": 1, "mode": "formats",
                    "session": "active", "cleanup_completed": True, "passed": False, "failures": [message]})
                return child
            with contextlib.redirect_stdout(stream), \
                    mock.patch.object(sys, "argv", [str(SCRIPT), "formats", "--active", "--output-dir", directory]), \
                    mock.patch.object(sys, "platform", "win32"), \
                    mock.patch.object(ORACLE.subprocess, "Popen", side_effect=start), \
                    mock.patch.object(ORACLE, "visible_modal", return_value=None):
                self.assertEqual(ORACLE.main(), 1)
            stream.flush()
            text = raw.getvalue().decode("cp1252")
            self.assertIn("FAIL formats/CJK:", text)
            self.assertIn("\\u4e2d\\u6587", text)
            self.assertEqual(json.loads((output/"results.json").read_text(encoding="utf-8"))["failures"], [message])

    def test_main_refuses_nonfinite_timeout_before_starting_excel(self) -> None:
        for timeout in ("nan", "inf", "-inf"):
            with self.subTest(timeout=timeout), mock.patch.object(sys, "argv", [str(SCRIPT), "probe", f"--timeout={timeout}"]), \
                    mock.patch.object(sys, "platform", "win32"), mock.patch.object(ORACLE.subprocess, "Popen") as popen, \
                    mock.patch.object(ORACLE, "write_json", side_effect=AssertionError("validation reached output")):
                with self.assertRaisesRegex(SystemExit, "finite"):
                    ORACLE.main()
                popen.assert_not_called()

    def test_text_function_resolves_once_and_dispatches_unchanged_codes_as_en_us(self) -> None:
        dispatch = mock.Mock()
        dispatch.GetIDsOfNames.return_value = 47
        dispatch.InvokeTypes.side_effect = ["first", "second"]
        localized = mock.Mock(side_effect=AssertionError("localized Text fallback used"))
        app = SimpleNamespace(WorksheetFunction=SimpleNamespace(_oleobj_=dispatch, Text=localized))
        com = SimpleNamespace(DISPATCH_METHOD=1, VT_VARIANT=12)
        codes = ('dddd, mmmm d, yyyy', '[Red]#,##0.00;[Blue](0.00);"d";@')
        with mock.patch.dict(sys.modules, {"pythoncom": com}):
            text = ORACLE.text_function(app)
            self.assertEqual(text(45292, codes[0]), "first")
            self.assertEqual(text(-1.25, codes[1]), "second")
        dispatch.GetIDsOfNames.assert_called_once_with("Text")
        self.assertEqual(dispatch.InvokeTypes.call_args_list, [
            mock.call(47, 1033, 1, (12, 0), ((12, 1), (12, 1)), 45292, codes[0]),
            mock.call(47, 1033, 1, (12, 0), ((12, 1), (12, 1)), -1.25, codes[1]),
        ])
        localized.assert_not_called()

    def test_formats_share_one_en_us_dispatch_for_calibration_and_original_cases(self) -> None:
        cases = [
            {"id": "colored", "code": '[Red]0.00;[Blue](0.00)', "value": -1.25, "expected": "(1.25)"},
            {"id": "quoted", "code": '0.00" dddd"', "value": 2.5, "expected": "2.50 dddd"},
        ]
        text = mock.Mock(side_effect=["1,234.50", "Monday, January 1, 2024", "(1.25)", "2.50 dddd"])
        book = mock.Mock()
        app = SimpleNamespace(UseSystemSeparators=True, DecimalSeparator=",", ThousandsSeparator=" ",
                              Workbooks=SimpleNamespace(Add=lambda: book))
        result = {"failures": []}
        with tempfile.TemporaryDirectory() as directory, \
                mock.patch.object(ORACLE, "read_cases", return_value=cases), \
                mock.patch.object(ORACLE, "text_function", return_value=text) as factory:
            ORACLE.run_formats(app, SimpleNamespace(cases=Path("cases.json"), output_dir=Path(directory)), result)
        factory.assert_called_once_with(app)
        self.assertEqual(text.call_args_list, [
            mock.call(1234.5, "#,##0.00"), mock.call(45292, "dddd, mmmm d, yyyy"),
            mock.call(-1.25, cases[0]["code"]), mock.call(2.5, cases[1]["code"]),
        ])
        self.assertEqual(result["failures"], [])
        self.assertTrue(all(case["passed"] for case in result["cases"]))
        book.Close.assert_called_once_with(SaveChanges=False)

    def test_formats_reject_empty_duplicate_and_nonfinite_cases(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "cases.json"
            case = {"id": "round", "code": "0.00", "value": 1.005, "expected": "1.01"}
            for cases in ([], [case, case], [{**case, "value": float("nan")}], [{**case, "expected": None}]):
                path.write_text(json.dumps({"schema_version": 1, "cases": cases}), encoding="utf-8")
                with self.subTest(cases=cases), self.assertRaises(ValueError):
                    ORACLE.read_cases(path, "formats")

    def test_formats_preserve_zero_false_empty_and_date_system(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "cases.json"
            cases = [{"id": str(i), "code": "General", "value": value, "expected": text, "date_system": "1904"}
                     for i, (value, text) in enumerate(((0, "0"), (False, "FALSE"), ("", "")))]
            path.write_text(json.dumps({"schema_version": 1, "cases": cases}), encoding="utf-8")
            self.assertEqual(ORACLE.read_cases(path, "formats"), cases)

    def test_format_value_kind_cannot_silently_change_a_scalar(self) -> None:
        base = {"id": "typed", "code": "General", "expected": "NaN"}
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "cases.json"
            for kind, value in (("number", "1"), ("number", True), ("text", 1),
                                ("boolean", 0), ("blank", ""), ("nonfinite", "nan"),
                                ("unknown", None), (["number"], 1)):
                path.write_text(json.dumps({"schema_version": 1, "cases": [
                    {**base, "value_kind": kind, "value": value}]}), encoding="utf-8")
                with self.subTest(kind=kind, value=value), self.assertRaisesRegex(ValueError, "typed.*value_kind"):
                    ORACLE.read_cases(path, "formats")

    def test_nonfinite_format_cases_reach_excel_as_numbers_and_refusals_are_failures(self) -> None:
        cases = [{"id": literal, "code": "General", "value_kind": "nonfinite",
                  "value": literal, "expected": literal} for literal in ("NaN", "Infinity", "-Infinity")]
        values = []

        def text(value, code):
            if code == "#,##0.00":
                return "1,234.50"
            if code == "dddd, mmmm d, yyyy":
                return "Monday, January 1, 2024"
            values.append(value)
            if not isinstance(value, float) or math.isfinite(value):
                self.fail("nonfinite numeric input was passed as literal text")
            raise ValueError("Excel refuses this numeric value")

        book = mock.Mock()
        app = SimpleNamespace(UseSystemSeparators=True, DecimalSeparator=",", ThousandsSeparator=" ",
                              WorksheetFunction=SimpleNamespace(Text=text),
                              Workbooks=SimpleNamespace(Add=lambda: book))
        result = {"failures": []}
        with mock.patch.object(ORACLE, "read_cases", return_value=cases), mock.patch.object(ORACLE, "write_json"), \
                mock.patch.object(ORACLE, "text_function", return_value=app.WorksheetFunction.Text):
            ORACLE.run_formats(app, SimpleNamespace(cases=Path("cases.json"), output_dir=Path(".")), result)
        self.assertEqual(len(values), 3)
        self.assertTrue(math.isnan(values[0]))
        self.assertEqual(values[1:], [math.inf, -math.inf])
        self.assertEqual(len(result["failures"]), 3)
        self.assertTrue(all(not answer["passed"] and "error" in answer for answer in result["cases"]))
        book.Close.assert_called_once_with(SaveChanges=False)
        self.assertEqual((app.UseSystemSeparators, app.DecimalSeparator, app.ThousandsSeparator),
                         (True, ",", " "))

    def test_formats_restore_excel_separators_when_open_or_calibration_fails(self) -> None:
        for fail_open in (False, True):
            book = mock.Mock()
            original = {"UseSystemSeparators": True, "DecimalSeparator": ",", "ThousandsSeparator": " "}
            app = SimpleNamespace(**original, WorksheetFunction=SimpleNamespace(Text=lambda *_: "français"),
                                  Workbooks=SimpleNamespace(Add=mock.Mock(
                                      side_effect=RuntimeError("open failed") if fail_open else None,
                                      return_value=book)))
            cases = [{"id": "number", "code": "0.00", "value": 1.0, "expected": "1.00"}]
            with self.subTest(fail_open=fail_open), mock.patch.object(ORACLE, "read_cases", return_value=cases), mock.patch.object(ORACLE, "write_json"), \
                    mock.patch.object(ORACLE, "text_function", return_value=app.WorksheetFunction.Text):
                with self.assertRaises((ValueError, RuntimeError)):
                    ORACLE.run_formats(app, SimpleNamespace(cases=Path("cases.json"), output_dir=Path(".")), {"failures": []})
                self.assertEqual({name: getattr(app, name) for name in original}, original)
                if fail_open:
                    book.Close.assert_not_called()
                else:
                    book.Close.assert_called_once_with(SaveChanges=False)

    def test_formats_persist_separator_recovery_before_first_mutation(self) -> None:
        original = {"UseSystemSeparators": True, "DecimalSeparator": ",", "ThousandsSeparator": " "}
        with tempfile.TemporaryDirectory() as directory:
            result_path = Path(directory)/"results.json"
            class CheckedApplication(SimpleNamespace):
                def __setattr__(self, name, value):
                    if name == "UseSystemSeparators" and value is False:
                        saved = json.loads(result_path.read_text(encoding="utf-8"))
                        assert saved["application_separators_before"] == original
                        raise RuntimeError("first separator setter failed")
                    super().__setattr__(name, value)
            app = CheckedApplication(**original)
            args = SimpleNamespace(cases=Path("cases.json"), output_dir=Path(directory))
            cases = [{"id": "number", "code": "0.00", "value": 1.0, "expected": "1.00"}]
            result = {"failures": [], "cleanup_completed": False}
            with mock.patch.object(ORACLE, "read_cases", return_value=cases), \
                    mock.patch.object(ORACLE, "text_function", return_value=mock.Mock()):
                with self.assertRaisesRegex(RuntimeError, "first separator setter failed"):
                    ORACLE.run_formats(app, args, result)
            self.assertEqual({name: getattr(app, name) for name in original}, original)
            saved = json.loads(result_path.read_text(encoding="utf-8"))
            self.assertFalse(saved["cleanup_completed"])
            self.assertEqual(saved["application_separators_before"], original)

    def test_repair_log_overwrite_is_detected_not_only_new_paths(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            log = root / "error000001.xml"
            log.write_text("old", encoding="utf-8")
            before = ORACLE.repair_logs(root)
            log.write_text("Excel repaired cell records", encoding="utf-8")
            unrelated = root / "notes.xml"
            unrelated.write_text("unrelated", encoding="utf-8")
            changed = ORACLE.changed_repair_logs(before, root)
            self.assertEqual(changed, [log])

    def test_difference_locates_nested_shifted_reference(self) -> None:
        expected = {"sheets": [{"name": "Data", "tables": [{"range": "$I$1:$J$4"}]}]}
        actual = {"sheets": [{"name": "Data", "tables": [{"range": "$H$1:$I$3"}]}]}
        difference = ORACLE.first_difference(expected, actual)
        self.assertIn("$.sheets[0].tables[0].range", difference)
        self.assertIn("$I$1:$J$4", difference)
        self.assertIsNone(ORACLE.first_difference(expected, expected))

    def test_snapshot_normalizes_reference_qualifiers_without_rewriting_quoted_text(self) -> None:
        before, _, left = snapshot_book("source.xlsx")
        after, _, right = snapshot_book("output.xlsx")
        left.HasFormula = right.HasFormula = True
        left.Formula = '=\'[source.xlsx]Data\'!A1+"[source.xlsx]"'
        right.Formula = '=\'[output.xlsx]Data\'!A1+"[source.xlsx]"'
        with mock.patch.object(ORACLE, "special_cells", return_value=[]):
            self.assertEqual(ORACLE.snapshot(before), ORACLE.snapshot(after))
            right.Formula = '=\'[output.xlsx]Data\'!A1+"[output.xlsx]"'
            self.assertNotEqual(ORACLE.snapshot(before), ORACLE.snapshot(after))

    def test_color_conversion_and_refusal(self) -> None:
        self.assertEqual(ORACLE.rgb_to_excel("#123456"), 0x563412)
        self.assertEqual(ORACLE.rgb_to_excel("FFFFFF"), 0xFFFFFF)
        for value in ("#123", "#gg0000", "#00112233", 123):
            with self.subTest(value=value), self.assertRaises(ValueError):
                ORACLE.rgb_to_excel(value)

    def test_fidelity_refuses_missing_rich_and_missing_exchange_half(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with self.assertRaisesRegex(ValueError, "rich"):
                ORACLE.fidelity_inputs(root, None)
            rich = root / "rich.xlsx"
            rich.touch()
            with self.assertRaisesRegex(ValueError, "from-rust.xlsx"):
                ORACLE.fidelity_inputs(root, rich)

    def test_instance_guard_rejects_reused_pid_and_other_images(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            state = Path(directory) / "instance.json"
            identity = {"pid": 12, "created": "old", "image": "C:/Office/EXCEL.EXE"}
            state.write_text(json.dumps({**identity, "owned": True}), encoding="utf-8")
            with mock.patch.object(ORACLE, "process_identity", return_value=identity):
                self.assertEqual(ORACLE.excel_instance(state), identity)
            with mock.patch.object(ORACLE, "process_identity", return_value={**identity, "created": "new"}):
                self.assertIsNone(ORACLE.excel_instance(state))
            identity["image"] = "C:/Windows/notepad.exe"
            state.write_text(json.dumps({**identity, "owned": True}), encoding="utf-8")
            with mock.patch.object(ORACLE, "process_identity", return_value=identity):
                self.assertIsNone(ORACLE.excel_instance(state))

    def test_open_checked_closes_a_repaired_workbook(self) -> None:
        book = mock.Mock(RepairMode=True)
        app = mock.Mock()
        app.Workbooks.Open.return_value = book
        app.Workbooks.Count = 0
        with mock.patch.object(ORACLE, "repair_logs", return_value={}), mock.patch.object(ORACLE, "changed_repair_logs", return_value=[]):
            with self.assertRaisesRegex(AssertionError, "RepairMode"):
                with ORACLE.open_checked(app, Path("bad.xlsx")):
                    self.fail("a repaired workbook entered the check")
        book.Close.assert_called_once_with(SaveChanges=False)
        self.assertEqual(app.Workbooks.Open.call_args.kwargs["UpdateLinks"], 0)
        self.assertEqual(app.Workbooks.Open.call_args.kwargs["CorruptLoad"], 0)

    def test_open_checked_closes_after_checker_failure(self) -> None:
        book = mock.Mock(RepairMode=False)
        app = mock.Mock()
        app.Workbooks.Open.return_value = book
        app.Workbooks.Count = 0
        with mock.patch.object(ORACLE, "repair_logs", return_value={}), mock.patch.object(ORACLE, "changed_repair_logs", return_value=[]):
            with self.assertRaisesRegex(AssertionError, "feature changed"):
                with ORACLE.open_checked(app, Path("valid.xlsx")):
                    raise AssertionError("feature changed")
        book.Close.assert_called_once_with(SaveChanges=False)

    def test_termination_rechecks_identity_on_the_same_handle(self) -> None:
        identity = {"pid": 12, "created": "old", "image": "C:/Office/EXCEL.EXE"}
        api = mock.Mock()
        with mock.patch.dict(sys.modules, {"win32api": api}), mock.patch.object(ORACLE, "excel_instance", return_value=identity):
            with mock.patch.object(ORACLE, "identity_from_handle", return_value={**identity, "created": "new"}):
                ORACLE.stop_owned_excel(Path("instance.json"))
            api.TerminateProcess.assert_not_called()
            with mock.patch.object(ORACLE, "identity_from_handle", return_value=identity):
                ORACLE.stop_owned_excel(Path("instance.json"))
            api.TerminateProcess.assert_called_once_with(api.OpenProcess.return_value, 1)



class GroupedStylesOracle(unittest.TestCase):
    def test_grouped_cases_open_once_and_checkpoint_each_answer(self):
        cases = [
            {"id": "first", "file": "a.xlsx", "sheet": "Styled", "cell": "A1",
             "expected": {"font": "#000000"}},
            {"id": "other", "file": "b.xlsx", "sheet": "Styled", "cell": "B1",
             "expected": {"font": "#000000"}},
            {"id": "last", "file": "a.xlsx", "sheet": "Styled", "cell": "C1",
             "expected": {"font": "#000000"}},
        ]
        cell = SimpleNamespace(Font=SimpleNamespace(Color=0))
        sheet = SimpleNamespace(Range=mock.Mock(return_value=cell))
        book = SimpleNamespace(Worksheets=SimpleNamespace(Item=mock.Mock(return_value=sheet)))
        opened, closed, checkpoints = [], [], []

        @contextlib.contextmanager
        def open_once(_app, path):
            opened.append(path.name)
            try:
                yield book
            finally:
                closed.append(path.name)

        original_write = ORACLE.write_json

        def checkpoint(path, value):
            checkpoints.append((len(value.get("cases", [])),
                                value.get("styles_progress", {}).get("completed")))
            original_write(path, value)

        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            manifest = folder / "cases.json"
            manifest.write_text(json.dumps({"schema_version": 1, "cases": cases}), encoding="utf-8")
            args = SimpleNamespace(cases=manifest, output_dir=folder, exchange=folder)
            result = {"failures": []}
            with mock.patch.object(ORACLE, "open_checked", side_effect=open_once), \
                    mock.patch.object(ORACLE, "write_json", side_effect=checkpoint):
                ORACLE.run_styles(mock.Mock(), args, result)
            self.assertTrue((folder / "results.json").exists(), "case checkpoints were not written")
            saved = json.loads((folder / "results.json").read_text(encoding="utf-8"))
        self.assertEqual(opened, ["a.xlsx", "b.xlsx"])
        self.assertEqual(closed, opened)
        self.assertEqual([answer["id"] for answer in result["cases"]],
                         ["first", "other", "last"])
        self.assertEqual([answer["id"] for answer in saved["cases"]],
                         ["first", "other", "last"])
        self.assertTrue(all(answer["passed"] for answer in result["cases"]))
        self.assertFalse(result["failures"])
        self.assertEqual(list(dict.fromkeys(count for count, progress in checkpoints
                                             if progress == count and count > 0)),
                         [1, 2, 3])

    def test_bad_manifest_refuses_before_opening_any_style_workbook(self):
        cases = [
            {"id": "good", "file": "a.xlsx", "sheet": "Styled", "cell": "A1",
             "expected": {"font": "#000000"}},
            {"id": "bad", "file": "b.xlsx", "sheet": "Styled", "cell": "B1",
             "expected": {"font": "not-a-color"}},
        ]
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            manifest = folder / "cases.json"
            manifest.write_text(json.dumps({"schema_version": 1, "cases": cases}), encoding="utf-8")
            args = SimpleNamespace(cases=manifest, output_dir=folder, exchange=folder)
            with mock.patch.object(ORACLE, "open_checked") as opened:
                with self.assertRaises(ValueError):
                    ORACLE.run_styles(mock.Mock(), args, {"failures": []})
            opened.assert_not_called()


class ActiveDesktopOracle(unittest.TestCase):
    def fake_app(self):
        app = SimpleNamespace(Visible=True, DisplayAlerts=True, EnableEvents=True,
                              AskToUpdateLinks=True, AutomationSecurity=1, Hwnd=123,
                              Version="16.0", Build="20430", OperatingSystem="Windows",
                              LanguageSettings=object(), Ready=True, ActiveWindow=mock.Mock(),
                              Workbooks=Collection([mock.Mock(FullName="C:/user.xlsx")]), Quit=mock.Mock())
        return app

    @contextlib.contextmanager
    def fake_com(self, app):
        client = SimpleNamespace(GetActiveObject=mock.Mock(return_value=app),
                                 DispatchEx=mock.Mock(side_effect=AssertionError("active mode spawned Excel")))
        modules = {"pythoncom": mock.Mock(), "win32com": SimpleNamespace(client=client),
                   "win32com.client": client, "win32process": mock.Mock()}
        modules["win32process"].GetWindowThreadProcessId.return_value = (1, 12)
        with mock.patch.dict(sys.modules, modules), \
                mock.patch.object(ORACLE, "process_identity", return_value={"pid": 12, "created": "old", "image": "C:/Office/EXCEL.EXE"}), \
                mock.patch.object(ORACLE, "wait_ready"), mock.patch.object(ORACLE, "indexed_property", return_value=1):
            yield client

    def test_active_flag_is_explicit_and_isolated_remains_default(self):
        self.assertFalse(ORACLE.parser().parse_args(["probe"]).active)
        self.assertTrue(ORACLE.parser().parse_args(["probe", "--active"]).active)

    def test_active_probe_preserves_existing_workbooks_visibility_and_preferences(self):
        app = self.fake_app()
        before = (app.Visible, app.DisplayAlerts, app.EnableEvents, app.AskToUpdateLinks, app.AutomationSecurity)
        with tempfile.TemporaryDirectory() as directory, self.fake_com(app) as client:
            args = SimpleNamespace(active=True, mode="probe", state=Path(directory)/"instance.json", output_dir=Path(directory))
            self.assertEqual(ORACLE.worker(args), 0)
            result = json.loads((args.output_dir/"results.json").read_text(encoding="utf-8"))
            self.assertEqual(result["session"], "active")
            self.assertTrue(result["cleanup_completed"])
            self.assertFalse(json.loads(args.state.read_text(encoding="utf-8"))["owned"])
        client.GetActiveObject.assert_called_once_with("Excel.Application")
        client.DispatchEx.assert_not_called()
        app.Quit.assert_not_called()
        app.Workbooks.items[0].Close.assert_not_called()
        self.assertEqual((app.Visible, app.DisplayAlerts, app.EnableEvents, app.AskToUpdateLinks, app.AutomationSecurity), before)
        app.ActiveWindow.Activate.assert_not_called()

    def test_active_operation_restores_preferences_and_focus_even_after_failure(self):
        for failure in (False, True):
            app = self.fake_app()
            user_book = app.Workbooks.items[0]
            concurrent_book = mock.Mock(FullName="C:/opened-concurrently.xlsx")
            def run(*_):
                self.assertEqual((app.DisplayAlerts, app.EnableEvents, app.AskToUpdateLinks, app.AutomationSecurity), (False, False, False, 3))
                app.Workbooks.items.append(concurrent_book)
                if failure:
                    raise ValueError("test oracle failure")
            with self.subTest(failure=failure), tempfile.TemporaryDirectory() as directory, self.fake_com(app), \
                    mock.patch.object(ORACLE, "run_formats", side_effect=run):
                args = SimpleNamespace(active=True, mode="formats", state=Path(directory)/"instance.json", output_dir=Path(directory))
                self.assertEqual(ORACLE.worker(args), int(failure))
                result = json.loads((args.output_dir/"results.json").read_text(encoding="utf-8"))
                self.assertTrue(result["cleanup_completed"])
            self.assertEqual((app.Visible, app.DisplayAlerts, app.EnableEvents, app.AskToUpdateLinks, app.AutomationSecurity), (True, True, True, True, 1))
            app.ActiveWindow.Activate.assert_called_once_with()
            user_book.Close.assert_not_called()
            concurrent_book.Close.assert_not_called()
            app.Quit.assert_not_called()

    def test_active_missing_session_never_spawns_a_fallback_instance(self):
        app = self.fake_app()
        with tempfile.TemporaryDirectory() as directory, self.fake_com(app) as client:
            client.GetActiveObject.side_effect = RuntimeError("no active Excel")
            args = SimpleNamespace(active=True, mode="probe", state=Path(directory)/"instance.json", output_dir=Path(directory))
            self.assertEqual(ORACLE.worker(args), 1)
            result = json.loads((args.output_dir/"results.json").read_text(encoding="utf-8"))
            self.assertIn("no active Excel", " ".join(result["failures"]))
        client.DispatchEx.assert_not_called()
        app.Quit.assert_not_called()

    def test_already_open_input_is_refused_without_closing_it(self):
        path = Path("user.xlsx").resolve()
        user_book = mock.Mock(FullName=str(path), RepairMode=False)
        workbooks = Collection([user_book])
        workbooks.Open = mock.Mock(return_value=user_book)
        app = SimpleNamespace(Workbooks=workbooks)
        with mock.patch.object(ORACLE, "repair_logs", return_value={}), \
                mock.patch.object(ORACLE, "changed_repair_logs", return_value=[]):
            with self.assertRaisesRegex(ValueError, "already open"):
                with ORACLE.open_checked(app, path):
                    pass
        workbooks.Open.assert_not_called()
        user_book.Close.assert_not_called()

    def test_active_state_can_never_authorize_process_termination(self):
        identity = {"pid": 12, "created": "old", "image": "C:/Office/EXCEL.EXE"}
        api = mock.Mock()
        with tempfile.TemporaryDirectory() as directory, mock.patch.dict(sys.modules, {"win32api": api}):
            state = Path(directory)/"instance.json"
            state.write_text(json.dumps({**identity, "owned": False}), encoding="utf-8")
            with mock.patch.object(ORACLE, "process_identity", return_value=identity):
                ORACLE.stop_owned_excel(state)
        api.OpenProcess.assert_not_called()
        api.TerminateProcess.assert_not_called()

    def test_active_timeout_never_stops_excel_and_reports_unrestored_state(self):
        child = mock.Mock(returncode=-15)
        child.poll.side_effect = [None, None]
        with tempfile.TemporaryDirectory() as directory, contextlib.redirect_stdout(io.StringIO()), \
                mock.patch.object(sys, "argv", [str(SCRIPT), "probe", "--active", "--timeout", "1", "--output-dir", directory]), \
                mock.patch.object(sys, "platform", "win32"), \
                mock.patch.object(ORACLE.subprocess, "Popen", return_value=child), \
                mock.patch.object(ORACLE.time, "monotonic", side_effect=[0, 2]), \
                mock.patch.object(ORACLE, "visible_modal", return_value=None), \
                mock.patch.object(ORACLE, "stop_owned_excel") as stop:
            self.assertEqual(ORACLE.main(), 1)
            result = json.loads((Path(directory)/"results.json").read_text(encoding="utf-8"))
            self.assertIn("attached", result["cleanup_limitation"])
            self.assertIn("preferences", result["cleanup_limitation"])
            self.assertIn("not terminated", result["cleanup_limitation"])
            stop.assert_not_called()
            child.terminate.assert_called_once_with()


    def test_active_records_recovery_values_before_first_preference_change(self):
        with tempfile.TemporaryDirectory() as directory:
            result_path = Path(directory)/"results.json"
            class CheckedApplication(SimpleNamespace):
                def __setattr__(self, name, value):
                    if name == "DisplayAlerts" and value is False:
                        saved = json.loads(result_path.read_text(encoding="utf-8"))
                        assert saved["application_preferences_before"] == {
                            "DisplayAlerts": True, "EnableEvents": True,
                            "AskToUpdateLinks": True, "AutomationSecurity": 1,
                        }
                    super().__setattr__(name, value)
            app = CheckedApplication(**vars(self.fake_app()))
            with self.fake_com(app), mock.patch.object(ORACLE, "run_formats"):
                args = SimpleNamespace(active=True, mode="formats", state=Path(directory)/"instance.json", output_dir=Path(directory))
                self.assertEqual(ORACLE.worker(args), 0)
            self.assertTrue(json.loads(result_path.read_text(encoding="utf-8"))["cleanup_completed"])

    def test_active_partial_setup_restores_every_attempted_preference(self):
        class PartialFailure(SimpleNamespace):
            def __setattr__(self, name, value):
                super().__setattr__(name, value)
                if name == "AutomationSecurity" and value == 3:
                    raise RuntimeError("setting partially applied")
        app = PartialFailure(**vars(self.fake_app()))
        with tempfile.TemporaryDirectory() as directory, self.fake_com(app), \
                mock.patch.object(ORACLE, "run_formats") as run:
            args = SimpleNamespace(active=True, mode="formats", state=Path(directory)/"instance.json", output_dir=Path(directory))
            self.assertEqual(ORACLE.worker(args), 1)
            result = json.loads((args.output_dir/"results.json").read_text(encoding="utf-8"))
            self.assertTrue(result["cleanup_completed"])
            run.assert_not_called()
        self.assertEqual((app.DisplayAlerts, app.EnableEvents, app.AskToUpdateLinks, app.AutomationSecurity), (True, True, True, 1))
        app.Quit.assert_not_called()


    def test_active_worker_uses_real_readiness_with_existing_workbooks(self):
        app = self.fake_app()
        readiness = ORACLE.wait_ready
        with tempfile.TemporaryDirectory() as directory, self.fake_com(app), \
                mock.patch.object(ORACLE, "wait_ready", readiness), \
                mock.patch.dict(sys.modules, {"pywintypes": SimpleNamespace(com_error=RuntimeError)}), \
                mock.patch.object(ORACLE.time, "monotonic", side_effect=[0, 0, 11]), \
                mock.patch.object(ORACLE.time, "sleep") as sleep:
            args = SimpleNamespace(active=True, mode="probe", state=Path(directory)/"instance.json", output_dir=Path(directory))
            self.assertEqual(ORACLE.worker(args), 0)
            result = json.loads((args.output_dir/"results.json").read_text(encoding="utf-8"))
            self.assertTrue(result["passed"])
            self.assertTrue(result["cleanup_completed"])
            sleep.assert_not_called()
        app.Workbooks.items[0].Close.assert_not_called()
        app.Quit.assert_not_called()

    def test_isolated_readiness_still_requires_zero_workbooks(self):
        for count in (0, 1):
            with self.subTest(count=count), \
                    mock.patch.dict(sys.modules, {"pythoncom": mock.Mock(), "pywintypes": SimpleNamespace(com_error=RuntimeError)}), \
                    mock.patch.object(ORACLE.time, "monotonic", side_effect=[0, 0, 11]), \
                    mock.patch.object(ORACLE.time, "sleep"):
                app = SimpleNamespace(Ready=True, Workbooks=SimpleNamespace(Count=count))
                if count:
                    with self.assertRaisesRegex(RuntimeError, "isolated Excel.*zero workbooks"):
                        ORACLE.wait_ready(app)
                else:
                    ORACLE.wait_ready(app)

    def test_active_readiness_still_requires_ready_application(self):
        with mock.patch.dict(sys.modules, {"pythoncom": mock.Mock(), "pywintypes": SimpleNamespace(com_error=RuntimeError)}), \
                mock.patch.object(ORACLE.time, "monotonic", side_effect=[0, 0, 11]), \
                mock.patch.object(ORACLE.time, "sleep"):
            app = SimpleNamespace(Ready=False, Workbooks=SimpleNamespace(Count=1))
            with self.assertRaisesRegex(RuntimeError, "active Excel.*ready"):
                ORACLE.wait_ready(app, active=True)


class RichSaveOracle(unittest.TestCase):
    @staticmethod
    def observed(top=7.8):
        return {"sheets": [{"name": "Data", "cells": [{"value": "kept", "formula": "=1", "style": {"font": 0}}],
            "comments": [{"cell": "$A$1", "author": "Author", "text": "Note", "from": "$B$1", "to": "$D$4",
                          "geometry": {"Left": 10.2, "Top": top, "Width": 108.0, "Height": 69.0, "Placement": 3}}],
            "charts": [{"geometry": {"Top": 12.0}}]}], "names": ["Data!$A$1"]}

    def run_rich(self, root, *, control=None, candidate_before=None, candidate_after=None, failure=None):
        before = self.observed()
        control = self.observed(7.2) if control is None else control
        candidate_before = before if candidate_before is None else candidate_before
        candidate_after = control if candidate_after is None else candidate_after
        observations = {"source-control": before, "source-candidate": candidate_before,
                        "rich_excel.control.xlsx": control, "rich_excel.pending.xlsx": candidate_after}
        books = {name: mock.Mock(name=name) for name in observations}
        user_book = mock.Mock(name="user workbook")
        app = SimpleNamespace(Workbooks=Collection([user_book]), Quit=mock.Mock(), DisplayAlerts=True)
        opened, closed, saved = [], [], []
        for name, book in books.items():
            def save(path, *, _name=name, **kwargs):
                saved.append((_name, Path(path).name, kwargs))
                Path(path).write_bytes(b"native save")
                if failure == _name + "/save":
                    raise RuntimeError(failure)
            book.SaveAs.side_effect = save
        @contextlib.contextmanager
        def open_book(_app, path):
            name = path.name
            if name == "rich.xlsx":
                name = "source-control" if "source-control" not in opened else "source-candidate"
            opened.append(name)
            try:
                yield books[name]
            finally:
                closed.append(name)
        def observe(book):
            name = next(name for name, held in books.items() if held is book)
            if failure == name + "/snapshot":
                raise RuntimeError(failure)
            return observations[name]
        write = ORACLE.write_json
        def checkpoint(path, value):
            if path.name == "rich_excel.pending.json" and failure == "note/write":
                path.write_text("partial", encoding="utf-8")
                raise RuntimeError(failure)
            write(path, value)
        args = SimpleNamespace(exchange=root, rich=root/"rich.xlsx", workbook=[], output_dir=root)
        result = {"excel": {"version": "16.0", "build": "native-test"}, "failures": [],
                  "workbooks": {str(args.rich.resolve()): before}}
        self.rich_state = result, books, opened, closed, saved, app, user_book
        with mock.patch.object(ORACLE, "fidelity_inputs", return_value=[]), \
                mock.patch.object(ORACLE, "inspect_workbooks"), mock.patch.object(ORACLE, "run_shift"), \
                mock.patch.object(ORACLE, "open_checked", side_effect=open_book), \
                mock.patch.object(ORACLE, "snapshot", side_effect=observe), \
                mock.patch.object(ORACLE, "write_json", side_effect=checkpoint):
            ORACLE.run_fidelity(app, args, result)
        return result

    def test_rich_save_compares_independent_equal_save_legs_and_records_exact_normalization(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            result = self.run_rich(root)
            _, _, opened, closed, saved, app, user = self.rich_state
            self.assertEqual(opened, ["source-control", "rich_excel.control.xlsx",
                                      "source-candidate", "rich_excel.pending.xlsx"])
            self.assertCountEqual(closed, opened)
            self.assertEqual([item[:2] for item in saved], [("source-control", "rich_excel.control.xlsx"),
                                                           ("source-candidate", "rich_excel.pending.xlsx")])
            self.assertEqual(result["rich_save"]["control_before_save"], self.observed())
            self.assertEqual(result["rich_save"]["candidate_before_save"], self.observed())
            self.assertEqual(result["rich_save"]["control"], self.observed(7.2))
            self.assertEqual(result["rich_save"]["candidate"], self.observed(7.2))
            self.assertEqual(result["rich_save"]["normalization"], [{
                "path": "$.sheets[0].comments[0].geometry.Top", "before": 7.8, "after": 7.2}])
            self.assertFalse((root/"rich_excel.xlsx").exists())
            note = json.loads((root/"rich_excel.pending.json").read_text(encoding="utf-8"))
            self.assertEqual(note["native_save"]["normalization"], result["rich_save"]["normalization"])
            self.assertTrue(app.DisplayAlerts)
            user.Close.assert_not_called()
            app.Quit.assert_not_called()

    def test_rich_save_rejects_changed_placement_in_equal_native_saved_legs(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            after = self.observed(7.2)
            after["sheets"][0]["comments"][0]["geometry"]["Placement"] = 2
            with self.assertRaisesRegex(AssertionError, r"source-retention: .*geometry\.Placement"):
                self.run_rich(root, control=after)
            result, _, opened, closed, _, app, user = self.rich_state
            self.assertCountEqual(closed, opened)
            self.assertNotIn("fixture_candidate", result)
            self.assertFalse((root/"rich_excel.pending.xlsx").exists())
            self.assertFalse((root/"rich_excel.pending.json").exists())
            user.Close.assert_not_called()
            app.Quit.assert_not_called()

    def test_native_save_normalization_preserves_unknown_geometry_properties(self):
        before, after = self.observed(), self.observed(7.2)
        before["sheets"][0]["comments"][0]["geometry"]["FutureProperty"] = "retained"
        geometry = after["sheets"][0]["comments"][0]["geometry"]
        geometry["FutureProperty"] = "retained"
        self.assertEqual(ORACLE.native_save_normalization(before, after), [{
            "path": "$.sheets[0].comments[0].geometry.Top", "before": 7.8, "after": 7.2}])
        for change in (lambda: geometry.update(FutureProperty="lost"),
                       lambda: geometry.pop("FutureProperty"),
                       lambda: geometry.update(FutureProperty="retained", Unexpected="added")):
            with self.subTest(change=change):
                change()
                with self.assertRaisesRegex(AssertionError, "source-retention"):
                    ORACLE.native_save_normalization(before, after)

    def test_native_save_normalization_requires_all_four_finite_coordinates(self):
        for name in ("Left", "Top", "Width", "Height"):
            for value in (None, True, "7.2", float("nan"), float("inf")):
                with self.subTest(property=name, value=value):
                    before, after = self.observed(), self.observed(7.2)
                    geometry = after["sheets"][0]["comments"][0]["geometry"]
                    if value is None:
                        geometry.pop(name)
                    else:
                        geometry[name] = value
                    with self.assertRaisesRegex(AssertionError, "source-retention"):
                        ORACLE.native_save_normalization(before, after)

    def test_rich_save_rejects_identical_non_geometry_loss_in_both_native_legs(self):
        changes = {
            "cell": lambda s: s["cells"][0].update(value="lost"),
            "formula": lambda s: s["cells"][0].update(formula="=2"),
            "style": lambda s: s["cells"][0]["style"].update(font=1),
            "text": lambda s: s["comments"][0].update(text="lost"),
            "owner": lambda s: s["comments"][0].update(cell="$A$2"),
            "anchor": lambda s: s["comments"][0].update(to="$E$4"),
            "comment": lambda s: s.update(comments=[]),
            "chart geometry": lambda s: s["charts"][0]["geometry"].update(Top=11.4),
            "geometry missing": lambda s: s["comments"][0].pop("geometry"),
        }
        for label, change in changes.items():
            with self.subTest(change=label), tempfile.TemporaryDirectory() as folder:
                root = Path(folder)
                after = self.observed(7.2)
                change(after["sheets"][0])
                with self.assertRaisesRegex(AssertionError, "source-retention"):
                    self.run_rich(root, control=after)
                self.assertFalse((root/"rich_excel.pending.xlsx").exists())
                self.assertFalse((root/"rich_excel.pending.json").exists())
                self.assertNotIn("fixture_candidate", self.rich_state[0])

    def test_rich_save_rejects_unequal_native_source_and_saved_legs(self):
        for phase, keyword in (("before-save", "candidate_before"), ("after-save", "candidate_after")):
            with self.subTest(phase=phase), tempfile.TemporaryDirectory() as folder:
                with self.assertRaisesRegex(AssertionError, f"Excel-save/{phase}"):
                    self.run_rich(Path(folder), **{keyword: self.observed(8.4)})
                result, _, opened, closed, _, app, user = self.rich_state
                self.assertCountEqual(closed, opened)
                self.assertNotIn("fixture_candidate", result)
                user.Close.assert_not_called()
                app.Quit.assert_not_called()

    def test_rich_save_failures_close_only_owned_books_and_remove_pending_pair(self):
        for failure in ("source-control/save", "rich_excel.control.xlsx/snapshot",
                        "source-candidate/save", "rich_excel.pending.xlsx/snapshot", "note/write"):
            with self.subTest(failure=failure), tempfile.TemporaryDirectory() as folder:
                root = Path(folder)
                with self.assertRaisesRegex(RuntimeError, failure):
                    self.run_rich(root, failure=failure)
                result, _, opened, closed, _, app, user = self.rich_state
                self.assertCountEqual(closed, opened)
                self.assertNotIn("fixture_candidate", result)
                self.assertFalse((root/"rich_excel.pending.xlsx").exists())
                self.assertFalse((root/"rich_excel.pending.json").exists())
                self.assertFalse((root/"rich_excel.xlsx").exists())
                user.Close.assert_not_called()
                app.Quit.assert_not_called()

    def test_rich_save_candidate_is_not_published_when_worker_cleanup_fails(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            result = self.run_rich(root)
            ORACLE.finish_fixture(result, root, False)
            self.assertNotIn("fixture_candidate", result)
            self.assertFalse((root/"rich_excel.pending.xlsx").exists())
            self.assertFalse((root/"rich_excel.pending.json").exists())
            self.assertFalse((root/"rich_excel.xlsx").exists())
            self.assertTrue((root/"rich_excel.control.xlsx").is_file())

    def test_rich_save_refuses_existing_control_before_open_or_supervisor_start(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            control = root/"rich_excel.control.xlsx"
            control.write_bytes(b"prior control")
            with self.assertRaisesRegex(ValueError, "fresh"):
                self.run_rich(root)
            self.assertEqual(self.rich_state[2], [])
            self.assertEqual(control.read_bytes(), b"prior control")
            with mock.patch.object(sys, "argv", [str(SCRIPT), "fidelity", "--output-dir", folder]), \
                    mock.patch.object(sys, "platform", "win32"), \
                    mock.patch.object(ORACLE.subprocess, "Popen") as popen:
                with self.assertRaisesRegex(SystemExit, "fresh"):
                    ORACLE.main()
                popen.assert_not_called()
            self.assertEqual(control.read_bytes(), b"prior control")


class JsonWriteTests(unittest.TestCase):
    def test_write_json_roundtrips_unpaired_utf16_and_emoji_without_com(self):
        payload = {"high": "\ud83d", "low": "\ude00", "emoji": "😀", "plain": "aé"}
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "results.json"
            ORACLE.write_json(path, payload)
            encoded = path.read_text(encoding="utf-8")
            self.assertEqual(json.loads(encoded), payload)
            self.assertIn("\\ud83d", encoded)
            self.assertIn("\\ude00", encoded)
            self.assertFalse(path.with_suffix(".json.tmp").exists())

    def test_windows_reader_lock_retries_atomic_checkpoint(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "state.json"
            path.write_text('{"old": true}', encoding="utf-8")
            replace = Path.replace
            attempts = []

            def locked_once(source, destination):
                attempts.append(source)
                if len(attempts) == 1:
                    self.assertEqual(json.loads(path.read_text(encoding="utf-8")), {"old": True})
                    error = PermissionError("supervisor has checkpoint open")
                    error.winerror = 5
                    raise error
                return replace(source, destination)

            with mock.patch.object(Path, "replace", new=locked_once), mock.patch.object(ORACLE.time, "sleep"):
                ORACLE.write_json(path, {"owned": False, "pid": 42})
            self.assertEqual(len(attempts), 2)
            self.assertEqual(json.loads(path.read_text(encoding="utf-8")), {"owned": False, "pid": 42})
            self.assertFalse(path.with_suffix(".json.tmp").exists())

    def test_permanent_reader_lock_refuses_with_original_checkpoint_intact(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "state.json"
            path.write_text('{"old": true}', encoding="utf-8")
            error = PermissionError("locked")
            error.winerror = 32
            with mock.patch.object(Path, "replace", side_effect=error) as replace, mock.patch.object(ORACLE.time, "sleep"):
                with self.assertRaises(PermissionError) as caught:
                    ORACLE.write_json(path, {"new": True})
            self.assertIs(caught.exception, error)
            self.assertEqual(replace.call_count, 20)
            self.assertEqual(json.loads(path.read_text(encoding="utf-8")), {"old": True})

    def test_unrelated_permission_errors_are_not_retried(self):
        with tempfile.TemporaryDirectory() as folder:
            error = PermissionError("not a Windows sharing conflict")
            with mock.patch.object(Path, "replace", side_effect=error) as replace:
                with self.assertRaises(PermissionError) as caught:
                    ORACLE.write_json(Path(folder) / "state.json", {"owned": False})
            self.assertIs(caught.exception, error)
            self.assertEqual(replace.call_count, 1)


class FunctionOracle(unittest.TestCase):
    def manifest(self, root):
        paths = [root / "functions-input-1900.xlsx", root / "functions-input-1904.xlsx"]
        for path, content in zip(paths, (b"source-1900", b"source-1904")):
            path.write_bytes(content)
        cases = [
            {"id": "calibrate-utf16", "file": paths[0].name, "date_system": "1900", "sheet": "Cases",
             "cell": "B1", "formula": "LEN(\"A😀Z\")", "wire_formula": "LEN(\"A😀Z\")",
             "mode": "calibration", "gates": [],
             "control": {"kind": "equals", "value": 4, "gate": "utf16_version1"}},
            {"id": "gated", "file": paths[0].name, "date_system": "1900", "sheet": "Cases",
             "cell": "B2", "formula": "FIND(\"Z\",\"A😀Z\")", "wire_formula": "FIND(\"Z\",\"A😀Z\")",
             "mode": "cached_equal", "gates": ["utf16_version1"]},
            {"id": "random", "file": paths[0].name, "date_system": "1900", "sheet": "Cases",
             "cell": "B3", "formula": "RAND()", "wire_formula": "RAND()", "mode": "behavior_only", "gates": [],
             "control": {"kind": "number_bounds", "lower": 0, "upper": 1, "upper_inclusive": False}},
            {"id": "date1904", "file": paths[1].name, "date_system": "1904", "sheet": "Cases",
             "cell": "B1", "formula": "DATE(2024,1,1)", "wire_formula": "DATE(2024,1,1)",
             "mode": "cached_equal", "gates": []},
        ]
        payload = {"schema_version": 1, "kind": "p5_function_oracle_inputs", "native_answers": False,
                   "mode_counts": {"cached_equal": 2, "behavior_only": 1, "calibration": 1},
                   "workbooks": [{"file": path.name, "date_system": year,
                                  "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                                  "cases": sum(case["file"] == path.name for case in cases)}
                                 for path, year in zip(paths, ("1900", "1904"))], "cases": cases}
        manifest = root / "cases.json"
        manifest.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        return manifest, paths

    def test_manifest_intake_refuses_tampering_and_fabricated_native_answers(self):
        with tempfile.TemporaryDirectory() as directory:
            manifest, paths = self.manifest(Path(directory))
            self.assertEqual(len(ORACLE.read_function_manifest(manifest)["cases"]), 4)
            paths[0].write_bytes(b"tampered")
            with self.assertRaisesRegex(ValueError, "sha256"):
                ORACLE.read_function_manifest(manifest)
            paths[0].write_bytes(b"source-1900")
            payload = json.loads(manifest.read_text(encoding="utf-8"))
            payload["cases"][1]["expected"] = 3
            manifest.write_text(json.dumps(payload), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "native answer"):
                ORACLE.read_function_manifest(manifest)

    def test_owned_workbook_calculation_records_all_cases_and_stages_both_files(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            manifest, paths = self.manifest(root)
            outputs = root / "output"
            outputs.mkdir()
            app = mock.Mock()
            app.CalculateFull.side_effect = AssertionError("application-wide calculation is forbidden")
            app.WorksheetFunction.IsError.return_value = False
            books = {}
            for path, year in zip(paths, ("1900", "1904")):
                book = mock.Mock(Date1904=year == "1904")
                cells = {}
                for case in json.loads(manifest.read_text(encoding="utf-8"))["cases"]:
                    if case["file"] == path.name:
                        cells[case["cell"]] = mock.Mock(HasFormula=True, Formula="=" + case["wire_formula"],
                                                         Value2=None, Text="")
                sheet = mock.Mock(Name="Cases")
                sheet.Range.side_effect = cells.__getitem__
                source = mock.Mock(Name="Inputs")
                book.Worksheets.Count = 2
                book.Worksheets.Item.side_effect = lambda key, source=source, sheet=sheet: (
                    source if key == 1 else sheet if key in (2, "Cases") else AssertionError(key))
                values = {"B1": 4 if year == "1900" else 45291, "B2": 4, "B3": 0.25}
                sheet.Calculate.side_effect = lambda cells=cells, values=values: [
                    setattr(cell, "Value2", values[address]) for address, cell in cells.items()]
                book.SaveAs.side_effect = lambda target, **kwargs: Path(target).write_bytes(b"cached workbook")
                books[path.resolve()] = (book, source, sheet)

            @contextlib.contextmanager
            def owned_open(_app, path):
                yield books[path.resolve()][0]

            args = SimpleNamespace(cases=manifest, output_dir=outputs)
            result = {"failures": [], "excel": {"version": "16.0", "build": "20430"}}
            saved_cache = lambda _path, selected: {
                (case["sheet"], case["cell"]): {"type": "n", "value_text": str(
                    {"calibrate-utf16": 4, "gated": 4, "random": .25, "date1904": 45291}[case["id"]])} for case in selected}
            with mock.patch.object(ORACLE, "open_checked", side_effect=owned_open), \
                    mock.patch.object(ORACLE, "saved_function_cache", side_effect=saved_cache, create=True), \
                    mock.patch.object(ORACLE, "evaluate_function", return_value=mock.Mock()) as evaluate_factory:
                ORACLE.run_functions(app, args, result)
            self.assertEqual(result["failures"], [])
            self.assertEqual(evaluate_factory.call_args_list, [mock.call(sheet) for _, _, sheet in books.values()])
            self.assertEqual(len(result["functions"]["cases"]), 4)
            self.assertTrue(result["functions"]["calibrations"]["utf16_version1"]["passed"])
            self.assertEqual(result["functions"]["cases"][1]["gate_status"], {"utf16_version1": True})
            self.assertEqual(result["functions"]["cases"][1]["saved_cache"], {"type": "n", "value_text": "4"})
            self.assertEqual(result["functions"]["cases"][1]["native_cache_comparisons"]["actual"]["status"], "equal")
            self.assertEqual([source.Calculate.call_count for _, source, _ in books.values()], [1, 1])
            self.assertEqual([sheet.Calculate.call_count for _, _, sheet in books.values()], [1, 1])
            self.assertTrue(all(book.Calculate.call_count == 0 for book, _, _ in books.values()))
            self.assertEqual([book.SaveAs.call_count for book, _, _ in books.values()], [1, 1])
            app.CalculateFull.assert_not_called()
            self.assertEqual(len(result["function_fixture_candidates"]), 2)
            self.assertTrue(all(Path(path).is_file() for path in result["function_fixture_candidates"]))
            mismatch_output = root / "mismatched-native-cache"
            mismatch_output.mkdir()
            mismatch = {"failures": [], "excel": result["excel"]}
            wrong_cache = lambda _path, selected: {
                (case["sheet"], case["cell"]): {"type": "n", "value_text": "1"} for case in selected}
            with mock.patch.object(ORACLE, "open_checked", side_effect=owned_open), \
                    mock.patch.object(ORACLE, "saved_function_cache", side_effect=wrong_cache), \
                    mock.patch.object(ORACLE, "evaluate_function", return_value=mock.Mock()) as evaluate_factory:
                ORACLE.run_functions(app, SimpleNamespace(cases=manifest, output_dir=mismatch_output), mismatch)
            self.assertIn("saved cache", " ".join(mismatch["failures"]))
            self.assertEqual(len(mismatch["functions"]["cases"]), 4)
            self.assertTrue(all("actual" in answer and "after_save" in answer for answer in mismatch["functions"]["cases"]))
            failed_output = root / "failed"
            failed_output.mkdir()
            books[paths[0].resolve()][2].Calculate.side_effect = lambda: setattr(
                books[paths[0].resolve()][2].Range("B1"), "Value2", 5)
            failed = {"failures": [], "excel": {"version": "16.0", "build": "20430"}}
            with mock.patch.object(ORACLE, "open_checked", side_effect=owned_open), \
                    mock.patch.object(ORACLE, "saved_function_cache", side_effect=saved_cache, create=True), \
                    mock.patch.object(ORACLE, "evaluate_function", return_value=mock.Mock()) as evaluate_factory:
                ORACLE.run_functions(app, SimpleNamespace(cases=manifest, output_dir=failed_output), failed)
            self.assertEqual(len(failed["functions"]["cases"]), 4)
            self.assertFalse(failed["functions"]["calibrations"]["utf16_version1"]["passed"])
            self.assertEqual(failed["functions"]["cases"][1]["gate_status"], {"utf16_version1": False})
            self.assertIn("calibration control failed", " ".join(failed["failures"]))
            ORACLE.finish_function_fixtures(failed, failed_output, False)
            self.assertFalse(list(failed_output.glob("*.xlsx")))

    def test_fixture_publication_requires_clean_worker_and_never_overwrites(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            pending = [root / f"functions-cached-{year}.pending.xlsx" for year in ("1900", "1904")]
            for path in pending:
                path.write_bytes(b"cached workbook")
            result = {"excel": {"version": "16.0", "build": "20430"},
                      "function_fixture_candidates": [str(path) for path in pending],
                      "functions": {"manifest_sha256": "abc", "workbooks": []}}
            ORACLE.finish_function_fixtures(result, root, False)
            self.assertFalse(any(path.exists() for path in pending))
            for path in pending:
                path.write_bytes(b"cached workbook")
            result["function_fixture_candidates"] = [str(path) for path in pending]
            (root / "functions-cached-1904.xlsx").write_bytes(b"existing")
            with self.assertRaisesRegex(ValueError, "existing"):
                ORACLE.finish_function_fixtures(result, root, True)
            self.assertTrue(all(path.exists() for path in pending))
            self.assertEqual((root / "functions-cached-1904.xlsx").read_bytes(), b"existing")

    def test_existing_output_refuses_before_starting_excel_worker(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            manifest, _ = self.manifest(root)
            output = root / "output"
            output.mkdir()
            (output / "functions-cached-1900.xlsx").write_bytes(b"existing")
            with mock.patch.object(sys, "argv", [str(SCRIPT), "functions", "--cases", str(manifest),
                                                "--output-dir", str(output), "--active"]), \
                    mock.patch.object(sys, "platform", "win32"), \
                    mock.patch.object(ORACLE.subprocess, "Popen") as worker:
                with self.assertRaisesRegex(SystemExit, "already exists"):
                    ORACLE.main()
            worker.assert_not_called()

    def test_active_dispatch_preserves_user_workbook_and_restores_preferences(self):
        helper = ActiveDesktopOracle()
        app = helper.fake_app()
        app.CalculateFull = mock.Mock()
        user_book = app.Workbooks.items[0]
        with tempfile.TemporaryDirectory() as directory, helper.fake_com(app), \
                mock.patch.object(ORACLE, "run_functions") as run:
            root = Path(directory)
            args = SimpleNamespace(active=True, mode="functions", state=root / "instance.json",
                                   output_dir=root, cases=root / "cases.json")
            self.assertEqual(ORACLE.worker(args), 0)
            self.assertTrue(json.loads((root / "results.json").read_text(encoding="utf-8"))["cleanup_completed"])
            run.assert_called_once()
        user_book.Close.assert_not_called()
        app.CalculateFull.assert_not_called()
        app.Quit.assert_not_called()
        self.assertEqual((app.DisplayAlerts, app.EnableEvents, app.AskToUpdateLinks, app.AutomationSecurity),
                         (True, True, True, 1))
        app.ActiveWindow.Activate.assert_called_once_with()


class FunctionNativeCache(unittest.TestCase):
    def test_nonfinite_transport_retains_actual_payload_without_classifying_error(self):
        import struct
        for bits in ("7ff000000000000a", "7ff1ccf385ebc8a0", "fff0000000000000"):
            raw = struct.unpack(">d", bytes.fromhex(bits))[0]
            observed = ORACLE.function_value(raw)
            self.assertEqual(observed["variant"], "float")
            self.assertEqual(observed["ieee754_hex"], bits)
            self.assertIsNone(observed["value"])
            self.assertNotIn("error", observed)
            json.dumps(observed, allow_nan=False)
        self.assertNotEqual(ORACLE.function_value(.1 + .2)["ieee754_hex"], ORACLE.function_value(.3)["ieee754_hex"])
        self.assertNotEqual(ORACLE.function_value(True)["variant"], ORACLE.function_value(1)["variant"])

    def test_saved_xml_preserves_serial60_decimal_spelling_errors_and_empty_text(self):
        import zipfile
        from openpyxl import Workbook
        main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        package = "http://schemas.openxmlformats.org/package/2006/relationships"
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "native.xlsx"
            book = Workbook()
            book.active.title = "Cases"
            book.active["A1"] = "=60"
            book.active["A1"].number_format = "yyyy-mm-dd"
            book.save(path)
            with zipfile.ZipFile(path) as archive:
                members = {name: archive.read(name) for name in archive.namelist()}
            members["xl/worksheets/sheet1.xml"] = f'<worksheet xmlns="{main}"><sheetData><row r="1"><c r="A1" s="1"><f>60</f><v>60</v></c><c r="B1"><f>1/3</f><v>0.33333333333333331</v></c><c r="C1" t="e"><f>0^0</f><v>#NUM!</v></c><c r="D1" t="str"><f>""</f><v/></c><c r="E1"><f>1</f><v>INF</v></c></row></sheetData></worksheet>'.encode()
            with zipfile.ZipFile(path, "w") as archive:
                for name, content in members.items():
                    archive.writestr(name, content)
            requested = [{"sheet": "Cases", "cell": f"{column}1"} for column in "ABCDE"]
            caches = ORACLE.saved_function_cache(path, requested)
            self.assertEqual([caches[("Cases", f"{column}1")]["value_text"] for column in "ABCDE"], ["60", "0.33333333333333331", "#NUM!", None, "INF"])
            self.assertEqual(caches[("Cases", "C1")]["type"], "e")
            self.assertEqual(caches[("Cases", "A1")]["attributes"]["s"], "1")
            with self.assertRaisesRegex(ValueError, "missing.*Z1"):
                ORACLE.saved_function_cache(path, [{"sheet": "Cases", "cell": "Z1"}])

    def test_function_cell_preserves_formula_and_formula2_as_separate_native_observations(self) -> None:
        # These strings are a pure transport control, not a prediction about
        # which spelling Excel will return for a particular workbook.
        cell = SimpleNamespace(Value2=1.0, Formula="=A1:A5", Formula2="=@A1:A5", Text="1")
        app = SimpleNamespace(WorksheetFunction=SimpleNamespace(IsError=mock.Mock(return_value=False)))
        observed = ORACLE.function_cell(app, cell, "B2", mock.Mock())
        self.assertEqual(observed["formula"], "=A1:A5")
        self.assertEqual(observed["formula2"], "=@A1:A5")

    def test_function_cell_accepts_only_a_single_boolean_array_transport(self):
        cell = SimpleNamespace(Value2="a" * 32767, Formula='=REPT("a",32767)',
                               Formula2='=REPT("a",32767)', Text="a")
        functions = SimpleNamespace(IsError=mock.Mock(return_value=(False,)))
        app = SimpleNamespace(WorksheetFunction=functions)
        evaluate = mock.Mock(side_effect=AssertionError("do not re-evaluate a proved Boolean"))
        observed = ORACLE.function_cell(app, cell, "B240", evaluate)
        self.assertIs(observed["is_error"], False)
        self.assertEqual(observed["value2"]["value"], "a" * 32767)
        self.assertEqual(observed["is_error_transport"], {"variant": "tuple", "value": [False]})
        functions.IsError.assert_called_once_with(cell)
        evaluate.assert_not_called()
        for invalid in ((), (False, True), (0,), ((False,),), [False], 0, None):
            with self.subTest(transport=invalid):
                functions.IsError.return_value = invalid
                with self.assertRaisesRegex(TypeError, "ISERROR"):
                    ORACLE.function_cell(app, cell, "B240", evaluate)
        evaluate.assert_not_called()

    def test_error_identity_comes_from_owned_cell_functions_not_nan_bits(self):
        import struct
        cell = SimpleNamespace(Value2=struct.unpack(">d", bytes.fromhex("7ff000000000000a"))[0], Formula="=ROUND(1,1)", Formula2="=ROUND(1,1)", Text="#NOMBRE!")
        functions = mock.Mock()
        functions.IsError.return_value = True
        functions.Error_Type.return_value = 6.0
        app = SimpleNamespace(WorksheetFunction=functions)
        evaluate = mock.Mock(side_effect=AssertionError("known method must not fall back"))
        observed = ORACLE.function_cell(app, cell, "B2", evaluate)
        self.assertTrue(observed["is_error"])
        self.assertEqual(observed["error_type"], {"number": 6, "literal": "#NUM!", "source": "worksheet_function"})
        functions.IsError.assert_called_once_with(cell)
        functions.Error_Type.assert_called_once_with(cell)
        functions.Error_Type.side_effect = RuntimeError("new error not classified")
        unknown = ORACLE.function_cell(app, cell, "B2", evaluate)
        self.assertIn("error_type_refusal", unknown)
        self.assertEqual(unknown["value2"]["ieee754_hex"], "7ff000000000000a")

    def test_error_type_fallback_is_owned_and_only_for_missing_com_member(self):
        evaluate = mock.Mock(return_value=6.0)
        sheet = SimpleNamespace(Evaluate=mock.Mock(side_effect=AssertionError("default locale must not be used")))
        cell = SimpleNamespace(Value2=-2146826252, Formula="=0^0", Formula2="=0^0", Text="#NOMBRE!", Parent=sheet)
        functions = SimpleNamespace(IsError=mock.Mock(return_value=True))
        app = SimpleNamespace(WorksheetFunction=functions)
        actual = ORACLE.function_cell(app, cell, "B2", evaluate)
        self.assertEqual(actual["error_type"], {"number": 6, "literal": "#NUM!", "source": "worksheet_evaluate", "lcid": 1033})
        evaluate.assert_called_once_with("ERROR.TYPE(B2)")
        evaluate.reset_mock()
        with self.assertRaisesRegex(ValueError, "address"):
            ORACLE.function_cell(app, cell, "B2)+NOW()", evaluate)
        evaluate.assert_not_called()
        functions.Error_Type = mock.Mock(side_effect=RuntimeError("native method refused"))
        refused = ORACLE.function_cell(app, cell, "B2", evaluate)
        self.assertIn("error_type_refusal", refused)
        evaluate.assert_not_called()
        del functions.Error_Type
        for invalid in (True, float("nan"), -2146826252, 9):
            evaluate.return_value = invalid
            refused = ORACLE.function_cell(app, cell, "B2", evaluate)
            self.assertNotIn("error_type", refused)
            self.assertIn("error_type_refusal", refused)
        functions.IsError.return_value = False
        evaluate.reset_mock()
        ORACLE.function_cell(app, cell, "B2", evaluate)
        evaluate.assert_not_called()

        sheet.Evaluate.assert_not_called()

    def test_cache_type_drives_exact_comparison_and_preserves_limited_error_proof(self):
        import struct
        def actual(value, error=None):
            result = {"value2": ORACLE.function_value(value), "is_error": error is not None}
            if error is not None:
                result["error_type"] = {"number": error, "literal": "#NUM!" if error == 6 else "#DIV/0!"}
            return result
        for value, cached in [(4.0, {"type": "n", "value_text": "1"}),
                              (1, {"type": "b", "value_text": "1"}),
                              (True, {"type": "n", "value_text": "1"}),
                              ("wrong", {"type": "str", "value_text": "right"})]:
            self.assertEqual(ORACLE.compare_function_cache(actual(value), cached)["status"], "different")
        self.assertEqual(ORACLE.compare_function_cache(actual(1 / 3), {"type": "n", "value_text": "0.33333333333333331"})["status"], "equal")
        self.assertEqual(ORACLE.compare_function_cache(actual(""), {"type": "str", "value_text": None})["status"], "equal")
        nan = struct.unpack(">d", bytes.fromhex("7ff000000000000a"))[0]
        self.assertEqual(ORACLE.compare_function_cache(actual(nan, 6), {"type": "e", "value_text": "#NUM!"})["status"], "equal")
        self.assertEqual(ORACLE.compare_function_cache(actual(nan, 6), {"type": "e", "value_text": "#DIV/0!"})["status"], "different")
        unknown = {"value2": ORACLE.function_value(nan), "is_error": True, "error_type_refusal": "new error"}
        self.assertEqual(ORACLE.compare_function_cache(unknown, {"type": "e", "value_text": "#CALC!"})["status"], "cache-authoritative")


    def test_native_half_surrogate_cache_uses_ooxml_x_escape(self):
        # Excel 16.0 build 20430.0, p5-unicode-absent: raw COM Value2
        # contains a lone UTF-16 unit; saved formula t=str uses _xHHHH_.
        for native, saved in [
            ("\ud83d", "_xD83D_"),
            ("\ude00", "_xDE00_"),
            ("AX\ude00Z", "AX_xDE00_Z"),
            ("\U0001f600", "_xD83D__xDE00_"),
            ("_xD83D_", "_x005F_xD83D_"),
        ]:
            actual = {"value2": ORACLE.function_value(native), "is_error": False}
            cache = {"type": "str", "value_text": saved}
            with self.subTest(saved=saved):
                self.assertEqual(ORACLE.compare_function_cache(actual, cache)["status"], "equal")
        actual = {"value2": ORACLE.function_value("A"), "is_error": False}
        self.assertEqual(ORACLE.compare_function_cache(actual,
            {"type": "str", "value_text": "_xD83D_"})["status"], "different")


class FunctionErrorLocale(unittest.TestCase):
    def test_evaluate_resolves_once_and_pins_en_us_for_every_original_expression(self):
        dispatch = mock.Mock()
        dispatch.GetIDsOfNames.return_value = 73
        dispatch.InvokeTypes.side_effect = [2.0, 6.0]
        sheet = SimpleNamespace(_oleobj_=dispatch, Evaluate=mock.Mock(side_effect=AssertionError('default locale is forbidden')))
        com = SimpleNamespace(VT_VARIANT=12, DISPATCH_METHOD=1)
        with mock.patch.dict(sys.modules, pythoncom=com):
            evaluate = ORACLE.evaluate_function(sheet)
            self.assertEqual(evaluate('ERROR.TYPE(B4)'), 2.0)
            self.assertEqual(evaluate('ERROR.TYPE(B5)'), 6.0)
        dispatch.GetIDsOfNames.assert_called_once_with('Evaluate')
        self.assertEqual(dispatch.InvokeTypes.call_args_list, [
            mock.call(73, 1033, 1, (12, 0), ((12, 1),), 'ERROR.TYPE(B4)'),
            mock.call(73, 1033, 1, (12, 0), ((12, 1),), 'ERROR.TYPE(B5)')])
        sheet.Evaluate.assert_not_called()

    def test_error_fallback_uses_the_resolved_owned_evaluator_and_records_its_locale(self):
        evaluate = mock.Mock(return_value=6.0)
        sheet = SimpleNamespace(Evaluate=mock.Mock(side_effect=AssertionError('default locale is forbidden')))
        cell = SimpleNamespace(Value2=-2146826252, Formula='=0^0', Formula2='=0^0', Text='#NOMBRE!', Parent=sheet)
        app = SimpleNamespace(WorksheetFunction=SimpleNamespace(IsError=mock.Mock(return_value=True)))
        result = ORACLE.function_cell(app, cell, 'B5', evaluate)
        self.assertEqual(result['error_type'], {'number': 6, 'literal': '#NUM!', 'source': 'worksheet_evaluate', 'lcid': 1033})
        evaluate.assert_called_once_with('ERROR.TYPE(B5)')
        sheet.Evaluate.assert_not_called()
        evaluate.reset_mock()
        with self.assertRaisesRegex(ValueError, 'address'):
            ORACLE.function_cell(app, cell, 'B5)+1', evaluate)
        evaluate.assert_not_called()



class FunctionFixtureCoverage(unittest.TestCase):
    def test_published_notes_name_only_checked_subset_and_keep_locale_gate_scope(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            pending = [output / f"functions-cached-{year}.pending.xlsx" for year in ("1900", "1904")]
            for path in pending:
                path.write_bytes(b"whole workbook with additional unchecked formulas")
            cases = [{"id": "selected-error-1900", "date_system": "1900", "gate_status": {},
                      "native_cache_comparisons": {"actual": {"status": "equal"}, "after_save": {"status": "equal"}}},
                     {"id": "selected-new-error-1904", "date_system": "1904", "gate_status": {},
                      "native_cache_comparisons": {"actual": {"status": "cache-authoritative"},
                                                   "after_save": {"status": "cache-authoritative"}}}]
            result = {"excel": {"version": "16.0", "build": "20430", "ui_language": 1036},
                      "function_fixture_candidates": [str(path) for path in pending],
                      "functions": {"manifest": "errors-only.json", "manifest_sha256": "abc",
                                    "execution_only": True, "rust_equivalence_checked": False,
                                    "cases": cases, "calibrations": {}, "workbooks": [],
                                    "cache_authoritative_cases": ["selected-new-error-1904"]}}
            ORACLE.finish_function_fixtures(result, output, True)
            for year, expected in zip(("1900", "1904"), cases):
                path = output / f"functions-cached-{year}.xlsx"
                note = json.loads(path.with_suffix('.json').read_text(encoding='utf-8'))
                self.assertEqual(note['sha256'], hashlib.sha256(path.read_bytes()).hexdigest())
                self.assertEqual(note['manifest'], 'errors-only.json')
                self.assertTrue(note['execution_only'])
                self.assertFalse(note['rust_equivalence_checked'])
                self.assertEqual(note['oracle'], 'native worksheet calculation and Excel SaveAs; application locale recorded in excel')
                coverage = note['coverage']
                self.assertEqual(coverage['scope'], 'selected manifest cases; other workbook formulas are not validated')
                self.assertEqual(coverage['manifest_case_count'], 2)
                self.assertEqual(coverage['checked_case_ids'], [expected['id']])
                self.assertEqual(coverage['checked_case_count'], 1)
                self.assertEqual(coverage['case_gate_status'], {expected['id']: {}})
                self.assertEqual(coverage['calibration_status'], {})
                self.assertNotIn('en_us_formula_text', coverage['calibration_status'])
                self.assertEqual(coverage['cache_authoritative_case_ids'], [expected['id']] if year == '1904' else [])
                self.assertEqual(path.read_bytes(), b"whole workbook with additional unchecked formulas")

    def test_failed_full_gate_stays_failed_and_never_publishes_a_cache_fixture(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            pending = output / 'functions-cached-1900.pending.xlsx'
            pending.write_bytes(b'native French TEXT cache')
            result = {'passed': False, 'failures': ['en_us_formula_text failed'],
                      'function_fixture_candidates': [str(pending)],
                      'functions': {'calibrations': {'en_us_formula_text': {'passed': False}},
                                    'cases': [{'id': 'text', 'gate_status': {'en_us_formula_text': False},
                                               'saved_cache': {'type': 'str', 'value_text': 'dddd, janvier d, yyyy'}}]}}
            before = json.dumps(result['functions'], sort_keys=True)
            ORACLE.finish_function_fixtures(result, output, False)
            self.assertFalse(result['passed'])
            self.assertEqual(result['failures'], ['en_us_formula_text failed'])
            self.assertEqual(json.dumps(result['functions'], sort_keys=True), before)
            self.assertNotIn('function_fixtures', result)
            self.assertFalse(list(output.glob('functions-cached-*.xlsx')))






class PivotOracle(unittest.TestCase):

    def test_verify_authored_pivots_requires_hash_and_unique_owned_sheets(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            source = folder / "authored.xlsx"
            source.write_bytes(b"owned input")
            payload = {"schema_version": 1, "kind": "p6_pivot_oracle_inputs", "native_answers": False,
                       "verify_existing": {"file": source.name,
                                           "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                                           "pivots": [{"id": "one", "sheet": "Report", "name": "P6_one"}]}}
            path = folder / "cases.json"
            path.write_text(json.dumps(payload), encoding="utf-8")
            self.assertEqual(ORACLE.read_pivot_manifest(path), payload)
            payload["verify_existing"]["pivots"].append({"id": "two", "sheet": "Report", "name": "P6_two"})
            path.write_text(json.dumps(payload), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "unique"):
                ORACLE.read_pivot_manifest(path)
            payload["verify_existing"]["pivots"].pop()
            payload["verify_existing"]["sha256"] = "0" * 64
            path.write_text(json.dumps(payload), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "sha256"):
                ORACLE.read_pivot_manifest(path)

    def test_pivot_item_order_refuses_unrepresented_sort_modes(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = self.manifest(Path(directory))
            authored = json.loads(path.read_text(encoding="utf-8"))
            for order in ("manual", "ASC", "", None, True, []):
                authored["cases"][0]["item_order"] = order
                path.write_text(json.dumps(authored), encoding="utf-8")
                with self.subTest(order=order), self.assertRaisesRegex(ValueError, "item_order"):
                    ORACLE.read_pivot_manifest(path)

    def test_pivot_order_uses_source_name_and_checks_native_readback(self) -> None:
        calls = []
        class Field:
            SourceName = "Original source"
            AutoSortOrder = 0
            def AutoSort(self, order, source):
                calls.append((order, source))
                self.AutoSortOrder = order
        field = Field()
        pivot = SimpleNamespace(PivotFields=lambda name: field)
        case = {"id": "one", "rows": ["Caption"], "columns": [], "item_order": "descending"}
        self.assertEqual(ORACLE.apply_pivot_order(pivot, case), [{"field": "Original source", "order": 2}])
        self.assertEqual(calls, [(2, "Original source")])
        field.AutoSort = lambda order, source: None
        case["item_order"] = "ascending"
        with self.assertRaisesRegex(AssertionError, "readback"):
            ORACLE.apply_pivot_order(pivot, case)
        del case["item_order"]
        self.assertEqual(ORACLE.apply_pivot_order(pivot, case), [])

    def test_explicit_native_pivot_captions_validate_and_read_back(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = self.manifest(Path(directory))
            authored = json.loads(path.read_text(encoding="utf-8"))
            authored["cases"][0].update(data_caption="Values", grand_total_caption="Grand Total")
            path.write_text(json.dumps(authored), encoding="utf-8")
            self.assertEqual(ORACLE.read_pivot_manifest(path)["cases"][0]["data_caption"], "Values")
            authored["cases"][0]["data_caption"] = ""
            path.write_text(json.dumps(authored), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "caption"):
                ORACLE.read_pivot_manifest(path)
        field = SimpleNamespace(Caption="Valeurs")
        pivot = SimpleNamespace(DataPivotField=field, GrandTotalName="Total général")
        observed = ORACLE.apply_pivot_captions(
            pivot, {"id": "one", "data_caption": "Values", "grand_total_caption": "Grand Total"})
        self.assertEqual(observed, {"data_caption": "Values", "grand_total_caption": "Grand Total"})
    def manifest(self, directory: Path) -> Path:
        source = directory / "pivot-input.xlsx"
        source.write_bytes(b"owned pivot source")
        payload = {"schema_version": 1, "kind": "p6_pivot_oracle_inputs",
                   "native_answers": False,
                   "workbook": {"file": source.name, "sha256": hashlib.sha256(source.read_bytes()).hexdigest()},
                   "source_tables": [{"sheet": "Data", "range": "A1:B3",
                                      "rows": [["Group", "Value"], ["East", 1], ["West", 2]]}],
                   "cases": [{"id": "one", "sheet": "Output", "anchor": "A3",
                              "source": {"sheet": "Data", "range": "A1:B3"},
                              "rows": ["Group"], "columns": [],
                              "values": [{"field": "Value", "aggregate": "sum", "caption": "Value Sum"}],
                              "subtotals": True, "row_grand_totals": True,
                              "column_grand_totals": True}]}
        manifest = directory / "cases.json"
        manifest.write_text(json.dumps(payload), encoding="utf-8")
        return manifest

    def test_manifest_proves_input_hash_and_refuses_unproven_answers(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = self.manifest(Path(directory))
            self.assertEqual(ORACLE.read_pivot_manifest(path)["cases"][0]["id"], "one")
            payload = json.loads(path.read_text(encoding="utf-8"))
            for edit, reason in [
                (lambda data: data["workbook"].update(sha256="0" * 64), "sha256"),
                (lambda data: data["cases"][0].update(expected=3), "native answers"),
                (lambda data: data["cases"][0].update(rows=["Missing"]), "axes"),
                (lambda data: data["cases"].append(dict(data["cases"][0])), "duplicate"),
            ]:
                changed = json.loads(json.dumps(payload))
                edit(changed)
                path.write_text(json.dumps(changed), encoding="utf-8")
                with self.subTest(reason=reason), self.assertRaisesRegex(ValueError, reason):
                    ORACLE.read_pivot_manifest(path)

    def test_manifest_admits_all_eleven_typed_pivot_aggregates(self) -> None:
        names = ("sum", "count", "average", "max", "min", "product",
                 "countNumbers", "stdDev", "stdDevP", "var", "varP")
        with tempfile.TemporaryDirectory() as directory:
            path = self.manifest(Path(directory))
            original = json.loads(path.read_text(encoding="utf-8"))
            for name in names:
                payload = json.loads(json.dumps(original))
                payload["cases"][0]["values"][0]["aggregate"] = name
                path.write_text(json.dumps(payload), encoding="utf-8", newline="\n")
                try:
                    observed = ORACLE.read_pivot_manifest(path)
                except ValueError as error:
                    with self.subTest(aggregate=name):
                        self.fail(f"native oracle refused registered aggregate {name}: {error}")
                else:
                    with self.subTest(aggregate=name):
                        self.assertEqual(observed["cases"][0]["values"][0]["aggregate"], name)

    def test_snapshot_captures_ordered_fields_and_typed_grid(self) -> None:
        cells = [["Group", "Value Sum"], ["East", 1.5]]
        area = SimpleNamespace(Address="$A$3:$B$4", Rows=SimpleNamespace(Count=2),
                               Columns=SimpleNamespace(Count=2),
                               Cells=lambda row, column: SimpleNamespace(Value2=cells[row - 1][column - 1]))
        field = lambda name, position, function=-4157: SimpleNamespace(
            Name=name, Caption=name, Position=position, Function=function,
            NumberFormat="#,##0.0000")
        pivot = SimpleNamespace(Name="P6_one", TableRange2=area,
                                RowFields=Collection([field("Group", 1)]), ColumnFields=Collection(),
                                DataFields=Collection([field("Value Sum", 1)]),
                                RowGrand=True, ColumnGrand=False)
        sheet = SimpleNamespace(Name="Output", PivotTables=lambda: Collection([pivot]))
        book = SimpleNamespace(
            Worksheets=SimpleNamespace(Item=lambda name: sheet),
            Application=SimpleNamespace(WorksheetFunction=SimpleNamespace(IsError=lambda cell: False)))
        observed = ORACLE.pivot_snapshot(book, [{"id": "one", "sheet": "Output"}])
        self.assertEqual(observed[0]["is_error"], [[False, False], [False, False]])
        self.assertEqual(observed[0]["table_range2"], "$A$3:$B$4")
        self.assertEqual(observed[0]["row_fields"],
                         [{"name": "Group", "caption": "Group", "position": 1}])
        self.assertEqual(observed[0]["value2"][1][1]["ieee754_hex"], "3ff8000000000000")
        self.assertEqual(observed[0]["data_fields"][0]["number_format"], "#,##0.0000")

    def test_candidate_is_published_only_after_cleanup(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            pending = output / "pivot-native.pending.xlsx"
            pending.write_bytes(b"fixture")
            result = {"pivot_fixture_candidate": str(pending), "pivots": {}}
            ORACLE.finish_pivot_fixture(result, output, False)
            self.assertFalse(pending.exists())
            self.assertFalse((output / "pivot-native.xlsx").exists())
            pending.write_bytes(b"fixture")
            result["pivot_fixture_candidate"] = str(pending)
            ORACLE.finish_pivot_fixture(result, output, True)
            self.assertEqual((output / "pivot-native.xlsx").read_bytes(), b"fixture")

    def test_changed_refresh_or_save_snapshot_fails_the_oracle(self) -> None:
        expected = [{"id": "one", "table_range2": "$A$3:$B$4",
                     "row_fields": [{"name": "Group", "position": 1}],
                     "value2": [[{"variant": "float", "value": 2.0}]]}]
        ORACLE.assert_pivot_snapshots(expected, json.loads(json.dumps(expected)), "SaveAs")
        for changed in (
            [{**expected[0], "table_range2": "$A$3:$C$4"}],
            [{**expected[0], "value2": [[{"variant": "float", "value": 3.0}]]}],
            [{**expected[0], "row_fields": [{"name": "Other", "position": 1}]}],
            [],
        ):
            with self.subTest(changed=changed), self.assertRaisesRegex(AssertionError, "changed"):
                ORACLE.assert_pivot_snapshots(expected, changed, "RefreshTable")

    def test_mode_refuses_non_attached_execution_before_com(self) -> None:
        with self.assertRaisesRegex(ValueError, "active attachment required"):
            ORACLE.run_pivots(None, SimpleNamespace(active=False), {})


class ExistingPivotOracle(unittest.TestCase):

    def test_independent_candidate_records_raw_and_control_differences(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = self.manifest(root)
            output = root / "out"
            output.mkdir()
            payload = json.loads(path.read_text(encoding="utf-8"))
            payload["compare_existing"] = payload.pop("check_existing")
            payload.pop("modified_members")
            path.write_text(json.dumps(payload), encoding="utf-8")
            class Book:
                def __init__(self, path):
                    pivot = SimpleNamespace(Name="P6_source_order_grand", RefreshTable=lambda: True)
                    self.Worksheets = SimpleNamespace(Item=lambda name: SimpleNamespace(PivotTables=lambda: Collection([pivot])))
                def SaveAs(self, target, **kwargs):
                    Path(target).write_bytes(b"owned")
            @contextlib.contextmanager
            def opened(_app, file):
                yield Book(file)
            snapshot = [{"id": "selected", "table_range2": "$A$3:$B$6", "value2": []}]
            args = SimpleNamespace(active=True, cases=path, output_dir=output)
            def compare(_before, _after, phase):
                if phase == "control/candidate RefreshTable":
                    raise AssertionError("native mismatch")
            with mock.patch.object(ORACLE, "open_checked", side_effect=opened), \
                 mock.patch.object(ORACLE, "pivot_snapshot", return_value=snapshot), \
                 mock.patch.object(ORACLE, "pivot_xml", return_value={}), \
                 mock.patch.object(ORACLE, "assert_pivot_cached", side_effect=AssertionError("raw mismatch")) as raw, \
                 mock.patch.object(ORACLE, "assert_pivot_snapshots", side_effect=compare) as cross:
                result = {}
                with self.assertRaisesRegex(AssertionError, "raw mismatch; native mismatch"):
                    ORACLE.run_pivots(None, args, result)
            self.assertEqual(raw.call_count, 1)
            self.assertEqual([call.args[2] for call in cross.call_args_list].count("control/candidate RefreshTable"), 1)
            self.assertIs(result["candidate_raw_cache_equal"], False)
            self.assertIs(result["control_candidate_equal"], False)

    def test_independent_rust_pivot_manifest_allows_distinct_owned_packages(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            paths = [folder / "pivot-control.xlsx", folder / "pivot-candidate.xlsx"]
            with ORACLE.zipfile.ZipFile(paths[0], "w") as archive:
                archive.writestr("xl/pivotTables/pivotTable1.xml", "<pivotTableDefinition/>")
            with ORACLE.zipfile.ZipFile(paths[1], "w") as archive:
                archive.writestr("xl/pivotTables/pivotTable2.xml", "<pivotTableDefinition/>")
            pair = [{"id": role, "file": source.name,
                     "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                     "sheet": "CaseOrder", "name": "P6_source_order_grand"}
                    for role, source in zip(("control", "candidate"), paths)]
            manifest = folder / "cases.json"
            manifest.write_text(json.dumps({"schema_version": 1, "kind": "p6_pivot_oracle_inputs",
                                            "native_answers": False, "compare_existing": pair}),
                                encoding="utf-8")
            self.assertEqual(ORACLE.read_pivot_manifest(manifest)["compare_existing"], pair)
            pair[1]["sha256"] = "0" * 64
            manifest.write_text(json.dumps({"schema_version": 1, "kind": "p6_pivot_oracle_inputs",
                                            "native_answers": False, "compare_existing": pair}),
                                encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "sha256 mismatch"):
                ORACLE.read_pivot_manifest(manifest)
    def test_label_observation_requires_explicit_true(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = self.manifest(Path(directory))
            original = json.loads(path.read_text(encoding="utf-8"))
            for flag in (False, "true", 1, None):
                payload = dict(original, observe_labels=flag)
                path.write_text(json.dumps(payload), encoding="utf-8")
                with self.subTest(flag=flag), self.assertRaisesRegex(ValueError, "observe_labels"):
                    ORACLE.read_pivot_manifest(path)
            path.write_text(json.dumps(dict(original, observe_labels=True)), encoding="utf-8")
            self.assertIs(ORACLE.read_pivot_manifest(path)["observe_labels"], True)

    def manifest(self, directory: Path) -> Path:
        full = directory / "pivot-full-native.xlsx"
        empty = directory / "pivot-empty-records.xlsx"
        parts = {
            "xl/pivotCache/pivotCacheDefinition5.xml": b'<pivotCacheDefinition recordCount="4"/>',
            "xl/pivotCache/pivotCacheRecords5.xml": b'<pivotCacheRecords count="4"/>',
            "xl/pivotTables/pivotTable5.xml": b'<pivotTableDefinition name="P6_source_order_grand"/>',
        }
        with ORACLE.zipfile.ZipFile(full, "w") as archive:
            for name, contents in parts.items():
                archive.writestr(name, contents)
        with ORACLE.zipfile.ZipFile(empty, "w") as archive:
            for name, contents in parts.items():
                if name.endswith("pivotCacheDefinition5.xml"):
                    contents = b'<pivotCacheDefinition recordCount="0" saveData="0" refreshOnLoad="1"/>'
                elif name.endswith("pivotCacheRecords5.xml"):
                    contents = b'<pivotCacheRecords count="0"/>'
                archive.writestr(name, contents)
        payload = {"schema_version": 1, "kind": "p6_pivot_oracle_inputs", "native_answers": False,
                   "check_existing": [
                       {"id": "control", "file": full.name,
                        "sha256": hashlib.sha256(full.read_bytes()).hexdigest(),
                        "sheet": "CaseOrder", "name": "P6_source_order_grand"},
                       {"id": "candidate", "file": empty.name,
                        "sha256": hashlib.sha256(empty.read_bytes()).hexdigest(),
                        "sheet": "CaseOrder", "name": "P6_source_order_grand"}],
                   "modified_members": ["xl/pivotCache/pivotCacheDefinition5.xml",
                                        "xl/pivotCache/pivotCacheRecords5.xml"]}
        path = directory / "cases.json"
        path.write_text(json.dumps(payload), encoding="utf-8")
        return path

    def test_existing_manifest_proves_hashes_pair_and_exact_member_delta(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = self.manifest(Path(directory))
            self.assertEqual([item["id"] for item in ORACLE.read_pivot_manifest(path)["check_existing"]],
                             ["control", "candidate"])
            original = json.loads(path.read_text(encoding="utf-8"))
            for edit, reason in (
                (lambda data: data["check_existing"][1].update(sha256="0" * 64), "sha256"),
                (lambda data: data["check_existing"][1].update(id="control"), "roles"),
                (lambda data: data["check_existing"][1].update(name="Other"), "same pivot"),
                (lambda data: data.update(modified_members=["xl/workbook.xml"]), "modified members"),
                (lambda data: data["check_existing"][1].update(expected=3), "native answers"),
            ):
                changed = json.loads(json.dumps(original))
                edit(changed)
                path.write_text(json.dumps(changed), encoding="utf-8")
                with self.subTest(reason=reason), self.assertRaisesRegex(ValueError, reason):
                    ORACLE.read_pivot_manifest(path)

    def test_existing_run_refreshes_only_selected_owned_pivots(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = self.manifest(root)
            output = root / "out"
            output.mkdir()
            refreshed = []
            saved = []

            class Book:
                def __init__(self, path):
                    pivot = SimpleNamespace(Name="P6_source_order_grand",
                                            RefreshTable=lambda: refreshed.append(path.name) or True)
                    sheet = SimpleNamespace(PivotTables=lambda: Collection([pivot]))
                    self.Worksheets = SimpleNamespace(Item=lambda name: sheet)

                def SaveAs(self, target, **kwargs):
                    saved.append(Path(target).name)
                    Path(target).write_bytes(b"owned saved workbook")

            @contextlib.contextmanager
            def opened(_app, file):
                yield Book(file)

            snapshot = [{"id": "selected", "table_range2": "$A$3:$B$6", "value2": []}]
            args = SimpleNamespace(active=True, cases=path, output_dir=output)
            with mock.patch.object(ORACLE, "open_checked", side_effect=opened), \
                 mock.patch.object(ORACLE, "pivot_snapshot", return_value=snapshot), \
                 mock.patch.object(ORACLE, "pivot_xml", return_value={}), \
                 mock.patch.object(ORACLE, "assert_pivot_cached", return_value=[]) as cached:
                result = {}
                ORACLE.run_pivots(None, args, result)
            self.assertEqual(refreshed, ["pivot-full-native.xlsx", "pivot-empty-records.xlsx"])
            self.assertEqual(saved, ["pivot-existing-control.pending.xlsx",
                                     "pivot-existing-candidate.pending.xlsx"])
            self.assertEqual(cached.call_count, 1)
            self.assertEqual(len(result["pivot_existing"]), 2)

    def test_existing_candidates_publish_only_after_cleanup(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            pending = [output / f"pivot-existing-{role}.pending.xlsx"
                       for role in ("control", "candidate")]
            for path in pending:
                path.write_bytes(b"owned")
            result = {"pivot_existing_candidates": [str(path) for path in pending],
                      "pivot_existing": []}
            ORACLE.finish_pivot_fixture(result, output, False)
            self.assertTrue(all(not path.exists() for path in pending))
            self.assertFalse(list(output.glob("pivot-existing-*.xlsx")))
            for path in pending:
                path.write_bytes(b"owned")
            result["pivot_existing_candidates"] = [str(path) for path in pending]
            ORACLE.finish_pivot_fixture(result, output, True)
            self.assertEqual(result["pivot_existing_published"],
                             ["pivot-existing-control.xlsx", "pivot-existing-candidate.xlsx"])
            self.assertTrue(all((output / name).read_bytes() == b"owned"
                                for name in result["pivot_existing_published"]))

    def test_failed_first_refresh_registers_both_candidates_for_cleanup(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = self.manifest(root)
            output = root / "out"
            output.mkdir()
            pivot = SimpleNamespace(Name="P6_source_order_grand", RefreshTable=lambda: False)
            sheet = SimpleNamespace(PivotTables=lambda: Collection([pivot]))
            book = SimpleNamespace(Worksheets=SimpleNamespace(Item=lambda name: sheet))

            @contextlib.contextmanager
            def opened(_app, _file):
                yield book

            result = {}
            with mock.patch.object(ORACLE, "open_checked", side_effect=opened), \
                 mock.patch.object(ORACLE, "pivot_snapshot", return_value=[]):
                with self.assertRaisesRegex(AssertionError, "RefreshTable returned false"):
                    ORACLE.run_pivots(None, SimpleNamespace(active=True, cases=path, output_dir=output), result)
            self.assertEqual(result["pivot_existing_candidates"],
                             [str(output / f"pivot-existing-{role}.pending.xlsx")
                              for role in ("control", "candidate")])
            ORACLE.finish_pivot_fixture(result, output, False)



class FunctionTextOriginalLocale(unittest.TestCase):
    def test_original_text_plan_is_explicit_and_never_localizes_currency(self):
        with tempfile.TemporaryDirectory() as directory:
            manifest, _, payload = FunctionTextFormatLocale().manifest(Path(directory))
            case = payload['cases'][0]
            case['wire_formula'] = case['formula'] = 'TEXT(-1234.5,"$#,##0.00;($#,##0.00)")'
            case['control']['value'] = '($1,234.50)'
            case['text_format_oracle'] = {'value_expression': '-1234.5', 'code': '$#,##0.00;($#,##0.00)'}
            manifest.write_text(json.dumps(payload), encoding='utf-8')
            ORACLE.read_function_manifest(manifest)
            sheet = mock.Mock()
            evaluate = mock.Mock(return_value='($1,234.50)')
            with mock.patch.object(ORACLE, 'local_number_format') as convert:
                answer = ORACLE.prepare_function_text_oracle(sheet, case, evaluate)
            self.assertEqual(answer['cell'], case['cell'])
            self.assertEqual(answer['kind'], 'original_native_cache')
            self.assertEqual(answer['en_us_original'], {'variant': 'str', 'value': '($1,234.50)'})
            evaluate.assert_called_once_with(case['wire_formula'])
            sheet.Range.assert_not_called(); convert.assert_not_called()
            # One omitted coordinate is not an implicit original-cell mode.
            case['text_format_oracle']['cell'] = 'C1'
            manifest.write_text(json.dumps(payload), encoding='utf-8')
            with self.assertRaises(ValueError): ORACLE.read_function_manifest(manifest)

    def test_original_text_mode_keeps_value_and_saved_cache_proofs(self):
        for fault in (None, 'en-us', 'after-save', 'original-cache'):
            with self.subTest(fault=fault), tempfile.TemporaryDirectory() as directory:
                root=Path(directory)
                manifest, files, payload=FunctionTextFormatLocale().manifest(root)
                case=payload['cases'][0]
                english='($1,234.50)'
                case['wire_formula']=case['formula']='TEXT(-1234.5,"$#,##0.00;($#,##0.00)")'
                case['control']['value']=english
                case['text_format_oracle']={'value_expression':'-1234.5','code':'$#,##0.00;($#,##0.00)'}
                manifest.write_text(json.dumps(payload),encoding='utf-8')
                output=root/'output';output.mkdir()
                app=mock.Mock();app.WorksheetFunction.IsError.return_value=False
                books={}
                for path,year in zip(files,('1900','1904')):
                    cells={}
                    for item in payload['cases']:
                        if item['file']==path.name:
                            value={'calibrate-utf16':english,'gated':4.0,'random':.25,'date1904':45291.0}[item['id']]
                            cells[item['cell']]=mock.Mock(HasFormula=True,Formula='='+item['wire_formula'],Value2=value,Text=str(value))
                    sheet=mock.Mock(Name='Cases');sheet.Range.side_effect=cells.__getitem__
                    book=mock.Mock(Date1904=year=='1904');book.Worksheets.Count=1
                    book.Worksheets.Item.side_effect=lambda key,sheet=sheet:sheet if key in (1,'Cases') else None
                    def save(target,cells=cells,year=year,**kwargs):
                        Path(target).write_bytes(b'candidate')
                        if fault=='after-save' and year=='1900':cells['B1'].Value2='changed'
                    book.SaveAs.side_effect=save;books[path.resolve()]=(book,sheet,cells,year)
                @contextlib.contextmanager
                def opened(_app,path):
                    book=books[path.resolve()][0]
                    try:yield book
                    finally:book.Close(SaveChanges=False)
                def caches(path,selected):
                    answer={}
                    for item in selected:
                        cell=item['cell'];year=item['date_system']
                        if year=='1900' and cell=='B1':value='lost' if fault=='original-cache' else ('changed' if fault=='after-save' else english);kind='str'
                        else:value=str({'B1':45291.0,'B2':4.0,'B3':.25}[cell]);kind='n'
                        answer[(item['sheet'],cell)]={'type':kind,'value_text':value}
                    return answer
                result={'failures':[],'excel':{'version':'16.0'}}
                with mock.patch.object(ORACLE,'open_checked',side_effect=opened), \
                        mock.patch.object(ORACLE,'saved_function_cache',side_effect=caches), \
                        mock.patch.object(ORACLE,'evaluate_function',return_value=lambda text:'wrong' if fault=='en-us' else english), \
                        mock.patch.object(ORACLE,'local_number_format') as convert:
                    ORACLE.run_functions(app,SimpleNamespace(cases=manifest,output_dir=output),result)
                self.assertEqual(bool(result['failures']),fault is not None,result['failures'])
                convert.assert_not_called();app.CalculateFull.assert_not_called();app.Quit.assert_not_called()
                self.assertTrue(all(book.Close.call_count==1 for book,_,_,_ in books.values()))
                self.assertEqual(books[files[0].resolve()][2]['B1'].Formula,'=TEXT(-1234.5,"$#,##0.00;($#,##0.00)")')
                if fault is None:
                    record=result['functions']['cases'][0]
                    self.assertEqual(record['text_oracle']['kind'],'original_native_cache')
                    self.assertEqual(record['text_oracle']['saved_cache'],record['saved_cache'])
                    self.assertEqual(result['functions']['calibrations']['en_us_formula_text']['scope'],'original_native_cache')
                else:
                    ORACLE.finish_function_fixtures(result,output,False)
                    self.assertFalse(list(output.glob('*.xlsx')))

if __name__ == "__main__":
    unittest.main()

class FunctionTextFormatLocale(unittest.TestCase):
    def test_native_format_conversion_dispatches_invariant_and_local_owners(self):
        dispatch = mock.Mock()
        dispatch.GetIDsOfNames.side_effect = lambda name: {"NumberFormat": 1, "NumberFormatLocal": 2}[name]
        dispatch.InvokeTypes.side_effect = [None, '[$-409]yyyy-mm-dd', '[$-409]aaaa-mm-jj']
        com = SimpleNamespace(DISPATCH_PROPERTYPUT=4, DISPATCH_PROPERTYGET=2, VT_EMPTY=0, VT_VARIANT=12)
        with mock.patch.dict(sys.modules, pythoncom=com):
            converted = ORACLE.local_number_format(SimpleNamespace(_oleobj_=dispatch), '[$-409]yyyy-mm-dd')
        self.assertEqual(converted['number_format_local'], '[$-409]aaaa-mm-jj')
        self.assertEqual(dispatch.InvokeTypes.call_args_list, [
            mock.call(1, 1033, 4, (0, 0), ((12, 1),), '[$-409]yyyy-mm-dd'),
            mock.call(1, 1033, 2, (12, 0), ()), mock.call(2, 1033, 2, (12, 0), ())])

    def manifest(self, root):
        path, files = FunctionOracle().manifest(root)
        value = json.loads(path.read_text(encoding='utf-8'))
        case = value['cases'][0]
        case.update(formula='TEXT(DATE(2024,1,1),"dddd, mmmm d, yyyy")',
                    wire_formula='TEXT(DATE(2024,1,1),"dddd, mmmm d, yyyy")',
                    control={'kind': 'equals', 'value': 'Monday, January 1, 2024', 'gate': 'en_us_formula_text'},
                    text_format_oracle={'value_expression': 'DATE(2024,1,1)', 'code': 'dddd, mmmm d, yyyy',
                                        'cell': 'C1', 'format_cell': 'D1'})
        value['cases'][1]['gates'] = ['en_us_formula_text']
        path.write_text(json.dumps(value), encoding='utf-8')
        return path, files, value

    def test_manifest_refuses_ambiguous_or_conflicting_companion_before_excel(self):
        with tempfile.TemporaryDirectory() as directory:
            path, _, payload = self.manifest(Path(directory))
            ORACLE.read_function_manifest(path)
            original = json.loads(json.dumps(payload))
            payload['cases'][0]['wire_formula']='TEXT(A:A,"dddd, mmmm d, yyyy")'
            payload['cases'][0]['text_format_oracle']['value_expression']='A:A'
            path.write_text(json.dumps(payload),encoding='utf-8')
            with self.assertRaises(ValueError):ORACLE.read_function_manifest(path)
            for key, value in [('code', 'yyyy'), ('cell', 'B2'), ('format_cell', 'C1')]:
                with self.subTest(key=key):
                    payload = json.loads(json.dumps(original))
                    payload['cases'][0]['text_format_oracle'][key] = value
                    path.write_text(json.dumps(payload), encoding='utf-8')
                    with self.assertRaises(ValueError):
                        ORACLE.read_function_manifest(path)

    def test_adapted_gate_keeps_original_cache_checks_and_provenance(self):
        for fault in (None, 'original-cache', 'companion-cache', 'en-us', 'after-save', 'occupied'):
            with self.subTest(fault=fault), tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                manifest, files, payload = self.manifest(root)
                output = root / 'output'; output.mkdir()
                app = mock.Mock(); app.WorksheetFunction.IsError.return_value = False
                original_text, english = 'dddd, janvier d, yyyy', 'Monday, January 1, 2024'
                books = {}
                for path, year in zip(files, ('1900', '1904')):
                    cells = {}
                    for case in payload['cases']:
                        if case['file'] == path.name:
                            value = {'calibrate-utf16': original_text, 'gated': 4.0, 'random': .25, 'date1904': 45291.0}[case['id']]
                            cells[case['cell']] = mock.Mock(HasFormula=True, Formula='='+case['wire_formula'], Value2=value, Text=str(value))
                    if year == '1900':
                        cells['C1'] = mock.Mock(HasFormula=False, Value2='owned data' if fault == 'occupied' else None, MergeCells=False, Text=english)
                        cells['D1'] = mock.Mock(HasFormula=False, Value2=None, MergeCells=False)
                        cells['C1'].Calculate.side_effect = lambda cells=cells: setattr(cells['C1'], 'Value2', english)
                    sheet = mock.Mock(Name='Cases'); sheet.Range.side_effect = cells.__getitem__
                    book = mock.Mock(Date1904=year=='1904'); book.Worksheets.Count=1
                    book.Worksheets.Item.side_effect=lambda key,sheet=sheet: sheet if key in (1,'Cases') else None
                    def save(target, book=book, cells=cells, year=year, **kwargs):
                        Path(target).write_bytes(b'candidate')
                        if fault=='after-save' and year=='1900': cells['C1'].Value2='changed'
                    book.SaveAs.side_effect=save;books[path.resolve()]=(book,sheet,cells,year)
                @contextlib.contextmanager
                def opened(_app,path):
                    book=books[path.resolve()][0]
                    try: yield book
                    finally: book.Close(SaveChanges=False)
                def caches(path,selected):
                    answer={}
                    for case in selected:
                        cell=case['cell']; year=case['date_system']
                        if year=='1900' and cell=='B1': value='lost' if fault=='original-cache' else original_text; kind='str'
                        elif cell=='C1': value='lost' if fault=='companion-cache' else ('changed' if fault=='after-save' else english); kind='str'
                        else: value=str({'B1':45291.0,'B2':4.0,'B3':.25}[cell]);kind='n'
                        answer[(case['sheet'],cell)]={'type':kind,'value_text':value}
                    return answer
                conversion={'input_code':'[$-409]dddd, mmmm d, yyyy','input_lcid':1033,
                            'invariant_readback':'[$-409]dddd, mmmm d, yyyy','number_format_local':'[$-409]jjjj, mmmm j, aaaa'}
                result={'failures':[],'excel':{'version':'16.0'}}
                with mock.patch.object(ORACLE,'open_checked',side_effect=opened), \
                        mock.patch.object(ORACLE,'saved_function_cache',side_effect=caches), \
                        mock.patch.object(ORACLE,'evaluate_function',return_value=lambda text:'wrong' if fault=='en-us' else english), \
                        mock.patch.object(ORACLE,'local_number_format',return_value=conversion,create=True):
                    ORACLE.run_functions(app,SimpleNamespace(cases=manifest,output_dir=output),result)
                self.assertEqual(bool(result['failures']),fault is not None,result['failures'])
                self.assertTrue(all(book.Close.call_count==1 for book,_,_,_ in books.values()))
                self.assertEqual(books[files[0].resolve()][2]['B1'].Formula,'=TEXT(DATE(2024,1,1),"dddd, mmmm d, yyyy")')
                app.CalculateFull.assert_not_called();app.Quit.assert_not_called()
                if fault is None:
                    case=result['functions']['cases'][0]
                    self.assertFalse(case['original_control_passed']);self.assertTrue(case['control_passed'])
                    self.assertEqual(case['saved_cache']['value_text'],original_text)
                    self.assertEqual(case['text_oracle']['saved_cache']['value_text'],english)
                    self.assertEqual(result['functions']['calibrations']['en_us_formula_text']['scope'],'native_number_format_local_companion')
                    ORACLE.finish_function_fixtures(result,output,True)
                    coverage=json.loads((output/'functions-cached-1900.json').read_text())['coverage']
                    self.assertEqual(coverage['text_oracle_case_ids'],['calibrate-utf16'])
                    self.assertIn('original caches retain native locale',coverage['english_cache_scope'])
                else:
                    ORACLE.finish_function_fixtures(result,output,False)
                    self.assertFalse(list(output.glob('*.xlsx')))


class PivotRawCacheComparison(unittest.TestCase):

    def test_native_saved_error_literal_is_authoritative_over_com_transport(self) -> None:
        import zipfile
        from openpyxl import Workbook
        main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "candidate.xlsx"
            native = root / "native-saved.xlsx"
            book = Workbook()
            book.active.title = "Result"
            book.save(source)
            with zipfile.ZipFile(source) as archive:
                members = {name: archive.read(name) for name in archive.namelist()}
            def owned_error(path: Path, literal: str) -> None:
                changed = dict(members)
                changed["xl/worksheets/sheet1.xml"] = (
                    f'<worksheet xmlns="{main}"><sheetData><row r="1">'
                    f'<c r="A1" t="e"><v>{literal}</v></c>'
                    '</row></sheetData></worksheet>').encode()
                with zipfile.ZipFile(path, "w") as archive:
                    for name, content in changed.items():
                        archive.writestr(name, content)
            owned_error(source, "#NUM!")
            owned_error(native, "#NUM!")
            # The native function oracle has observed NaN-like COM error
            # transports. ISERROR + saved t=e/#NUM! carry its identity.
            transport = ORACLE.function_value(float("nan"))
            refreshed = {"table_range2": "$A$1", "value2": [[transport]],
                         "is_error": [[True]]}
            observed = ORACLE.assert_pivot_cached(source, "Result", refreshed, native)
            self.assertEqual(observed[0][0], {"variant": "error", "value": "#NUM!"})
            owned_error(native, "#N/A")
            with self.assertRaisesRegex(AssertionError, "cached error literal"):
                ORACLE.assert_pivot_cached(source, "Result", refreshed, native)

    def test_typed_raw_cache_distinguishes_text_error_date_serial_and_numeric(self) -> None:
        import zipfile
        from openpyxl import Workbook
        main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "typed.xlsx"
            book = Workbook()
            book.active.title = "Result"
            book.save(path)
            with zipfile.ZipFile(path) as archive:
                members = {name: archive.read(name) for name in archive.namelist()}
            members["xl/worksheets/sheet1.xml"] = (
                f'<worksheet xmlns="{main}"><sheetData><row r="1">'
                '<c r="A1"><v>60</v></c><c r="B1" t="inlineStr"><is><t>#N/A</t></is></c>'
                '<c r="C1" t="e"><v>#N/A</v></c><c r="D1"><v>2.25</v></c>'
                '</row></sheetData></worksheet>').encode()
            with zipfile.ZipFile(path, "w") as archive:
                for name, content in members.items():
                    archive.writestr(name, content)
            refreshed = {"table_range2": "$A$1:$D$1", "is_error": [[False, False, True, False]], "value2": [
                [ORACLE.function_value(60.0), ORACLE.function_value("#N/A"),
                 ORACLE.function_value(-2146826246), ORACLE.function_value(2.25)]]}
            observed = ORACLE.assert_pivot_cached(path, "Result", refreshed)
            self.assertEqual(observed[0][0]["ieee754_hex"], ORACLE.function_value(60.0)["ieee754_hex"])
            changed = {**refreshed, "value2": [[ORACLE.function_value(60.0),
                       ORACLE.function_value(-2146826246), ORACLE.function_value(-2146826246),
                       ORACLE.function_value(2.25)]]}
            with self.assertRaisesRegex(AssertionError, "cached"):
                ORACLE.assert_pivot_cached(path, "Result", changed)
            wrong_kind = {**refreshed, "is_error": [[True, False, True, False]],
                          "value2": [[ORACLE.function_value(60.0), ORACLE.function_value("#N/A"),
                                      ORACLE.function_value(-2146826246), ORACLE.function_value(2.25)]]}
            with self.assertRaisesRegex(AssertionError, "error kind"):
                ORACLE.assert_pivot_cached(path, "Result", wrong_kind)

    def test_wrong_rust_cache_fails_before_excel_open_refresh(self) -> None:
        import openpyxl
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "rust-candidate.xlsx"
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "Result"
            sheet["A1"] = "Group"
            sheet["B1"] = "sum Metric"
            sheet["A2"] = "East"
            sheet["B2"] = 7.0
            book.save(path)
            refreshed = {"table_range2": "$A$1:$B$2", "value2": [
                [ORACLE.function_value("Group"), ORACLE.function_value("sum Metric")],
                [ORACLE.function_value("East"), ORACLE.function_value(8.0)],
            ]}
            refreshed["is_error"] = [[False, False], [False, False]]
            with self.assertRaisesRegex(AssertionError, "cached"):
                ORACLE.assert_pivot_cached(path, "Result", refreshed)
