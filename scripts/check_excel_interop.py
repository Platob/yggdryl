#!/usr/bin/env python3
"""Exchange Excel workbooks with openpyxl, both directions.

Self-consistency proves nothing about an exchange format, so this driver runs
the same rows through the Python implementation Excel users reach for:

1. ``cargo test --test interop excel::`` writes ``target/excel-interop/from-rust.xlsx``,
   and ``from-rust-styled.xlsx``: cells typed and styled through
   ``Workbook::set_entry`` and ``Workbook::set_style``.
2. openpyxl reads it with ``data_only=True`` and asserts every cell, the type
   it reads each as, and the number formats of the temporal cells; and it
   reads the styled workbook's bold font, solid fill, outline border, a
   custom number format, a built-in one and a date.
3. openpyxl writes ``target/excel-interop/from-openpyxl.xlsx`` with the same
   rows from Python natives - dates, datetimes, timedeltas, booleans, integers,
   floats and text - and a second sheet, plus ``from-openpyxl-1904.xlsx``
   under the 1904 date system, and ``from-openpyxl-styled.xlsx``: a bold
   header, a percent-formatted number and a second sheet.
   It also writes ``from-openpyxl-fidelity.xlsx``: a styled header (bold
   coloured font, fill, borders, centred), percent and currency formats, a
   column width, row heights, a hidden row holding no cell, merged cells, a
   frozen pane, two conditional formats, a data validation, a hyperlink to
   a URL and one to a place, an autofilter, a table, formulas on both
   sheets, a defined name, a rich inline string, a comment and a chart.
4. The same cargo target runs again; its reading half opens the workbooks
   and asserts the cells and the record rows, and it edits the styled
   workbook - a text cell, and a datetime to the millisecond, a time and a
   duration, whose formats the styles part lacks - and saves it as
   ``from-rust-edited.xlsx``. It opens the fidelity workbook, reads its
   layout, edits one cell of the sheet carrying every feature - so that
   sheet is written again from its cells - and saves it as
   ``from-rust-fidelity.xlsx``. That half prints ``SKIPPED`` when a file
   is absent, and this driver fails on that word, so a skipped half can
   never read as a pass.
   A separate copy, ``from-openpyxl-anchored-fidelity.xlsx``, authors an
   explicit VML cell anchor on the note. Rust opens a row inside every
   range and a column in front, then saves ``from-rust-shifted.xlsx``.
   The original CSS-only note remains a separate atomic Unsupported check:
   the core cannot derive omitted anchors without host font/display metrics.
5. openpyxl reads the edited workbook: the edits are there, each temporal
   under the format and as the Python type the crate's appended styles
   spell, and the bold font and the percent format openpyxl wrote survived
   the crate's save. It reads the fidelity workbook back and asserts every
   feature it wrote survived the crate's save beside the one edited cell,
   and it reads the shifted workbook and asserts every reference moved
   exactly: the formulas on both sheets, the conditional formats, the
   validation, the hyperlinks, the filter, the table, the merge, the frozen
   pane, the row and column sizes, the chart formulas and anchor, and the comment.
6. openpyxl writes a separate two-sheet workbook with a conditional format,
   validation and external hyperlink. Rust cuts their cells across sheets,
   saves the result, then applies its saved inverse after adopting that save.
   openpyxl checks both the moved and restored owners, formulas and link target.
7. openpyxl writes two named tables on one sheet, including a headless
   table and a table with a totals band, blank body row and literal LF/CRLF
   column names. Rust reads both as records and saves a separate cell edit;
   openpyxl then checks the tables and cell facts survived that save. A second
   two-table fixture lets Rust grow and shrink one table's body; openpyxl
   checks the new ListObject ranges, cells and stationary neighbor.
8. Both writers produce a merged two-row header with a nested field,
   literal LF/CRLF labels, dates and an entirely null body row. Each reader
   checks the other writer's values, header geometry and physical row count.

openpyxl is a checking tool of this script only - never a dependency of the
crate. Where the two disagree by design the assertions say so: openpyxl leaves
``_xHHHH_`` escapes unread and writes a literal one as it stands, which Excel
and this crate both read as the character it spells.
"""

from __future__ import annotations

import datetime
import subprocess
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
EXCHANGE = REPO / "rust" / "target" / "excel-interop"

HEADER = ["id", "symbol", "price", "live", "traded", "at", "took"]
LONG = "€ " + "x" * 300

# The rows both sides assert, matching rust/tests/interop/excel.rs.
ROWS: list[tuple[object, ...]] = [
    (1, " 12", 0.0, True, datetime.date(1900, 1, 1), datetime.datetime(2024, 1, 2, 3, 4, 5, 678000), datetime.timedelta(hours=1)),
    (2, "a ", -1.5, False, datetime.date(1900, 2, 28), datetime.datetime(1970, 1, 1), datetime.timedelta(0)),
    (3, "tab\tX_x0041_", 1e-9, True, datetime.date(1900, 3, 1), datetime.datetime(2000, 2, 29, 23, 59, 59), datetime.timedelta(hours=25)),
    (4, None, 1e15, False, None, datetime.datetime(9999, 12, 31), datetime.timedelta(milliseconds=1)),
    (5, LONG, 0.1 + 0.2, True, datetime.date(2024, 2, 29), datetime.datetime(2024, 2, 29, 12), datetime.timedelta(hours=12, milliseconds=500)),
]

# The third symbol the Rust writer wrote, as openpyxl reads it: the crate
# escapes a control character, a carriage return and a literal `_x0041_`
# under ST_Xstring, which openpyxl does not decode.
RUST_ESCAPED_SYMBOL = "SOH_x0001_CR_x000D_X_x005F_x0041_"


def run_cargo(allow_skip: bool) -> str:
    """Run the Rust half, failing on a skipped read unless it is expected."""
    result = subprocess.run(
        ["cargo", "test", "--locked", "--test", "interop", "excel::", "--", "--nocapture"],
        cwd=REPO / "rust",
        capture_output=True,
        text=True,
        check=False,
    )
    output = result.stdout + result.stderr
    if result.returncode != 0:
        raise SystemExit(f"cargo test failed:\n{output}")
    if not allow_skip and "SKIPPED" in output:
        raise SystemExit(f"the Rust reading half skipped:\n{output}")
    return output


def expect(condition: bool, message: str) -> None:
    if not condition:
        raise SystemExit(message)


def read_with_openpyxl(path: Path) -> None:
    """Assert the workbook this crate wrote, through the reference reader."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True)
    expect(workbook.sheetnames == ["Sheet1"], f"unexpected sheets: {workbook.sheetnames}")
    sheet = workbook["Sheet1"]
    header = [cell.value for cell in sheet[1]]
    expect(header == HEADER, f"unexpected header: {header}")
    for index, row in enumerate(ROWS, start=2):
        cells = [cell for cell in sheet[index]]
        values = [cell.value for cell in cells]
        expected = list(row)
        if row[0] == 3:
            expected[1] = RUST_ESCAPED_SYMBOL
        for name, got, want, cell in zip(HEADER, values, expected, cells):
            if name == "traded" and want is not None:
                expect(isinstance(got, datetime.datetime), f"row {index} {name}: {got!r} is not a datetime")
                got = got.date()
                expect(cell.number_format == "yyyy-mm-dd", f"row {index} {name}: format {cell.number_format!r}")
            if name == "at":
                expect(isinstance(got, datetime.datetime), f"row {index} {name}: {got!r} is not a datetime")
                expect("hh:mm:ss" in cell.number_format, f"row {index} {name}: format {cell.number_format!r}")
            if name == "took":
                expect(isinstance(got, datetime.timedelta), f"row {index} {name}: {got!r} is not a timedelta")
                expect(cell.number_format == "[h]:mm:ss", f"row {index} {name}: format {cell.number_format!r}")
            if name == "id":
                # A whole number is written as Excel writes it, and reads as
                # the int openpyxl gives one.
                expect(isinstance(got, int) and not isinstance(got, bool), f"row {index} id: {got!r}")
            if name == "live":
                expect(isinstance(got, bool), f"row {index} live: {got!r}")
            if name == "price" and want == 1e15:
                got = float(got)
            expect(got == want, f"row {index} {name}: expected {want!r}, got {got!r}")


STYLED_SHEET = "Styled"
EDITED_TEXT = "edited by yggdryl"


def read_styled_from_rust(path: Path) -> None:
    """Assert the styles this crate's ``set_style`` wrote, through openpyxl."""
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    sheet = workbook["Styled"]
    header = sheet["A1"]
    expect(header.value == "Total", f"A1: {header.value!r}")
    expect(header.font.b is True, f"A1 is not bold: {header.font!r}")
    expect(header.fill.fill_type == "solid", f"A1 fill: {header.fill!r}")
    expect(header.fill.fgColor.rgb == "FFFFFF00", f"A1 fill colour: {header.fill.fgColor!r}")
    for side in ("left", "right", "top", "bottom"):
        edge = getattr(header.border, side)
        expect(edge.style == "thin", f"A1 {side} border: {edge!r}")
    custom = sheet["B1"]
    expect(custom.value == 1234.5, f"B1: {custom.value!r}")
    expect(custom.number_format == '#,##0.000 "kg"', f"B1 format: {custom.number_format!r}")
    expect(custom.font.b is not True, "B1 took the header's font")
    percent = sheet["C1"]
    expect(percent.value == 0.125 and percent.number_format == "0.00%", f"C1: {percent.value!r} {percent.number_format!r}")
    day = sheet["D1"]
    expect(day.is_date and day.value == datetime.datetime(2024, 1, 2), f"D1: {day.value!r} {day.number_format!r}")
    # Only the header's fill is solid: the reserved fills stay in their place.
    expect(sheet["B1"].fill.fill_type is None, f"B1 fill: {sheet['B1'].fill!r}")


def write_styled_with_openpyxl(path: Path) -> None:
    """Write a workbook whose cells carry a font and a number format."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = STYLED_SHEET
    sheet["A1"] = "Header"
    sheet["A1"].font = Font(bold=True)
    sheet["A2"] = "rate"
    sheet["B2"] = 0.25
    sheet["B2"].number_format = "0.00%"
    other = workbook.create_sheet("Other")
    other["A1"] = "untouched"
    workbook.save(path)


def read_edited_with_openpyxl(path: Path) -> None:
    """Assert the styled workbook this crate edited kept what openpyxl wrote."""
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    expect(workbook.sheetnames == [STYLED_SHEET, "Other"], f"unexpected sheets: {workbook.sheetnames}")
    sheet = workbook[STYLED_SHEET]
    expect(sheet["C3"].value == EDITED_TEXT, f"the edit is missing: {sheet['C3'].value!r}")
    expect(sheet["A1"].value == "Header", f"A1: {sheet['A1'].value!r}")
    expect(sheet["A1"].font.bold is True, f"A1 lost its bold font: {sheet['A1'].font!r}")
    expect(sheet["B2"].value == 0.25, f"B2: {sheet['B2'].value!r}")
    expect(sheet["B2"].number_format == "0.00%", f"B2 lost its format: {sheet['B2'].number_format!r}")
    expect(sheet["A2"].value == "rate", f"A2: {sheet['A2'].value!r}")
    expect(workbook["Other"]["A1"].value == "untouched", "the other sheet changed")
    # The temporal cells the crate added, under the styles it appended.
    moment = sheet["D3"]
    expect(
        moment.number_format == "yyyy-mm-dd hh:mm:ss.000",
        f"D3's format: {moment.number_format!r}",
    )
    expect(isinstance(moment.value, datetime.datetime), f"D3: {moment.value!r}")
    expect(
        abs(moment.value - datetime.datetime(2024, 1, 1, 0, 0, 0, 250000)) < datetime.timedelta(milliseconds=1),
        f"D3: {moment.value!r}",
    )
    clock = sheet["E3"]
    expect(clock.number_format == "h:mm:ss", f"E3's format: {clock.number_format!r}")
    expect(clock.value == datetime.time(12, 34, 56), f"E3: {clock.value!r}")
    elapsed = sheet["F3"]
    expect(elapsed.number_format == "[h]:mm:ss", f"F3's format: {elapsed.number_format!r}")
    expect(elapsed.value == datetime.timedelta(hours=36), f"F3: {elapsed.value!r}")


FIDELITY_ROWS = [("Apple", 0.5, 1.25), ("Pear", 1.5, 2.5), ("Fig", 2.0, 3.75), ("Kiwi", 0.25, 0.5)]
FIDELITY_EDIT = "Kiwi, edited"
# What the crate puts in `Report!A3`, so that sheet is written again too.
FIDELITY_REPORT_EDIT = "written again"
HEADER_FONT = "FF1F4E79"
HEADER_FILL = "FFDDEBF7"


def write_fidelity_with_openpyxl(path: Path) -> None:
    """Write a workbook carrying every worksheet feature the crate keeps."""
    from openpyxl import Workbook
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.chart import BarChart, Reference
    from openpyxl.comments import Comment
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.hyperlink import Hyperlink
    from openpyxl.worksheet.table import Table, TableStyleInfo

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Name", "Rate", "Price", "Total"])
    for row, (name, rate, price) in enumerate(FIDELITY_ROWS, start=2):
        sheet.append([name, rate, price, f"=B{row}*C{row}"])
    thin = Side(style="thin", color="FF000000")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color=HEADER_FONT)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        cell.alignment = Alignment(horizontal="center")
    for row in range(2, 6):
        sheet.cell(row=row, column=2).number_format = "0.00%"
        sheet.cell(row=row, column=3).number_format = '"$"#,##0.00'
    sheet.column_dimensions["A"].width = 18
    sheet.row_dimensions[2].height = 24
    sheet.row_dimensions[8].height = 30
    sheet.row_dimensions[8].hidden = True
    sheet["F1"] = "merged"
    sheet.merge_cells("F1:G1")
    sheet.freeze_panes = "B2"
    sheet.conditional_formatting.add(
        "B2:B5",
        CellIsRule(operator="greaterThan", formula=["1"], fill=PatternFill("solid", bgColor="FFFFC7CE")),
    )
    sheet.conditional_formatting.add(
        "C2:C5",
        ColorScaleRule(start_type="min", start_color="FFF8696B", end_type="max", end_color="FF63BE7B"),
    )
    validation = DataValidation(type="whole", operator="between", formula1="0", formula2="100", allow_blank=True)
    validation.add("E2:E5")
    sheet.add_data_validation(validation)
    sheet["A2"].hyperlink = "https://example.com/apple"
    sheet["A3"].hyperlink = Hyperlink(ref="A3", location="Report!A1", display="Report")
    sheet.auto_filter.ref = "A1:D5"
    for reference, value in (("H1", "Item"), ("I1", "Cost"), ("H2", "Box"), ("I2", 3), ("H3", "Tape"), ("I3", 1)):
        sheet[reference] = value
    table = Table(displayName="Costs", ref="H1:I3")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    sheet["A4"].comment = Comment("Keep this note with the pear", "yggdryl")
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=3, min_row=1, max_row=5), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=5))
    chart.anchor = TwoCellAnchor(_from=AnchorMarker(col=10, row=1),
                                 to=AnchorMarker(col=18, row=11))
    sheet.add_chart(chart)
    sheet["A7"] = CellRichText(TextBlock(InlineFont(b=True), "bold"), " plain")
    report = workbook.create_sheet("Report")
    report["A1"] = "=SUM(Data!C2:C5)"
    report["A2"] = "=Rate*Data!C2"
    workbook.defined_names["Rate"] = DefinedName("Rate", attr_text="0.07")
    workbook.save(path)


# Explicit authored geometry, also observed in Excel 16.0 build 20430.0.
# This is a second fixture, never a CSS-to-cell-coordinate conversion.
FIDELITY_NOTE_ANCHOR = (0, 79, 0, 2, 2, 29, 3, 11)
FIDELITY_SHIFTED_NOTE_ANCHOR = (1, 79, 1, 2, 3, 29, 4, 11)


def fidelity_note_anchor(path: Path) -> tuple[int, ...] | None:
    """Read the one note's standard VML cell anchor; openpyxl ignores it."""
    from zipfile import ZipFile
    from xml.etree import ElementTree

    namespace = "urn:schemas-microsoft-com:office:excel"
    notes = []
    with ZipFile(path) as source:
        for name in source.namelist():
            if name.endswith(".vml"):
                root = ElementTree.fromstring(source.read(name))
                notes.extend(node for node in root.iter(f"{{{namespace}}}ClientData")
                             if node.get("ObjectType") == "Note")
    expect(len(notes) == 1, f"expected one VML note in {path.name}, got {len(notes)}")
    anchors = notes[0].findall(f"{{{namespace}}}Anchor")
    expect(len(anchors) <= 1, f"ambiguous note anchor in {path.name}")
    if not anchors:
        return None
    values = tuple(int(value.strip()) for value in (anchors[0].text or "").split(","))
    expect(len(values) == 8, f"malformed note anchor in {path.name}: {values}")
    return values


def write_anchored_fidelity(source_path: Path, target_path: Path) -> None:
    """Author one explicit note anchor, preserving every other package member."""
    from zipfile import ZipFile
    from xml.etree import ElementTree

    expect(source_path.resolve() != target_path.resolve(), "the CSS-only fixture must remain separate")
    expect(fidelity_note_anchor(source_path) is None, "the raw openpyxl note unexpectedly has an anchor")
    namespace = "urn:schemas-microsoft-com:office:excel"
    changed = []
    with ZipFile(source_path) as source, ZipFile(target_path, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename.endswith(".vml"):
                root = ElementTree.fromstring(data)
                notes = [node for node in root.iter(f"{{{namespace}}}ClientData")
                         if node.get("ObjectType") == "Note"]
                if notes:
                    expect(len(notes) == 1, "the fixture must contain exactly one note")
                    anchor = ElementTree.Element(f"{{{namespace}}}Anchor")
                    anchor.text = ", ".join(map(str, FIDELITY_NOTE_ANCHOR))
                    notes[0].insert(2, anchor)
                    data = ElementTree.tostring(root, encoding="utf-8")
                    changed.append(info.filename)
            target.writestr(info, data)
    expect(len(changed) == 1, f"anchored fixture changed unexpected parts: {changed}")
    with ZipFile(source_path) as source, ZipFile(target_path) as target:
        expect(source.namelist() == target.namelist(), "anchoring changed package membership")
        expect([name for name in source.namelist() if source.read(name) != target.read(name)] == changed,
               "anchoring changed another package member")
    expect(fidelity_note_anchor(source_path) is None, "anchoring altered the raw fixture")
    expect(fidelity_note_anchor(target_path) == FIDELITY_NOTE_ANCHOR, "the explicit anchor was not authored")


def assert_fidelity_objects(sheet, *, shifted: bool) -> None:
    """Check drawing references and the note attached to a moved cell."""
    address = "B5" if shifted else "A4"
    comments = [cell.coordinate for row in sheet for cell in row if cell.comment is not None]
    expect(comments == [address], f"comment locations: {comments!r}")
    comment = sheet[address].comment
    expect(comment is not None, f"{address} lost its comment")
    expect((comment.text, comment.author) == ("Keep this note with the pear", "yggdryl"),
           f"{address}'s comment: {comment!r}")
    expect(len(sheet._charts) == 1, f"charts: {len(sheet._charts)}")
    chart = sheet._charts[0]
    expect(len(chart.series) == 1, f"chart series: {len(chart.series)}")
    series = chart.series[0]
    expected = ("'Data'!D1", "'Data'!$D$2:$D$6", "'Data'!$B$2:$B$6") if shifted else (
        "'Data'!C1", "'Data'!$C$2:$C$5", "'Data'!$A$2:$A$5")
    actual = (series.tx.strRef.f, series.val.numRef.f, series.cat.numRef.f)
    # Optional quoting of this simple sheet name has no reference meaning;
    # absolute markers, endpoints and the single-cell title remain exact.
    expect(tuple(value.replace("'Data'!", "Data!") for value in actual) ==
           tuple(value.replace("'Data'!", "Data!") for value in expected), f"chart formulas: {actual!r}")
    anchor = chart.anchor
    actual = (anchor._from.col, anchor._from.row, anchor.to.col, anchor.to.row)
    expect(actual == ((11, 1, 19, 12) if shifted else (10, 1, 18, 11)),
           f"chart anchor: {actual!r}")


def read_fidelity_with_openpyxl(path: Path) -> None:
    """Assert every feature openpyxl wrote survived the crate's edit and save."""
    from openpyxl import load_workbook
    from openpyxl.cell.rich_text import CellRichText

    workbook = load_workbook(path, rich_text=True)
    expect(workbook.sheetnames == ["Data", "Report"], f"unexpected sheets: {workbook.sheetnames}")
    sheet = workbook["Data"]
    assert_fidelity_objects(sheet, shifted=False)
    expect(sheet["A5"].value == FIDELITY_EDIT, f"the edit is missing: {sheet['A5'].value!r}")
    expect(sheet["E2"].value == 50, f"the added cell is missing: {sheet['E2'].value!r}")
    for row, (name, rate, price) in enumerate(FIDELITY_ROWS, start=2):
        if row != 5:
            expect(sheet.cell(row=row, column=1).value == name, f"A{row}: {sheet.cell(row=row, column=1).value!r}")
        expect(sheet.cell(row=row, column=2).value == rate, f"B{row}: {sheet.cell(row=row, column=2).value!r}")
        expect(sheet.cell(row=row, column=3).value == price, f"C{row}: {sheet.cell(row=row, column=3).value!r}")
        expect(sheet.cell(row=row, column=4).value == f"=B{row}*C{row}", f"D{row}: {sheet.cell(row=row, column=4).value!r}")
        expect(sheet.cell(row=row, column=2).number_format == "0.00%", f"B{row}'s format")
        expect(sheet.cell(row=row, column=3).number_format == '"$"#,##0.00', f"C{row}'s format")
    for cell in sheet[1][:4]:
        expect(cell.font.bold is True, f"{cell.coordinate} lost its bold font")
        expect(cell.font.color is not None and cell.font.color.rgb == HEADER_FONT, f"{cell.coordinate}'s font colour: {cell.font.color!r}")
        expect(cell.fill.fgColor.rgb == HEADER_FILL, f"{cell.coordinate}'s fill: {cell.fill.fgColor!r}")
        expect(cell.border.left.style == "thin" and cell.border.bottom.style == "thin", f"{cell.coordinate}'s border")
        expect(cell.alignment.horizontal == "center", f"{cell.coordinate}'s alignment")
    expect(sheet.column_dimensions["A"].width == 18, f"column A's width: {sheet.column_dimensions['A'].width!r}")
    expect(sheet.row_dimensions[2].height == 24, f"row 2's height: {sheet.row_dimensions[2].height!r}")
    expect(sheet.row_dimensions[8].height == 30, f"row 8's height: {sheet.row_dimensions[8].height!r}")
    expect(sheet.row_dimensions[8].hidden is True, "row 8 is no longer hidden")
    expect([str(merged) for merged in sheet.merged_cells.ranges] == ["F1:G1"], f"merges: {sheet.merged_cells.ranges!r}")
    expect(sheet["F1"].value == "merged", f"F1: {sheet['F1'].value!r}")
    expect(sheet.freeze_panes == "B2", f"freeze panes: {sheet.freeze_panes!r}")
    rules = {
        str(formatting.sqref): [(rule.type, list(rule.formula or [])) for rule in formatting.rules]
        for formatting in sheet.conditional_formatting
    }
    expect(
        rules == {"B2:B5": [("cellIs", ["1"])], "C2:C5": [("colorScale", [])]},
        f"conditional formats: {rules!r}",
    )
    validations = sheet.data_validations.dataValidation
    expect(len(validations) == 1, f"data validations: {validations!r}")
    validation = validations[0]
    expect(
        (validation.type, str(validation.sqref), validation.formula1, validation.formula2)
        == ("whole", "E2:E5", "0", "100"),
        f"the data validation: {validation!r}",
    )
    expect(sheet["A2"].hyperlink is not None and sheet["A2"].hyperlink.target == "https://example.com/apple", f"A2's link: {sheet['A2'].hyperlink!r}")
    expect(sheet["A3"].hyperlink is not None and sheet["A3"].hyperlink.location == "Report!A1", f"A3's link: {sheet['A3'].hyperlink!r}")
    expect(sheet.auto_filter.ref == "A1:D5", f"autofilter: {sheet.auto_filter.ref!r}")
    expect(set(sheet.tables) == {"Costs"} and sheet.tables["Costs"].ref == "H1:I3", f"tables: {dict(sheet.tables.items())!r}")
    rich = sheet["A7"].value
    expect(isinstance(rich, CellRichText), f"A7 lost its runs: {rich!r}")
    runs = [(getattr(block, "font", None) is not None and bool(block.font.b), str(block.text if hasattr(block, "text") else block)) for block in rich]
    expect(runs == [(True, "bold"), (False, " plain")], f"A7's runs: {runs!r}")
    report = workbook["Report"]
    expect(report["A1"].value == "=SUM(Data!C2:C5)", f"Report!A1: {report['A1'].value!r}")
    expect(report["A2"].value == "=Rate*Data!C2", f"Report!A2: {report['A2'].value!r}")
    expect(report["A3"].value == FIDELITY_REPORT_EDIT, f"Report!A3: {report['A3'].value!r}")
    expect("Rate" in workbook.defined_names and workbook.defined_names["Rate"].attr_text == "0.07", "the defined name is gone")


def read_shifted_with_openpyxl(path: Path) -> None:
    """Assert every reference moved with the row and the column the crate opened.

    The crate opened a row at row 3 of ``Data`` - inside every range the
    sheet states - then a column at column A, in front of every column and
    inside the frozen one.
    """
    from openpyxl import load_workbook
    from openpyxl.cell.rich_text import CellRichText

    workbook = load_workbook(path, rich_text=True)
    expect(workbook.sheetnames == ["Data", "Report"], f"unexpected sheets: {workbook.sheetnames}")
    sheet = workbook["Data"]
    assert_fidelity_objects(sheet, shifted=True)
    expect(fidelity_note_anchor(path) == FIDELITY_SHIFTED_NOTE_ANCHOR,
           f"the note box did not follow its owner: {fidelity_note_anchor(path)}")
    header = [cell.value for cell in sheet[1][:5]]
    expect(header == [None, "Name", "Rate", "Price", "Total"], f"the header row: {header!r}")
    expect(all(cell.value is None for cell in sheet[3]), f"the opened row is not empty: {[cell.value for cell in sheet[3]]!r}")
    expect(all(sheet.cell(row=row, column=1).value is None for row in range(1, 10)), "the opened column is not empty")
    for row, (name, rate, price) in zip((2, 4, 5, 6), FIDELITY_ROWS):
        expect(sheet.cell(row=row, column=2).value == name, f"B{row}: {sheet.cell(row=row, column=2).value!r}")
        expect(sheet.cell(row=row, column=3).value == rate, f"C{row}: {sheet.cell(row=row, column=3).value!r}")
        expect(sheet.cell(row=row, column=4).value == price, f"D{row}: {sheet.cell(row=row, column=4).value!r}")
        expect(sheet.cell(row=row, column=5).value == f"=C{row}*D{row}", f"E{row}: {sheet.cell(row=row, column=5).value!r}")
        expect(sheet.cell(row=row, column=3).number_format == "0.00%", f"C{row}'s format")
    for cell in sheet[1][1:5]:
        expect(cell.font.bold is True, f"{cell.coordinate} lost its bold font")
    expect(sheet.column_dimensions["B"].width == 18, f"column B's width: {sheet.column_dimensions['B'].width!r}")
    expect(sheet.row_dimensions[2].height == 24, f"row 2's height: {sheet.row_dimensions[2].height!r}")
    expect(sheet.row_dimensions[9].height == 30, f"row 9's height: {sheet.row_dimensions[9].height!r}")
    expect(sheet.row_dimensions[9].hidden is True, "row 9 is no longer hidden")
    expect([str(merged) for merged in sheet.merged_cells.ranges] == ["G1:H1"], f"merges: {sheet.merged_cells.ranges!r}")
    expect(sheet["G1"].value == "merged", f"G1: {sheet['G1'].value!r}")
    expect(sheet.freeze_panes == "C2", f"freeze panes: {sheet.freeze_panes!r}")
    rules = {
        str(formatting.sqref): [(rule.type, list(rule.formula or [])) for rule in formatting.rules]
        for formatting in sheet.conditional_formatting
    }
    expect(
        rules == {"C2:C6": [("cellIs", ["1"])], "D2:D6": [("colorScale", [])]},
        f"conditional formats: {rules!r}",
    )
    validations = sheet.data_validations.dataValidation
    expect(len(validations) == 1, f"data validations: {validations!r}")
    validation = validations[0]
    expect(
        (validation.type, str(validation.sqref), validation.formula1, validation.formula2)
        == ("whole", "F2:F6", "0", "100"),
        f"the data validation: {validation!r}",
    )
    expect(sheet["B2"].hyperlink is not None and sheet["B2"].hyperlink.target == "https://example.com/apple", f"B2's link: {sheet['B2'].hyperlink!r}")
    expect(sheet["B4"].hyperlink is not None and sheet["B4"].hyperlink.location == "Report!A1", f"B4's link: {sheet['B4'].hyperlink!r}")
    expect(sheet["A2"].hyperlink is None and sheet["A3"].hyperlink is None, "a link stayed behind")
    expect(sheet.auto_filter.ref == "B1:E6", f"autofilter: {sheet.auto_filter.ref!r}")
    expect(set(sheet.tables) == {"Costs"}, f"tables: {dict(sheet.tables.items())!r}")
    table = sheet.tables["Costs"]
    expect(table.ref == "I1:J4", f"the table: {table.ref!r}")
    expect(table.autoFilter is None or table.autoFilter.ref == "I1:J4", f"the table's filter: {table.autoFilter!r}")
    expect([sheet["I1"].value, sheet["J1"].value, sheet["I2"].value, sheet["I4"].value] == ["Item", "Cost", "Box", "Tape"], "the table's cells")
    expect(isinstance(sheet["B8"].value, CellRichText), f"B8 lost its runs: {sheet['B8'].value!r}")
    report = workbook["Report"]
    expect(report["A1"].value == "=SUM(Data!D2:D6)", f"Report!A1: {report['A1'].value!r}")
    expect(report["A2"].value == "=Rate*Data!D2", f"Report!A2: {report['A2'].value!r}")
    expect("Rate" in workbook.defined_names and workbook.defined_names["Rate"].attr_text == "0.07", "the defined name is gone")


def write_carried_cut_with_openpyxl(path: Path) -> None:
    """Provide three owned worksheet features for the Rust cross-sheet cut."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    source = workbook.active
    source.title = "Data"
    source["B2"], source["B3"] = 1, 2
    source["C2"], source["C3"] = 5, 6
    source["D2"] = "link"
    source["D2"].hyperlink = "https://example.com/carried"
    source.conditional_formatting.add("B2:B3", FormulaRule(formula=["B2>0"]))
    validation = DataValidation(type="whole", operator="greaterThan", formula1="C2")
    validation.add("C2:C3")
    source.add_data_validation(validation)
    target = workbook.create_sheet("Other")
    target["A1"] = "keep"
    workbook.save(path)


def read_carried_cut_with_openpyxl(path: Path, *, moved: bool) -> None:
    """Read the moved or restored owner through the outside implementation."""
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    expect(workbook.sheetnames == ["Data", "Other"], f"carried sheets: {workbook.sheetnames!r}")
    source, target = workbook["Data"], workbook["Other"]
    expect(target["A1"].value == "keep", "the destination sentinel changed")
    owner = target if moved else source
    absent = source if moved else target
    first, second, validated, linked = (("J10", "J11", "K10:K11", "L10") if moved
                                         else ("B2", "B3", "C2:C3", "D2"))
    expect((owner[first].value, owner[second].value) == (1, 2),
           f"carried cells: {owner[first].value!r}, {owner[second].value!r}")
    for original, destination, value in (("B2", "J10", 1), ("B3", "J11", 2),
                                         ("C2", "K10", 5), ("C3", "K11", 6),
                                         ("D2", "L10", "link"), ("D3", "L11", None)):
        current, former = (destination, original) if moved else (original, destination)
        expect(owner[current].value == value,
               f"carried cell {current}: expected {value!r}, got {owner[current].value!r}")
        expect(absent[former].value is None,
               f"cut cell {former} stayed on the former owner: {absent[former].value!r}")
    rules = [(str(item.sqref), rule.type, list(rule.formula or []))
             for item in owner.conditional_formatting for rule in item.rules]
    expect(rules == [(f"{first}:{second}", "expression", [f"{first}>0"])],
           f"carried conditional format: {rules!r}")
    validations = [(item.type, item.operator, str(item.sqref), item.formula1)
                   for item in owner.data_validations.dataValidation]
    expect(validations == [("whole", "greaterThan", validated, "K10" if moved else "C2")],
           f"carried data validation: {validations!r}")
    link = owner[linked].hyperlink
    expect(link is not None and link.target == "https://example.com/carried",
           f"carried external hyperlink: {link!r}")
    expect(len(absent.conditional_formatting) == 0, "a conditional format stayed on the former owner")
    expect(len(absent.data_validations.dataValidation) == 0, "a validation stayed on the former owner")
    expect(absent["D2" if moved else "L10"].hyperlink is None,
           "a hyperlink stayed on the former owner")


def write_named_tables_with_openpyxl(path: Path) -> None:
    """Write two same-sheet tables with independent names and row bands."""
    import os
    from zipfile import ZipFile
    from openpyxl import Workbook, load_workbook
    from openpyxl.worksheet.table import Table, TableColumn

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    for address, value in {
        "A1": "Region\nYear", "B1": "Notes\r\nDetail",
        "A2": 2024, "B2": 3, "A4": 2025, "B4": 4,
        "A5": "Total", "B5": "=SUBTOTAL(109,B2:B4)",
        "D2": 2024, "E2": 12, "D3": 2025, "E3": 19,
    }.items():
        sheet[address] = value
    sheet.add_table(Table(
        displayName="Quantities", ref="A1:B5", totalsRowCount=1,
        tableColumns=[
            TableColumn(id=1, name="Region\nYear", totalsRowLabel="Total"),
            TableColumn(id=2, name="Notes\r\nDetail", totalsRowFunction="sum"),
        ],
    ))
    sheet.add_table(Table(
        displayName="Headless", ref="D2:E3", headerRowCount=0,
        tableColumns=[TableColumn(id=1, name="Year"), TableColumn(id=2, name="Amount")],
    ))
    workbook.save(path)

    # XML normalizes raw CR LF to LF, and CR CR LF to LF LF. The writer's
    # raw spelling can vary by platform; spell the cell's CR explicitly.
    if load_workbook(path).active["B1"].value != "Notes\r\nDetail":
        member = "xl/worksheets/sheet1.xml"
        new = b"Notes&#13;" + bytes((10,)) + b"Detail"
        staged = path.with_name(path.name + ".tmp")
        with ZipFile(path) as source, ZipFile(staged, "w") as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == member:
                    old = b"Notes" + bytes((13, 13, 10)) + b"Detail"
                    if data.count(old) != 1:
                        old = b"Notes" + bytes((13, 10)) + b"Detail"
                    expect(data.count(old) == 1, f"unexpected openpyxl CRLF encoding in {member}")
                    data = data.replace(old, new, 1)
                target.writestr(info, data)
        os.replace(staged, path)
    assert_named_tables_with_openpyxl(path)


def assert_named_tables_with_openpyxl(path: Path, *, edited: bool = False) -> None:
    """Read the outside writer's table parts before and after a Rust save."""
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    expect(workbook.sheetnames == ["Data"], f"named-table sheets: {workbook.sheetnames!r}")
    sheet = workbook["Data"]
    expect(dict(sheet.tables.items()) == {"Quantities": "A1:B5", "Headless": "D2:E3"},
           f"named tables/ranges: {dict(sheet.tables.items())!r}")
    quantities, headless = sheet.tables["Quantities"], sheet.tables["Headless"]
    expect((quantities.headerRowCount, quantities.totalsRowCount) == (1, 1),
           f"Quantities row bands: {quantities.headerRowCount}, {quantities.totalsRowCount}")
    expect([(column.totalsRowLabel, column.totalsRowFunction)
            for column in quantities.tableColumns] == [("Total", None), (None, "sum")],
           "Quantities totals metadata changed")
    expect(headless.headerRowCount == 0 and headless.totalsRowCount in (None, 0),
           f"Headless row bands: {headless.headerRowCount}, {headless.totalsRowCount}")
    expect([column.name for column in quantities.tableColumns] ==
           ["Region\nYear", "Notes\r\nDetail"],
           f"literal tableColumn names: {[column.name for column in quantities.tableColumns]!r}")
    expect([column.name for column in headless.tableColumns] == ["Year", "Amount"],
           f"headless tableColumn names: {[column.name for column in headless.tableColumns]!r}")
    # As with RUST_ESCAPED_SYMBOL, openpyxl leaves ST_Xstring escapes in
    # inline cell text unread (its tableColumn reader does decode them).
    # Pin the exact wire spelling here; Rust also reopens the saved file and
    # asserts the decoded CRLF, so a doubled/protected escape cannot pass.
    note = "Notes_x000D_\nDetail" if edited else "Notes\r\nDetail"
    headers = (sheet["A1"].value, sheet["B1"].value)
    expect(headers == ("Region\nYear", note), f"header cell line breaks: {headers!r}")
    expect([(sheet[f"A{row}"].value, sheet[f"B{row}"].value) for row in (2, 3, 4, 5)] ==
           [(2024, 3), (None, None), (2025, 4), ("Total", "=SUBTOTAL(109,B2:B4)")],
           "Quantities body, blank row or totals changed")
    expect([(sheet[f"D{row}"].value, sheet[f"E{row}"].value) for row in (2, 3)] ==
           [(2024, 12), (2025, 19)], "Headless body changed")
    if edited:
        expect(sheet["J8"].value == "edited by yggdryl", "the Rust edit is missing")


def write_named_resize_with_openpyxl(path: Path) -> None:
    """Author one resizable and one stationary same-sheet table."""
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    for address, value in {
        "A1": "id", "B1": "name", "A2": 1, "B2": "one",
        "A3": 2, "B3": "two", "C4": "neighbor",
        "E1": "kept_id", "F1": "kept_name",
        "E2": 101, "F2": "keep", "E3": 102, "F3": "held",
    }.items():
        sheet[address] = value
    sheet.add_table(Table(displayName="Names", ref="A1:B3"))
    sheet.add_table(Table(displayName="Stationary", ref="E1:F3"))
    workbook.save(path)
    assert_named_resize_with_openpyxl(path, "original")


def assert_named_resize_with_openpyxl(path: Path, mode: str) -> None:
    """Check the outside ListObject reader after Rust grows or shrinks it."""
    from openpyxl import load_workbook

    expected = {
        "original": ("A1:B3", [(1, "one"), (2, "two")]),
        "grow": ("A1:B5", [(9, "nine"), (10, "ten"), (11, "eleven"), (12, "twelve")]),
        "shrink": ("A1:B2", [(9, "nine")]),
    }
    wanted_range, wanted_rows = expected[mode]
    workbook = load_workbook(path)
    expect(workbook.sheetnames == ["Data"], f"resize sheets: {workbook.sheetnames!r}")
    sheet = workbook["Data"]
    ranges = dict(sheet.tables.items())
    expect(ranges == {"Names": wanted_range, "Stationary": "E1:F3"},
           f"{mode} table ranges: {ranges!r}")
    names = sheet.tables["Names"]
    stationary = sheet.tables["Stationary"]
    expect([column.name for column in names.tableColumns] == ["id", "name"],
           f"{mode} selected table columns changed")
    expect([column.name for column in stationary.tableColumns] == ["kept_id", "kept_name"],
           f"{mode} stationary table columns changed")
    if names.autoFilter is not None:
        expect(names.autoFilter.ref == wanted_range,
               f"{mode} selected table filter: {names.autoFilter.ref!r}")
    expect([(sheet[f"A{row}"].value, sheet[f"B{row}"].value)
            for row in range(2, 6)] == wanted_rows + [(None, None)] * (4 - len(wanted_rows)),
           f"{mode} selected body or cleared departing rows changed")
    expect(sheet["A1"].value == "id" and sheet["B1"].value == "name",
           f"{mode} selected header changed")
    expect(sheet["C4"].value == "neighbor", f"{mode} adjacent cell changed")
    expect([(sheet[f"E{row}"].value, sheet[f"F{row}"].value)
            for row in (2, 3)] == [(101, "keep"), (102, "held")],
           f"{mode} stationary table body changed")


def write_named_totals_resize_with_openpyxl(path: Path) -> None:
    """Author a resizable totals table beside one stationary named table."""
    from openpyxl import Workbook
    from openpyxl.worksheet.filters import AutoFilter
    from openpyxl.worksheet.table import Table, TableColumn

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    for address, value in {
        "A1": "id", "B1": "name", "A2": 1, "B2": "one",
        "A3": 2, "B3": "two", "C4": "left sentinel",
        "D1": "year", "E1": "qty", "D2": 2024, "E2": 3,
        "D3": 2025, "E3": 4, "D4": "Total",
        "E4": "=SUBTOTAL(109,[qty])", "G4": "right sentinel",
    }.items():
        sheet[address] = value
    sheet.add_table(Table(displayName="Names", ref="A1:B3"))
    sheet.add_table(Table(
        displayName="Quantities", ref="D1:E4", totalsRowCount=1,
        autoFilter=AutoFilter(ref="D1:E3"),
        tableColumns=[TableColumn(id=1, name="year", totalsRowLabel="Total"),
                      TableColumn(id=2, name="qty", totalsRowFunction="sum")],
    ))
    workbook.save(path)
    assert_named_totals_resize_with_openpyxl(path, "original")


def assert_named_totals_resize_with_openpyxl(path: Path, mode: str) -> None:
    """Prove table/filter bands, moved totals, and unrelated cells after Rust save."""
    from openpyxl import load_workbook

    expected = {
        "original": ("D1:E4", "D1:E3", [(2024, 3), (2025, 4)]),
        "shrink": ("D1:E3", "D1:E2", [(2030, 6)]),
        "equal": ("D1:E4", "D1:E3", [(2030, 6), (2031, 7)]),
        "grow": ("D1:E5", "D1:E4", [(2030, 6), (2031, 7), (2032, 8)]),
    }
    table_range, filter_range, body = expected[mode]
    workbook = load_workbook(path)
    expect(workbook.sheetnames == ["Data"], f"{mode} totals sheets: {workbook.sheetnames!r}")
    sheet = workbook["Data"]
    expect(dict(sheet.tables.items()) == {"Names": "A1:B3", "Quantities": table_range},
           f"{mode} totals table refs: {dict(sheet.tables.items())!r}")
    table = sheet.tables["Quantities"]
    expect((table.headerRowCount, table.totalsRowCount) == (1, 1),
           f"{mode} table row bands: {table.headerRowCount}, {table.totalsRowCount}")
    expect(table.autoFilter is not None and table.autoFilter.ref == filter_range,
           f"{mode} table body filter: {table.autoFilter!r}")
    expect([(column.name, column.totalsRowLabel, column.totalsRowFunction)
            for column in table.tableColumns] ==
           [("year", "Total", None), ("qty", None, "sum")],
           f"{mode} totals column metadata changed")
    totals_row = 2 + len(body)
    for row in range(2, 6):
        actual = sheet[f"D{row}"].value, sheet[f"E{row}"].value
        wanted = body[row - 2] if row - 2 < len(body) else (
            ("Total", "=SUBTOTAL(109,[qty])") if row == totals_row else (None, None))
        expect(actual == wanted, f"{mode} row {row}: {actual!r} != {wanted!r}")
    expect([(sheet[f"A{row}"].value, sheet[f"B{row}"].value)
            for row in (2, 3)] == [(1, "one"), (2, "two")],
           f"{mode} stationary table changed")
    expect((sheet["C4"].value, sheet["G4"].value) ==
           ("left sentinel", "right sentinel"), f"{mode} totals neighbors changed")


def write_rows_with_openpyxl(path: Path) -> None:
    """Write merged Rows(2) headers and a physical empty record row."""
    import os
    from zipfile import ZipFile
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["A1"] = "Sales\nGroup"
    sheet["C1"] = "Person\r\nName"
    sheet["A2"] = "Units"
    sheet["B2"] = "Day"
    sheet.merge_cells("A1:B1")
    sheet.merge_cells("C1:C2")
    sheet.append([2.0, datetime.date(2024, 1, 2), "Ann"])
    sheet.append([None, None, None])
    sheet.append([3.0, datetime.date(2024, 1, 3), "Bob"])
    workbook.save(path)

    # XML 1.0 normalizes a raw CR in element text. Keep the outside
    # producer's CRLF as character-reference CR plus literal LF.
    member = "xl/worksheets/sheet1.xml"
    staged = path.with_name(path.name + ".tmp")
    with ZipFile(path) as source, ZipFile(staged, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == member:
                old = b"Person" + bytes((13, 13, 10)) + b"Name"
                if data.count(old) != 1:
                    old = b"Person" + bytes((13, 10)) + b"Name"
                expect(data.count(old) == 1, "unexpected openpyxl CRLF encoding")
                data = data.replace(old, b"Person&#13;\nName", 1)
            target.writestr(info, data)
    os.replace(staged, path)
    inspect_rows_with_openpyxl(path, from_rust=False)


def inspect_rows_with_openpyxl(path: Path, *, from_rust: bool) -> None:
    """Check physical merged header, literal labels and no-value row."""
    from zipfile import ZipFile
    from xml.etree import ElementTree
    from openpyxl import load_workbook

    sheet = load_workbook(path).active
    expect(set(str(span) for span in sheet.merged_cells.ranges) ==
           {"A1:B1", "C1:C2"}, "Rows(2) merges changed")
    note = "Person_x000D_\nName" if from_rust else "Person\r\nName"
    expect((sheet["A1"].value, sheet["A2"].value, sheet["B2"].value,
            sheet["C1"].value) ==
           ("Sales\nGroup", "Units", "Day", note), "Rows(2) labels changed")
    expect([(sheet[f"A{row}"].value, sheet[f"B{row}"].value,
             sheet[f"C{row}"].value) for row in (3, 4, 5)] ==
           [(2, datetime.datetime(2024, 1, 2), "Ann"),
            (None, None, None),
            (3, datetime.datetime(2024, 1, 3), "Bob")],
           "Rows(2) values or all-null record changed")
    with ZipFile(path) as source:
        root = ElementTree.fromstring(source.read("xl/worksheets/sheet1.xml"))
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    rows = root.findall(f"{ns}sheetData/{ns}row")
    empty = [row for row in rows if row.get("r") == "4"]
    expect(len(empty) == 1 and len(empty[0]) == 0,
           "all-null record must be an explicit empty <row>, no phantom cell")


def write_with_openpyxl(path: Path, path_1904: Path) -> None:
    """Write the same rows from Python natives, and a 1904 workbook."""
    from openpyxl import Workbook
    from openpyxl.utils.datetime import CALENDAR_MAC_1904

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Trades"
    sheet.append(HEADER)
    for row in ROWS:
        sheet.append(list(row))
    for index in range(2, len(ROWS) + 2):
        sheet.cell(row=index, column=5).number_format = "yyyy-mm-dd"
        sheet.cell(row=index, column=6).number_format = "yyyy-mm-dd hh:mm:ss.000"
        sheet.cell(row=index, column=7).number_format = "[h]:mm:ss"
    notes = workbook.create_sheet("Notes")
    notes["A1"] = "written by openpyxl"
    workbook.save(path)

    mac = Workbook()
    mac.epoch = CALENDAR_MAC_1904
    sheet = mac.active
    sheet["A1"] = datetime.date(2024, 2, 29)
    sheet["A1"].number_format = "yyyy-mm-dd"
    sheet["A2"] = datetime.date(1904, 1, 1)
    sheet["A2"].number_format = "yyyy-mm-dd"
    mac.save(path_1904)

    # What was written reads back through openpyxl as itself, so the Rust
    # reading half checks the exchange rather than a broken fixture.
    from openpyxl import load_workbook

    again = load_workbook(path, data_only=True)
    expect(again.sheetnames == ["Trades", "Notes"], "the fixture did not round-trip")
    expect(load_workbook(path_1904).epoch == CALENDAR_MAC_1904, "the 1904 fixture lost its epoch")



def p6_openpyxl_pivot_exchange(path: Path, *, rust_authored: bool) -> None:
    """Check typed pivot metadata and the value format with openpyxl."""
    from openpyxl import load_workbook

    book = load_workbook(path, data_only=True)
    host = "RustPivot" if rust_authored else "CaseOrder"
    name = "P6_interop" if rust_authored else "P6_source_order_grand"
    pivots = [pivot for pivot in book[host]._pivots if pivot.name == name]
    expect(len(pivots) == 1, f"{path}: expected exactly one {name} pivot")
    pivot = pivots[0]
    expect(len(pivot.dataFields) == 1, f"{path}: expected one data field")
    expect(pivot.dataFields[0].subtotal == "sum", f"{path}: aggregate changed")
    expect(pivot.dataFields[0].numFmtId == 2, f"{path}: 0.00 number format lost")
    expect(book[host]["B4"].number_format == "0.00", f"{path}: cell format lost")


def write_openpyxl_pivot_exchange(path: Path) -> None:
    """Modify a committed Excel-authored pivot through openpyxl's pivot model."""
    from openpyxl import load_workbook

    source = REPO / "rust" / "tests" / "excel" / "fixtures" / "pivot_excel.xlsx"
    book = load_workbook(source)
    pivots = [pivot for pivot in book["CaseOrder"]._pivots
              if pivot.name == "P6_source_order_grand"]
    expect(len(pivots) == 1, "Excel-authored fixture is missing its pivot")
    expect(len(pivots[0].dataFields) == 1, "Excel-authored pivot changed shape")
    pivots[0].dataFields[0].numFmtId = 2  # built-in 0.00
    # Missing sortType is manual in OOXML. This outside edit explicitly
    # requests the automatic ordering the public PivotSpec can represent.
    for field in pivots[0].pivotFields:
        if field.axis in ("axisRow", "axisCol"):
            field.sortType = "ascending"
    for row in book["CaseOrder"]["B4:B6"]:
        for cell in row:
            cell.number_format = "0.00"
    book.save(path)


def main() -> int:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        raise SystemExit("openpyxl is not installed: pip install openpyxl") from None

    EXCHANGE.mkdir(parents=True, exist_ok=True)
    for name in (
        "from-rust.xlsx",
        "from-rust-styled.xlsx",
        "from-openpyxl.xlsx",
        "from-openpyxl-1904.xlsx",
        "from-openpyxl-styled.xlsx",
        "from-rust-edited.xlsx",
        "from-openpyxl-fidelity.xlsx",
        "from-openpyxl-anchored-fidelity.xlsx",
        "from-rust-fidelity.xlsx",
        "from-rust-shifted.xlsx",
        "from-openpyxl-carried.xlsx",
        "from-rust-carried.xlsx",
        "from-rust-carried-undone.xlsx",
        "from-openpyxl-regions.xlsx",
        "from-rust-regions-edited.xlsx",
        "from-openpyxl-named-resize.xlsx",
        "from-rust-named-grow.xlsx",
        "from-rust-named-shrink.xlsx",
        "from-openpyxl-named-totals.xlsx",
        "from-rust-named-totals-grow.xlsx",
        "from-rust-named-totals-equal.xlsx",
        "from-rust-named-totals-shrink.xlsx",
        "from-rust-rows.xlsx",
        "from-rust-pivot.xlsx",
        "from-openpyxl-pivot.xlsx",
        "from-rust-pivot-edited.xlsx",
        "from-openpyxl-rows.xlsx",
    ):
        (EXCHANGE / name).unlink(missing_ok=True)

    first = run_cargo(allow_skip=True)
    expect("excel-interop: wrote" in first, f"the Rust writing half did not report:\n{first}")
    expect("excel-interop: wrote pivot" in first, "Rust pivot writer did not report")
    p6_openpyxl_pivot_exchange(EXCHANGE / "from-rust-pivot.xlsx", rust_authored=True)
    read_with_openpyxl(EXCHANGE / "from-rust.xlsx")
    print("openpyxl read the workbook this crate wrote")
    expect("excel-interop: wrote styled" in first, f"the Rust styling half did not report:\n{first}")
    read_styled_from_rust(EXCHANGE / "from-rust-styled.xlsx")
    print("openpyxl read the bold, filled, bordered and formatted cells this crate styled")
    expect("excel-interop: wrote rows" in first, "Rust nested-row writer did not report")
    inspect_rows_with_openpyxl(EXCHANGE / "from-rust-rows.xlsx", from_rust=True)
    print("openpyxl read merged multiline headers and physical all-null records this crate wrote")

    write_openpyxl_pivot_exchange(EXCHANGE / "from-openpyxl-pivot.xlsx")
    write_with_openpyxl(EXCHANGE / "from-openpyxl.xlsx", EXCHANGE / "from-openpyxl-1904.xlsx")
    write_styled_with_openpyxl(EXCHANGE / "from-openpyxl-styled.xlsx")
    write_fidelity_with_openpyxl(EXCHANGE / "from-openpyxl-fidelity.xlsx")
    write_anchored_fidelity(EXCHANGE / "from-openpyxl-fidelity.xlsx",
                            EXCHANGE / "from-openpyxl-anchored-fidelity.xlsx")
    write_carried_cut_with_openpyxl(EXCHANGE / "from-openpyxl-carried.xlsx")
    write_named_tables_with_openpyxl(EXCHANGE / "from-openpyxl-regions.xlsx")
    write_named_resize_with_openpyxl(EXCHANGE / "from-openpyxl-named-resize.xlsx")
    write_named_totals_resize_with_openpyxl(EXCHANGE / "from-openpyxl-named-totals.xlsx")
    write_rows_with_openpyxl(EXCHANGE / "from-openpyxl-rows.xlsx")
    second = run_cargo(allow_skip=False)
    expect("excel-interop: read" in second, f"the Rust reading half did not report:\n{second}")
    expect("excel-interop: read 1904" in second, f"the 1904 half did not report:\n{second}")
    expect("excel-interop: edited" in second, f"the Rust editing half did not report:\n{second}")
    expect("excel-interop: fidelity" in second, f"the Rust fidelity half did not report:\n{second}")
    expect("excel-interop: shifted" in second, f"the Rust shifting half did not report:\n{second}")
    expect("excel-interop: missing-anchor refusal" in second,
           f"the CSS-only atomic refusal half did not report:\n{second}")
    expect("excel-interop: carried cut" in second, f"the Rust carried-cut half did not report:\n{second}")
    expect("excel-interop: named tables" in second, f"the Rust named-table half did not report:\n{second}")
    expect("excel-interop: resized named table" in second,
           f"the Rust named-table resize half did not report:\n{second}")
    expect("excel-interop: resized named totals" in second,
           f"the Rust totals-table resize half did not report:\n{second}")
    expect("excel-interop: read pivot" in second, "Rust pivot reader did not report")
    expect("excel-interop: read rows" in second, "Rust nested-row reader did not report")
    print("this crate read the workbooks openpyxl wrote, including merged multiline headers and all-null records")

    p6_openpyxl_pivot_exchange(EXCHANGE / "from-rust-pivot-edited.xlsx", rust_authored=False)
    read_edited_with_openpyxl(EXCHANGE / "from-rust-edited.xlsx")
    print("openpyxl read the styled workbook this crate edited, its styles kept")

    read_fidelity_with_openpyxl(EXCHANGE / "from-rust-fidelity.xlsx")
    print("openpyxl read the workbook this crate edited, every feature it wrote kept")

    read_shifted_with_openpyxl(EXCHANGE / "from-rust-shifted.xlsx")
    print("openpyxl read every reference this crate moved with a row and a column")
    read_carried_cut_with_openpyxl(EXCHANGE / "from-rust-carried.xlsx", moved=True)
    read_carried_cut_with_openpyxl(EXCHANGE / "from-rust-carried-undone.xlsx", moved=False)
    print("openpyxl read cross-sheet CF, validation and hyperlink ownership and saved undo")
    assert_named_tables_with_openpyxl(EXCHANGE / "from-rust-regions-edited.xlsx", edited=True)
    print("openpyxl read both named tables after this crate saved an unrelated cell edit")
    for mode in ("grow", "shrink"):
        assert_named_resize_with_openpyxl(EXCHANGE / f"from-rust-named-{mode}.xlsx", mode)
    print("openpyxl read Rust-grown and Rust-shrunk named tables beside a stationary table")
    for mode in ("shrink", "equal", "grow"):
        assert_named_totals_resize_with_openpyxl(
            EXCHANGE / f"from-rust-named-totals-{mode}.xlsx", mode)
    print("openpyxl read Rust-resized totals rows, table filters and stationary neighbors")
    return 0


if __name__ == "__main__":
    sys.exit(main())
