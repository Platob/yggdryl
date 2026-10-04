#!/usr/bin/env python3
"""Local-only Microsoft Excel oracle; never runs in CI or starts cargo.

Windows setup: python -m pip install pywin32 openpyxl

    python scripts/check_excel_desktop.py probe
    python scripts/check_excel_desktop.py inspect --workbook PATH [--workbook PATH]
    python scripts/check_excel_desktop.py fidelity --rich PATH
    python scripts/check_excel_desktop.py shift --shift-original PATH --shifted PATH
    python scripts/check_excel_desktop.py formats --cases PATH
    python scripts/check_excel_desktop.py styles [--cases PATH]
    python scripts/check_excel_desktop.py functions --cases PATH
    python scripts/check_excel_desktop.py pivots --cases PATH --active

All modes write versioned results.json under --output-dir. Fidelity consumes
every half of scripts/check_excel_interop.py, plus --rich; it compares Excel's
own row-3/column-A insert with from-rust-shifted.xlsx. On success only, it saves
rich_excel.xlsx and its version note for review before committing the fixture.
Shift runs only the structural comparison, including the same feature-coverage
check. It compares Excel's edited reference with the opened Rust output, then
saves and reopens both through Excel for a second exact comparison. Both save
cycles are equal: Excel itself can quantize note offsets during serialization.

Format cases are exported by Rust, never extracted from Rust source:
{"schema_version":1,"cases":[{"id":"round","code":"0.00","value":1.005,
"expected":"1.01","date_system":"1900"}]}. Values may be numbers, strings,
booleans or null. A value_kind of nonfinite carries NaN/Infinity/-Infinity as
JSON strings; they cross COM as numbers, never as text. Every case receives an
Excel answer or a named failure, including an unrepresentable numeric input.
Style cases use the same envelope: {"id":"tint","file":"book.xlsx",
"sheet":"Sheet1","cell":"A1","expected":{"font":"#123456",
"fill":"#FFFFFF","bold":true,"borders":{"left":"#000000"}}}.
Relative style file paths resolve beside the case file; expected colors must
come from StyleSheet::resolve, not from this oracle's own reading.
Function cases name two authored workbooks in a schema-versioned manifest. The
oracle calculates only worksheets belonging to each newly opened input, records
Excel formula/value observations and calibration outcomes, then stages cached
SaveAs copies. The supervisor publishes copies only after worker cleanup.

By default a child process owns a DispatchEx instance. Explicit --active
attaches to an existing Excel application without starting another instance.
Attached mode preserves its visibility and user workbooks; it restores the
preferences it changes and the original active window. Probe changes neither.
Already-open input workbooks are refused so a user's workbook is never closed.
The parent enforces --timeout and detects visible modal dialogs. It may stop
only an isolated Excel instance, after checking PID, creation time and image.
Attached Excel is never quit or terminated. A killed worker cannot run its
finally blocks; the report explicitly names possible unrestored preferences,
focus or still-open oracle workbooks. Unknown process identities are never guessed.

COM contracts: https://learn.microsoft.com/en-us/office/vba/api/excel.workbooks.open
https://learn.microsoft.com/en-us/office/vba/api/excel.application.automationsecurity
https://learn.microsoft.com/en-us/office/vba/api/excel.worksheetfunction.text
"""

from __future__ import annotations

import argparse
import contextlib
import datetime as dt
import gc
import hashlib
import json
import math
from pathlib import Path
import re
import posixpath
import struct
import xml.etree.ElementTree as ET
import zipfile
import subprocess
import sys
import tempfile
import time
import traceback
from typing import Any, Callable

REPO = Path(__file__).resolve().parents[1]
EXCHANGE = REPO / "rust" / "target" / "excel-interop"
OUTPUT = REPO / "rust" / "target" / "excel-desktop"
EXCHANGE_NAMES = (
    "from-rust.xlsx", "from-rust-styled.xlsx", "from-openpyxl.xlsx",
    "from-openpyxl-1904.xlsx", "from-openpyxl-styled.xlsx", "from-rust-edited.xlsx",
    "from-openpyxl-fidelity.xlsx", "from-openpyxl-anchored-fidelity.xlsx",
    "from-rust-fidelity.xlsx", "from-rust-shifted.xlsx",
)
EDGES = {"left": 7, "top": 8, "bottom": 9, "right": 10}
FONT_PROPERTIES = ("Name", "Size", "Bold", "Italic", "Underline", "Strikethrough", "Color",
                   "OutlineFont", "Shadow", "Subscript", "Superscript")
# Keep the character-by-character COM scan bounded independently of UsedRange.
MAX_CELL_CHARACTERS = 32_767
MAX_SNAPSHOT_CHARACTERS = 100_000


def write_json(path: Path, value: Any) -> None:
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, indent=2, ensure_ascii=True, allow_nan=False) + "\n", encoding="utf-8")
    for attempt in range(20):
        try:
            temporary.replace(path)
            return
        except PermissionError as error:
            # A Windows reader can briefly deny replacement while the
            # supervisor reads this checkpoint. Keep the previous JSON whole.
            if getattr(error, "winerror", None) not in (5, 32) or attempt == 19:
                raise
            time.sleep(0.01)


def read_cases(path: Path, mode: str) -> list[dict[str, Any]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict) or payload.get("schema_version") != 1:
        raise ValueError("cases: expected schema_version 1")
    cases = payload.get("cases")
    if not isinstance(cases, list) or not cases:
        raise ValueError("cases: expected a nonempty cases array")
    ids: set[str] = set()
    for case in cases:
        if not isinstance(case, dict) or not isinstance(case.get("id"), str) or not case["id"]:
            raise ValueError("cases: each case needs a nonempty string id")
        if case["id"] in ids:
            raise ValueError(f"cases: duplicate id {case['id']!r}")
        ids.add(case["id"])
        if mode == "formats":
            if not isinstance(case.get("code"), str) or not isinstance(case.get("expected"), str):
                raise ValueError(f"{case['id']}: code and expected must be strings")
            if "value" not in case or not isinstance(case["value"], (str, int, float, bool, type(None))):
                raise ValueError(f"{case['id']}: expected a JSON scalar value")
            if isinstance(case["value"], (int, float)) and not math.isfinite(case["value"]):
                raise ValueError(f"{case['id']}: nonfinite value")
            format_value(case)
            if case.get("date_system", "1900") not in ("1900", "1904"):
                raise ValueError(f"{case['id']}: date_system must be 1900 or 1904")
        else:
            for key in ("file", "sheet", "cell"):
                if not isinstance(case.get(key), str) or not case[key]:
                    raise ValueError(f"{case['id']}: missing {key}")
            expected = case.get("expected")
            if not isinstance(expected, dict) or not expected:
                raise ValueError(f"{case['id']}: expected a nonempty style expectation")
            if set(expected) - {"font", "fill", "bold", "italic", "number_format", "borders"}:
                raise ValueError(f"{case['id']}: unknown style expectation")
            for key in ("font", "fill"):
                if key in expected:
                    rgb_to_excel(expected[key])
            for edge, color in expected.get("borders", {}).items():
                if edge not in EDGES:
                    raise ValueError(f"{case['id']}: unknown border {edge}")
                rgb_to_excel(color)
    return cases


def read_function_manifest(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if (not isinstance(payload, dict) or payload.get("schema_version") != 1
            or payload.get("kind") != "p5_function_oracle_inputs" or payload.get("native_answers") is not False):
        raise ValueError("functions: expected schema-versioned inputs without native answers")
    workbooks = payload.get("workbooks")
    cases = payload.get("cases")
    if not isinstance(workbooks, list) or not workbooks or not isinstance(cases, list) or not cases:
        raise ValueError("functions: expected nonempty workbooks and cases")
    by_file = {}
    for workbook in workbooks:
        if not isinstance(workbook, dict):
            raise ValueError("functions: expected workbook metadata")
        name, year, digest = (workbook.get(key) for key in ("file", "date_system", "sha256"))
        if (not isinstance(name, str) or not re.fullmatch(r"functions-input-(1900|1904)\.xlsx", name)
                or name in by_file or year not in ("1900", "1904")
                or name != f"functions-input-{year}.xlsx"):
            raise ValueError("functions: expected distinct 1900/1904 input workbook names")
        file = path.parent / name
        if not file.is_file() or not isinstance(digest, str) or hashlib.sha256(file.read_bytes()).hexdigest() != digest:
            raise ValueError(f"functions/{name}: missing file or sha256 mismatch")
        by_file[name] = workbook
    if set(by_file) != {"functions-input-1900.xlsx", "functions-input-1904.xlsx"}:
        raise ValueError("functions: expected both authored 1900 and 1904 inputs")
    ids, coordinates = set(), set()
    counts = {"cached_equal": 0, "behavior_only": 0, "calibration": 0}
    calibration_gates = set()
    for case in cases:
        if not isinstance(case, dict):
            raise ValueError("functions: expected case objects")
        identifier, name, sheet, cell = (case.get(key) for key in ("id", "file", "sheet", "cell"))
        if (not isinstance(identifier, str) or not identifier or identifier in ids
                or name not in by_file or not isinstance(sheet, str) or not sheet
                or not isinstance(cell, str) or not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", cell)):
            raise ValueError(f"functions: invalid or duplicate case {identifier!r}")
        coordinate = (name, sheet, cell)
        if coordinate in coordinates:
            raise ValueError(f"functions/{identifier}: duplicate formula cell {coordinate!r}")
        ids.add(identifier)
        coordinates.add(coordinate)
        if (case.get("date_system") != by_file[name]["date_system"]
                or not isinstance(case.get("wire_formula"), str) or not case["wire_formula"]
                or case.get("mode") not in counts
                or not isinstance(case.get("gates"), list)
                or any(not isinstance(gate, str) or not gate for gate in case["gates"])):
            raise ValueError(f"functions/{identifier}: invalid formula, mode, date system or gates")
        if "expected" in case or "actual" in case:
            raise ValueError(f"functions/{identifier}: native answer must be observed, not supplied")
        counts[case["mode"]] += 1
        if case["mode"] == "calibration":
            control = case.get("control")
            if (not isinstance(control, dict) or control.get("kind") != "equals"
                    or not isinstance(control.get("gate"), str) or "value" not in control
                    or control["gate"] in calibration_gates):
                raise ValueError(f"functions/{identifier}: invalid calibration control")
            calibration_gates.add(control["gate"])
        elif case["mode"] == "behavior_only":
            control = case.get("control")
            if not isinstance(control, dict) or control.get("kind") not in (
                    "number_bounds", "integer_bounds", "clock_observation"):
                raise ValueError(f"functions/{identifier}: invalid behavior control")
    if payload.get("mode_counts") != counts:
        raise ValueError("functions: mode counts disagree with cases")
    if any(gate not in calibration_gates for case in cases for gate in case["gates"]):
        raise ValueError("functions: case names an absent calibration gate")
    for name, workbook in by_file.items():
        if workbook.get("cases") != sum(case["file"] == name for case in cases):
            raise ValueError(f"functions/{name}: case count disagrees with manifest")
    companions = set()
    for case in cases:
        plan = case.get("text_format_oracle")
        if plan is None:
            continue
        if (not isinstance(plan, dict) or set(plan) not in ({"value_expression", "code"}, {"value_expression", "code", "cell", "format_cell"})
                or any(not isinstance(value, str) or not value for value in plan.values())
                or case["mode"] == "behavior_only"
                or not re.fullmatch(r"-?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[Ee][+-]?[0-9]+)?|DATE\([+-]?[0-9]+,[+-]?[0-9]+,[+-]?[0-9]+\)", plan["value_expression"])
                or function_text_formula(plan["value_expression"], plan["code"]) != case["wire_formula"]):
            raise ValueError(f"functions/{case['id']}: TEXT oracle must describe the exact authored formula")
        for key in (() if "cell" not in plan else ("cell", "format_cell")):
            address = plan[key]
            coordinate = (case["file"], case["sheet"], address)
            if (not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", address)
                    or coordinate in coordinates or coordinate in companions):
                raise ValueError(f"functions/{case['id']}: TEXT companion conflicts with a selected cell")
            companions.add(coordinate)
    return payload


# https://learn.microsoft.com/en-us/office/vba/api/excel.xlconsolidationfunction
PIVOT_AGGREGATES = {
    "sum": -4157, "count": -4112, "average": -4106,
    "max": -4136, "min": -4139, "product": -4149,
    "countNumbers": -4113, "stdDev": -4155, "stdDevP": -4156,
    "var": -4164, "varP": -4165,
}

def read_pivot_manifest(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if (not isinstance(payload, dict) or payload.get("schema_version") != 1
            or payload.get("kind") != "p6_pivot_oracle_inputs"
            or payload.get("native_answers") is not False):
        raise ValueError("pivots: expected schema-versioned inputs without native answers")
    if "verify_existing" in payload:
        if set(payload) != {"schema_version", "kind", "native_answers", "verify_existing"}:
            raise ValueError("pivots: verify_existing has no other input mode or native answers")
        authored = payload["verify_existing"]
        if (not isinstance(authored, dict) or set(authored) != {"file", "sha256", "pivots"}
                or not isinstance(authored["file"], str)
                or Path(authored["file"]).name != authored["file"]
                or not authored["file"].endswith(".xlsx")):
            raise ValueError("pivots: invalid authored workbook identity")
        source = path.parent / authored["file"]
        if (not source.is_file() or not isinstance(authored["sha256"], str)
                or hashlib.sha256(source.read_bytes()).hexdigest() != authored["sha256"]):
            raise ValueError("pivots: authored workbook sha256 mismatch")
        cases = authored["pivots"]
        if (not isinstance(cases, list) or not cases
                or any(not isinstance(case, dict) or set(case) != {"id", "sheet", "name"}
                       or any(not isinstance(value, str) or not value for value in case.values())
                       for case in cases)
                or any(len({case[key] for case in cases}) != len(cases) for key in ("id", "sheet", "name"))):
            raise ValueError("pivots: expected unique authored pivot identities and sheets")
        return payload
    if "compare_existing" in payload:
        if any(key in payload for key in ("check_existing", "workbook", "source_tables", "cases", "modified_members", "observe_labels")):
            raise ValueError("pivots: compare_existing has no other pivot input mode")
        pair = payload["compare_existing"]
        if (not isinstance(pair, list) or len(pair) != 2
                or [item.get("id") if isinstance(item, dict) else None for item in pair]
                   != ["control", "candidate"]):
            raise ValueError("pivots: compare_existing requires control/candidate roles")
        chosen = None
        for item in pair:
            if (set(item) != {"id", "file", "sha256", "sheet", "name"}
                    or not all(isinstance(item[key], str) and item[key] for key in item)
                    or Path(item["file"]).name != item["file"]
                    or not item["file"].endswith(".xlsx")):
                raise ValueError("pivots: compare_existing identity invalid")
            identity = item["sheet"], item["name"]
            if chosen is not None and identity != chosen:
                raise ValueError("pivots: compare_existing must select the same pivot")
            chosen = identity
            source = path.parent / item["file"]
            if (not source.is_file() or not re.fullmatch(r"[0-9a-f]{64}", item["sha256"])
                    or hashlib.sha256(source.read_bytes()).hexdigest() != item["sha256"]):
                raise ValueError(f"pivots/{item['id']}: missing workbook or sha256 mismatch")
        return payload
    if "check_existing" in payload:
        if any(key in payload for key in ("workbook", "source_tables", "cases")):
            raise ValueError("pivots: check_existing cannot also author new pivots")
        if "observe_labels" in payload and payload["observe_labels"] is not True:
            raise ValueError("pivots: observe_labels requires true")
        pair = payload["check_existing"]
        if (not isinstance(pair, list) or len(pair) != 2
                or [item.get("id") if isinstance(item, dict) else None for item in pair]
                   != ["control", "candidate"]):
            raise ValueError("pivots: check_existing requires control/candidate roles")
        chosen = None
        files = []
        for item in pair:
            if (set(item) != {"id", "file", "sha256", "sheet", "name"}
                    or not all(isinstance(item[key], str) and item[key] for key in item)
                    or Path(item["file"]).name != item["file"]
                    or not item["file"].endswith(".xlsx")):
                raise ValueError("pivots: check_existing identity or native answers invalid")
            identity = (item["sheet"], item["name"])
            if chosen is not None and identity != chosen:
                raise ValueError("pivots: check_existing must select the same pivot")
            chosen = identity
            source = path.parent / item["file"]
            if (not source.is_file() or not re.fullmatch(r"[0-9a-f]{64}", item["sha256"])
                    or hashlib.sha256(source.read_bytes()).hexdigest() != item["sha256"]):
                raise ValueError(f"pivots/{item['id']}: missing workbook or sha256 mismatch")
            files.append(source)
        members = payload.get("modified_members")
        if (not isinstance(members, list) or not members
                or any(not isinstance(name, str) or not name.startswith("xl/") for name in members)
                or len(set(members)) != len(members)):
            raise ValueError("pivots: invalid modified members")
        with zipfile.ZipFile(files[0]) as control, zipfile.ZipFile(files[1]) as candidate:
            if control.namelist() != candidate.namelist():
                raise ValueError("pivots: modified members changed package member inventory")
            changed = {name for name in control.namelist()
                       if control.read(name) != candidate.read(name)}
        if changed != set(members):
            raise ValueError("pivots: modified members disagree with package bytes")
        return payload
    workbook = payload.get("workbook")
    if not isinstance(workbook, dict) or workbook.get("file") != "pivot-input.xlsx":
        raise ValueError("pivots: expected pivot-input.xlsx")
    source = path.parent / workbook["file"]
    digest = workbook.get("sha256")
    if (not source.is_file() or not isinstance(digest, str)
            or not re.fullmatch(r"[0-9a-f]{64}", digest)
            or hashlib.sha256(source.read_bytes()).hexdigest() != digest):
        raise ValueError("pivots: missing source workbook or sha256 mismatch")
    tables, cases = payload.get("source_tables"), payload.get("cases")
    if not isinstance(tables, list) or not tables or not isinstance(cases, list) or not cases:
        raise ValueError("pivots: expected source tables and cases")
    fields = {}
    for table in tables:
        if (not isinstance(table, dict) or not isinstance(table.get("sheet"), str)
                or not isinstance(table.get("range"), str) or not isinstance(table.get("rows"), list)
                or not table["rows"] or not isinstance(table["rows"][0], list)):
            raise ValueError("pivots: invalid source table")
        header = table["rows"][0]
        if (table["sheet"] in fields or not header
                or any(not isinstance(field, str) or not field for field in header)
                or len(set(header)) != len(header)
                or any(not isinstance(row, list) or len(row) != len(header) for row in table["rows"])):
            raise ValueError("pivots: invalid or duplicate source table fields")
        fields[table["sheet"]] = set(header)
    ids, destinations = set(), set()
    for case in cases:
        if not isinstance(case, dict):
            raise ValueError("pivots: expected case objects")
        identifier, sheet, anchor = (case.get(key) for key in ("id", "sheet", "anchor"))
        source_spec = case.get("source")
        if not isinstance(identifier, str) or not re.fullmatch(r"[a-z][a-z0-9_]+", identifier):
            raise ValueError(f"pivots: invalid case id {identifier!r}")
        if identifier in ids:
            raise ValueError(f"pivots: duplicate id {identifier!r}")
        if (not isinstance(sheet, str) or not sheet
                or not isinstance(anchor, str) or not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", anchor)
                or not isinstance(source_spec, dict)
                or not isinstance(source_spec.get("sheet"), str)
                or source_spec.get("sheet") not in fields
                or not isinstance(source_spec.get("range"), str)):
            raise ValueError(f"pivots: invalid case identity/source {identifier!r}")
        if sheet in destinations:
            raise ValueError(f"pivots: duplicate destination {sheet!r}")
        if source_spec["range"] != next(item["range"] for item in tables if item["sheet"] == source_spec["sheet"]):
            raise ValueError(f"pivots/{identifier}: source range disagrees with authored table")
        rows, columns, values = (case.get(key) for key in ("rows", "columns", "values"))
        if (not isinstance(rows, list) or not isinstance(columns, list) or not isinstance(values, list)
                or not rows or not values or any(not isinstance(axis, str) for axis in rows + columns)
                or len(set(rows + columns)) != len(rows + columns)
                or any(axis not in fields[source_spec["sheet"]] for axis in rows + columns)
                or any(not isinstance(item, dict) or not isinstance(item.get("field"), str)
                       or item.get("field") not in fields[source_spec["sheet"]]
                       or item.get("aggregate") not in PIVOT_AGGREGATES
                       or not isinstance(item.get("caption"), str) or not item["caption"] for item in values)
                or any(type(case.get(key)) is not bool for key in
                       ("subtotals", "row_grand_totals", "column_grand_totals"))):
            raise ValueError(f"pivots/{identifier}: invalid axes, values or totals")
        for caption in ("data_caption", "grand_total_caption"):
            if caption in case and (not isinstance(case[caption], str)
                                    or not 1 <= len(case[caption]) <= 255):
                raise ValueError(f"pivots/{identifier}: invalid {caption} caption")
        if "item_order" in case and case["item_order"] not in ("ascending", "descending"):
            raise ValueError(f"pivots/{identifier}: invalid item_order")
        if any(key in case for key in ("actual", "expected", "native_answer")):
            raise ValueError(f"pivots/{identifier}: native answers must be observed")
        ids.add(identifier)
        destinations.add(sheet)
    return payload


def function_value(value: Any) -> dict[str, Any]:
    """Preserve COM transport; saved XML and native ERROR.TYPE establish meaning."""
    if type(value) is float:
        finite = math.isfinite(value)
        classification = "finite" if finite else "nan" if math.isnan(value) else "positive-infinity" if value > 0 else "negative-infinity"
        return {"variant": "float", "value": value if finite else None, "hex": value.hex(),
                "ieee754_hex": struct.pack(">d", value).hex(), "classification": classification}
    if type(value) in (bool, int, str):
        return {"variant": type(value).__name__, "value": value}
    if value is None:
        return {"variant": "empty", "value": None}
    raise TypeError(f"unsupported Excel Value2 variant {type(value).__name__}")


FUNCTION_ERRORS = ("#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#GETTING_DATA")


def function_cell(app: Any, cell: Any, address: str, evaluate: Callable[[str], Any]) -> dict[str, Any]:
    if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", address):
        raise ValueError(f"invalid owned formula cell address {address!r}")
    value = function_value(cell.Value2)
    result = {"formula": str(cell.Formula), "formula2": str(cell.Formula2),
              "value2": value, "display": str(cell.Text)}
    # Pass the owned Range, not the marshalled NaN/HRESULT: Excel16 can expose
    # a #NUM! cell as either transport. No error identity is guessed from bits.
    error = app.WorksheetFunction.IsError(cell)
    # Excel16 returns a one-element SAFEARRAY for the owned Range containing
    # a 32767-character result. Retain that observed transport and accept only
    # its single typed Boolean; no truthiness or multi-cell reduction is valid.
    if type(error) is tuple and len(error) == 1 and type(error[0]) is bool:
        result["is_error_transport"] = {"variant": "tuple", "value": list(error)}
        error = error[0]
    if type(error) is not bool:
        raise TypeError("ISERROR must return a Boolean")
    result["is_error"] = error
    if error:
        try:
            try:
                classify = app.WorksheetFunction.Error_Type
            except AttributeError:
                # Some Excel COM type libraries do not expose ERROR.TYPE.
                # Evaluate only the validated address on its owned worksheet.
                number = evaluate(f"ERROR.TYPE({address})")
                origin = "worksheet_evaluate"
            else:
                number = classify(cell)
                origin = "worksheet_function"
            if type(number) not in (int, float) or not math.isfinite(number) or number != int(number) or not 1 <= number <= len(FUNCTION_ERRORS):
                raise ValueError(f"unrepresented ERROR.TYPE result {number!r}")
            result["error_type"] = {"number": int(number), "literal": FUNCTION_ERRORS[int(number) - 1], "source": origin}
            if origin == "worksheet_evaluate":
                result["error_type"]["lcid"] = 1033
        except Exception as exception:
            result["error_type_refusal"] = {"type": type(exception).__name__, "message": str(exception)}
    return result



def decode_ooxml_cached_text(text: str) -> str:
    """Read OOXML _xHHHH_ UTF-16 units, including lone surrogate values.

    One pass preserves a literal `_xHHHH_` whose underscore was `_x005F_`.
    Paired escapes become the same Unicode scalar COM reports; lone halves
    remain Python surrogate code points for exact native comparison.
    """
    pieces = []
    end = 0
    for escape in re.finditer(r"_x([0-9A-Fa-f]{4})_", text):
        pieces.append(text[end:escape.start()])
        pieces.append(chr(int(escape.group(1), 16)))
        end = escape.end()
    pieces.append(text[end:])
    return "".join(pieces).encode("utf-16-le", errors="surrogatepass").decode(
        "utf-16-le", errors="surrogatepass")


def compare_function_cache(actual: dict[str, Any], cached: dict[str, Any]) -> dict[str, str]:
    """Prove native transport vs saved type; retain an explicit unknown-error limit."""
    kind, text = cached["type"], cached["value_text"]
    value = actual["value2"]
    if kind == "e":
        if not actual["is_error"]:
            return {"status": "different", "reason": "native non-error became a saved error"}
        if "error_type" not in actual:
            return {"status": "cache-authoritative", "reason": "native ISERROR confirmed an error; ERROR.TYPE could not prove its literal"}
        equal = actual["error_type"]["literal"] == text
    elif actual["is_error"]:
        return {"status": "different", "reason": "native error became a saved non-error"}
    elif kind == "n":
        try:
            numeric = float(text)
        except (TypeError, ValueError):
            return {"status": "different", "reason": "saved numeric spelling cannot be interpreted"}
        if value["variant"] == "float":
            equal = struct.pack(">d", numeric).hex() == value["ieee754_hex"]
        else:
            equal = value["variant"] == "int" and math.isfinite(numeric) and value["value"] == numeric
    elif kind == "b":
        equal = value["variant"] == "bool" and text in ("0", "1") and value["value"] == (text == "1")
    elif kind == "str":
        equal = value["variant"] == "str" and value["value"] == decode_ooxml_cached_text(text or "")
    else:
        return {"status": "different", "reason": f"unrepresented saved formula cache type {kind!r}"}
    return {"status": "equal" if equal else "different", "reason": "native value and saved cache type/value " + ("agree" if equal else "differ")}


def function_control(control: dict[str, Any], actual: Any) -> bool:
    kind = control["kind"]
    if kind == "equals":
        return first_difference(control["value"], actual) is None
    if isinstance(actual, bool) or not isinstance(actual, (int, float)) or not math.isfinite(actual):
        return False
    if kind == "clock_observation":
        return actual >= 0 and (not control["integral"] or actual == int(actual))
    if kind == "integer_bounds":
        return actual == int(actual) and control["lower"] <= actual <= control["upper"]
    return control["lower"] <= actual and (actual <= control["upper"] if control["upper_inclusive"] else actual < control["upper"])


def saved_function_cache(path: Path, cases: list[dict[str, Any]], *,
                         formulas_only: bool = True,
                         allow_absent: bool = False) -> dict[tuple[str, str], dict[str, Any]]:
    # One OOXML worksheet/cache reader serves both formula and pivot observations.
    # No style/date conversion: serial 60 and submillisecond numeric bits stay numeric.
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    document = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package = "http://schemas.openxmlformats.org/package/2006/relationships"
    requested: dict[str, set[str]] = {}
    for case in cases:
        targets = requested.setdefault(case["sheet"], set())
        if case["cell"] in targets:
            raise ValueError(f"duplicate native formula request {case['sheet']}!{case['cell']}")
        targets.add(case["cell"])
    result = {}
    with zipfile.ZipFile(path) as archive:
        strings = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            strings = ["".join(node.text or "" for node in
                       [*item.findall(f"{{{main}}}t"),
                        *(run.find(f"{{{main}}}t") for run in item.findall(f"{{{main}}}r"))]
                       ) for item in shared.findall(f"{{{main}}}si")]
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relations = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        for name, cells in requested.items():
            sheets = [sheet for sheet in workbook.findall(f"{{{main}}}sheets/{{{main}}}sheet") if sheet.attrib.get("name") == name]
            if len(sheets) != 1:
                raise ValueError(f"native workbook lacks one worksheet {name!r}")
            relation = sheets[0].attrib[f"{{{document}}}id"]
            links = [link for link in relations.findall(f"{{{package}}}Relationship") if link.attrib.get("Id") == relation]
            if len(links) != 1 or links[0].attrib.get("TargetMode") == "External" or links[0].attrib.get("Type") != document + "/worksheet":
                raise ValueError(f"native {name!r} relationship is not one internal worksheet")
            target = links[0].attrib["Target"]
            member = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
            if member.startswith("../"):
                raise ValueError(f"native {name!r} target escapes the package")
            sheet = ET.fromstring(archive.read(member))
            for cell in sheet.findall(f"{{{main}}}sheetData/{{{main}}}row/{{{main}}}c"):
                address = cell.attrib.get("r")
                if address not in cells:
                    continue
                key = name, address
                if key in result:
                    raise ValueError(f"native duplicate cell {name}!{address}")
                formula, cached = cell.find(f"{{{main}}}f"), cell.find(f"{{{main}}}v")
                if formulas_only and (formula is None or cached is None):
                    raise ValueError(f"native {name}!{address}: formula/cache missing")
                kind = cell.attrib.get("t", "n")
                text = None if cached is None else cached.text
                if not formulas_only:
                    if kind == "s":
                        if text is None or not text.isdecimal() or int(text) >= len(strings):
                            raise ValueError(f"native {name}!{address}: shared string index invalid")
                        kind, text = "str", strings[int(text)]
                    elif kind == "inlineStr":
                        inline = cell.find(f"{{{main}}}is")
                        if inline is None:
                            raise ValueError(f"native {name}!{address}: inline string missing")
                        kind, text = "str", "".join(node.text or "" for node in
                            [*inline.findall(f"{{{main}}}t"),
                             *(run.find(f"{{{main}}}t") for run in inline.findall(f"{{{main}}}r"))]
                            if node is not None)
                result[key] = {"type": kind, "value_text": text,
                               "formula_text": None if formula is None else formula.text,
                               "formula_attributes": {} if formula is None else dict(formula.attrib),
                               "attributes": dict(cell.attrib)}
            missing = cells - {address for sheet_name, address in result if sheet_name == name}
            if missing and not allow_absent:
                raise ValueError(f"native {name}: missing formula cache cells {sorted(missing)}")
    return result


def format_value(case: dict) -> Any:
    """Resolve an exported scalar kind without conflating numbers and text."""
    value, kind = case["value"], case.get("value_kind")
    if kind is None:
        return value
    valid = {"number": type(value) in (int, float), "text": isinstance(value, str),
             "boolean": isinstance(value, bool), "blank": value is None,
             "nonfinite": isinstance(value, str) and value in ("NaN", "Infinity", "-Infinity")}
    if not isinstance(kind, str) or not valid.get(kind, False):
        raise ValueError(f"{case['id']}: value_kind {kind!r} disagrees with value {value!r}")
    return float(value) if kind == "nonfinite" else value


def repair_logs(directory: Path | None = None) -> dict[Path, tuple[int, int, str]]:
    root = directory or Path(tempfile.gettempdir())
    return {path: (path.stat().st_mtime_ns, path.stat().st_size, hashlib.sha256(path.read_bytes()).hexdigest())
            for path in root.glob("error*.xml") if path.is_file()}


def changed_repair_logs(before: dict, directory: Path | None = None) -> list[Path]:
    return sorted(path for path, signature in repair_logs(directory).items() if before.get(path) != signature)


def first_difference(expected: Any, actual: Any, path: str = "$") -> str | None:
    if type(expected) is not type(actual):
        # Excel returns some exact integers as doubles; their value is the contract.
        if not isinstance(expected, bool) and not isinstance(actual, bool) and isinstance(expected, (int, float)) and isinstance(actual, (int, float)) and expected == actual:
            return None
        return f"{path}: expected {expected!r}, actual {actual!r}"
    if isinstance(expected, dict):
        if expected.keys() != actual.keys():
            return f"{path}: expected keys {sorted(expected)}, actual {sorted(actual)}"
        for key in expected:
            if found := first_difference(expected[key], actual[key], f"{path}.{key}"):
                return found
    elif isinstance(expected, (list, tuple)):
        if len(expected) != len(actual):
            return f"{path}: expected {len(expected)} entries, actual {len(actual)}"
        for index, (left, right) in enumerate(zip(expected, actual)):
            if found := first_difference(left, right, f"{path}[{index}]"):
                return found
    elif expected != actual:
        return f"{path}: expected {expected!r}, actual {actual!r}"
    return None


def rgb_to_excel(value: str) -> int:
    if not isinstance(value, str):
        raise ValueError(f"expected six RGB hex digits, got {value!r}")
    code = value.removeprefix("#")
    if len(code) != 6 or any(char not in "0123456789abcdefABCDEF" for char in code):
        raise ValueError(f"expected six RGB hex digits, got {value!r}")
    return int(code[0:2], 16) | int(code[2:4], 16) << 8 | int(code[4:6], 16) << 16


def fidelity_inputs(exchange: Path, rich: Path | None) -> list[Path]:
    if rich is None or not rich.is_file():
        raise ValueError("fidelity: --rich must name the exported rich_package() workbook")
    paths = [exchange / name for name in EXCHANGE_NAMES]
    for path in paths:
        if not path.is_file():
            raise ValueError(f"fidelity: missing exchange half {path}")
    return paths + [rich]


def collection(value: Any) -> list[Any]:
    return [value.Item(index) for index in range(1, int(value.Count) + 1)]


def indexed_property(value: Any, name: str, *indices: int) -> Any:
    import pythoncom
    return value._oleobj_.InvokeTypes(value._oleobj_.GetIDsOfNames(name), 0, pythoncom.DISPATCH_PROPERTYGET,
                                     (pythoncom.VT_VARIANT, 0), ((pythoncom.VT_VARIANT, 1),) * len(indices), *indices)


def characters(cell: Any, start: int, length: int | None = None) -> Any:
    import win32com.client
    indices = (start,) if length is None else (start, length)
    return win32com.client.Dispatch(indexed_property(cell, "Characters", *indices))


def text_runs(cell: Any, text: str, path: str, remaining: int) -> tuple[list[dict], int]:
    """Read each native Characters position once; retain coalesced UTF-16 offsets.

    Aggregate Font properties cannot identify runs, even when none is null.
    Excel's indexing is observed through Count and validated through each Text;
    unsupported indexing fails instead of silently losing a surrogate or run.
    https://learn.microsoft.com/en-us/office/vba/api/excel.range.characters
    """
    if len(text) > MAX_CELL_CHARACTERS:
        raise ValueError(f"{path}: rich text exceeds oracle's {MAX_CELL_CHARACTERS}-unit cell bound")
    try:
        encoded = text.encode("utf-16-le")
    except UnicodeEncodeError as error:
        raise ValueError(f"{path}: unsupported rich text: unpaired surrogate in cell value") from error
    units = len(encoded) // 2
    if units > MAX_CELL_CHARACTERS:
        raise ValueError(f"{path}: rich text exceeds oracle's {MAX_CELL_CHARACTERS}-unit cell bound")
    if units > remaining:
        raise ValueError(f"{path}: rich text exceeds oracle's {MAX_SNAPSHOT_CHARACTERS}-unit workbook bound")
    # Excel16 build20430 collapses CRLF to one LF only in Characters. Keep
    # source UTF-16 offsets separately so the literal Value2 and run coverage
    # remain exact even when the native display has fewer positions.
    display = text.replace("\r\n", "\n")
    display_units = units if display is text else len(display.encode("utf-16-le")) // 2
    try:
        whole = characters(cell, 1)
        count = whole.Count
        if type(count) is not int or count not in (len(display), display_units):
            raise ValueError("Characters.Count does not describe codepoint or UTF-16 indexing")
        # Whole Characters.Text can reject a long cell independently of the
        # requested positions. Each position below proves the exact Value2
        # bytes and its font; aggregate text supplies no additional evidence.
    except Exception as error:
        raise ValueError(f"{path}: unsupported rich text: {error}") from error
    runs: list[dict] = []
    offset = index = 1
    source_at = 0
    while source_at < len(text):
        character = text[source_at]
        source_length = 2 if character == "\r" and text[source_at:source_at + 2] == "\r\n" else 1
        piece = ("\n" if source_length == 2 else character).encode("utf-16-le")
        width = 2 if source_length == 2 else len(piece) // 2
        native_width = len(piece) // 2 if count == display_units else 1
        try:
            # Request a complete surrogate pair. Single-unit Text can expose
            # an empty high half and the whole scalar at the low-half position.
            part = characters(cell, index, native_width)
            displayed = part.Text
            observed_text = displayed.encode("utf-16-le", errors="surrogatepass") if isinstance(displayed, str) else None
            # Excel16 build20430 exposes Value2 CR as LF in Characters.Text
            # at the same position (Sheet1!B4 in the recorded desktop probe).
            # This validates only the display position: cells retain Value2
            # exactly, and neither dropped positions nor other changes pass.
            if observed_text != piece and not (piece == b"\r\x00" and observed_text == b"\n\x00"):
                raise ValueError("Characters.Text has unsupported indexing")
            font = part.Font
            observed = {name: getattr(font, name) for name in FONT_PROPERTIES}
            for name, value in observed.items():
                if value is None or not isinstance(value, (str, bool, int, float)) or (
                        isinstance(value, float) and not math.isfinite(value)):
                    raise ValueError(f"Font.{name} is not a single-character scalar")
        except Exception as error:
            raise ValueError(f"{path}: unsupported rich text at Characters({index}, {native_width}): {error}") from error
        if runs and runs[-1]["font"] == observed:
            runs[-1]["length_utf16"] += width
        else:
            runs.append({"start_utf16": offset, "length_utf16": width, "font": observed})
        offset += width
        index += native_width
        source_at += source_length
    return runs, units


def wait_ready(app: Any, *, active: bool = False) -> None:
    import pythoncom
    import pywintypes
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        try:
            if app.Ready and (active or app.Workbooks.Count == 0):
                return
        except pywintypes.com_error as error:
            if error.hresult not in (-2147418111, -2147417846):
                raise
        pythoncom.PumpWaitingMessages()
        time.sleep(0.1)
    if active:
        raise RuntimeError("active Excel did not become ready within 10s")
    raise RuntimeError("isolated Excel did not become ready with zero workbooks within 10s")


def canonical(items: list[Any]) -> list[Any]:
    return sorted(items, key=lambda value: json.dumps(value, sort_keys=True, ensure_ascii=False))


def special_cells(sheet: Any, kind: int) -> list[Any]:
    import pywintypes
    try:
        area = sheet.Cells.SpecialCells(kind)
    except pywintypes.com_error as error:
        if error.excepinfo and error.excepinfo[5] == -2146827284:  # Excel: no cells found.
            return []
        raise
    if int(area.CountLarge) > 20_000:
        raise ValueError(f"{sheet.Name}: more than 20000 special cells; fixture exceeds oracle bound")
    return list(area.Cells)


def snapshot(book: Any) -> dict[str, Any]:
    """Capture observable feature references, not Excel's incidental ZIP spelling."""
    def formula(value: Any) -> Any:
        if not isinstance(value, str):
            return value
        # Text literals are data even inside a formula; only qualify references.
        pattern = r'"(?:[^"]|"")*"|\[' + re.escape(book.Name) + r'\](?=[^!]*!)'
        return re.sub(pattern, lambda match: match[0] if match[0].startswith('"') else "[workbook]", value)

    def paint(value: Any) -> dict:
        font, fill = value.Font, value.Interior
        borders = {}
        for edge, index in {**EDGES, "diagonal_down": 5, "diagonal_up": 6}.items():
            border = value.Borders.Item(index)
            borders[edge] = None if border.LineStyle == -4142 else {
                "style": border.LineStyle, "weight": border.Weight, "color": border.Color}
        return {"font": {name: getattr(font, name) for name in FONT_PROPERTIES},
                "fill": {"pattern": fill.Pattern, "color": fill.Color, "pattern_color": fill.PatternColor},
                "borders": borders}

    def frozen(sheet: Any) -> list:
        active, visible = book.ActiveSheet, sheet.Visible
        try:
            if visible != -1:
                sheet.Visible = -1
            sheet.Activate()
            window = book.Windows.Item(1)
            return [bool(window.FreezePanes), int(window.SplitRow), int(window.SplitColumn)]
        finally:
            active.Activate()
            if visible != -1:
                sheet.Visible = visible

    def anchor(shape: Any) -> dict:
        return {"from": shape.TopLeftCell.Address, "to": shape.BottomRightCell.Address,
                "geometry": {name: getattr(shape, name) for name in ("Left", "Top", "Width", "Height", "Placement")}}

    sheets = []
    remaining_characters = MAX_SNAPSHOT_CHARACTERS
    for sheet in collection(book.Worksheets):
        used = sheet.UsedRange
        if int(used.CountLarge) > 20_000:
            raise ValueError(f"{sheet.Name}: used range exceeds oracle's 20000-cell bound")
        cells, styles, formulas, rich_text, merges = [], [], [], [], set()
        for cell in used.Cells:
            value = cell.Value2
            if value is not None or cell.HasFormula:
                cells.append([cell.Address, formula(cell.Formula) if cell.HasFormula else value])
            if cell.HasFormula:
                formulas.append(cell.Address)
            elif isinstance(value, str) and value:
                runs, inspected = text_runs(cell, value, f"{sheet.Name}!{cell.Address}", remaining_characters)
                remaining_characters -= inspected
                rich_text.append([cell.Address, runs])
            styles.append([cell.Address, {**paint(cell), "number_format": cell.NumberFormat,
                           "alignment": {name: getattr(cell, name) for name in (
                               "HorizontalAlignment", "VerticalAlignment", "WrapText", "ShrinkToFit", "Orientation", "IndentLevel")}}])
            if cell.MergeCells:
                merges.add(cell.MergeArea.Address)
        rules = []
        for rule in collection(sheet.Cells.FormatConditions):
            record = {"range": rule.AppliesTo.Address, "type": int(rule.Type)}
            if rule.Type in (1, 2):
                record["formula1"] = formula(rule.Formula1)
                if rule.Type == 1:
                    record["operator"] = int(rule.Operator)
                    if rule.Operator in (1, 2):
                        record["formula2"] = formula(rule.Formula2)
            rules.append(record)
        validations = []
        for cell in special_cells(sheet, -4174):
            validation = cell.Validation
            record = {"cell": cell.Address, "type": int(validation.Type), "ignore_blank": bool(validation.IgnoreBlank)}
            if validation.Type != 0:
                record["formula1"] = formula(validation.Formula1)
                if validation.Type not in (3, 7):
                    record["operator"] = int(validation.Operator)
                    if validation.Operator in (1, 2):
                        record["formula2"] = formula(validation.Formula2)
            validations.append(record)
        charts = []
        for chart in collection(sheet.ChartObjects()):
            charts.append({"name": chart.Name, **anchor(chart),
                           "series": [formula(series.Formula) for series in collection(chart.Chart.SeriesCollection())]})
        sheets.append({
            "name": sheet.Name, "cells": canonical(cells), "merges": sorted(merges),
            "formulas": sorted(formulas),
            "visible": int(sheet.Visible), "rich_text": canonical(rich_text), "styles": canonical(styles),
            "rows": [[int(row.Row), row.RowHeight, bool(row.Hidden), int(row.OutlineLevel)] for row in used.Rows],
            "columns": [[int(column.Column), column.ColumnWidth, bool(column.Hidden), int(column.OutlineLevel)] for column in used.Columns],
            "conditional_formats": canonical(rules), "validations": canonical(validations),
            "hyperlinks": canonical([{"cell": link.Range.Address, "address": link.Address, "location": link.SubAddress}
                                     for link in collection(sheet.Hyperlinks)]),
            "autofilter": sheet.AutoFilter.Range.Address if sheet.AutoFilterMode else None,
            "tables": canonical([{"name": table.Name, "range": table.Range.Address} for table in collection(sheet.ListObjects)]),
            "charts": canonical(charts),
            "comments": canonical([{"cell": comment.Parent.Address, "author": comment.Author, "text": comment.Text(),
                                    **anchor(comment.Shape)}
                                   for comment in collection(sheet.Comments)]),
            "pivots": canonical([{"name": pivot.Name, "range": pivot.TableRange2.Address}
                                 for pivot in collection(sheet.PivotTables())]),
            "frozen": frozen(sheet),
        })
    return {"date1904": bool(book.Date1904), "sheets": sheets,
            "sheet_order": [sheet.Name for sheet in collection(book.Sheets)],
            "active_sheet": book.ActiveSheet.Name,
            "chartsheets": [{"name": chart.Name, "visible": int(chart.Visible), "type": int(chart.ChartType),
                             "series": [formula(series.Formula) for series in collection(chart.SeriesCollection())]}
                            for chart in collection(book.Charts)],
            "names": canonical([{"name": name.Name, "formula": formula(name.RefersTo)} for name in collection(book.Names)])}


@contextlib.contextmanager
def open_checked(app: Any, path: Path):
    # Workbooks.Open can return an existing workbook. Never acquire cleanup
    # ownership of a document that the user already has open.
    resolved = path.resolve()
    for existing in collection(app.Workbooks):
        if Path(existing.FullName).resolve() == resolved:
            raise ValueError(f"{path.name}: input workbook is already open in Excel")
    before = repair_logs()
    book = None
    try:
        # xlNormalLoad refuses explicit repair; nonempty passwords suppress prompts.
        book = app.Workbooks.Open(str(path.resolve()), UpdateLinks=0, ReadOnly=True,
                                  Password="yggdryl-oracle-no-password", WriteResPassword="yggdryl-oracle-no-password",
                                  IgnoreReadOnlyRecommended=True, Notify=False, AddToMru=False, Local=False, CorruptLoad=0)
        # RepairMode is exposed by some Excel versions, absent from others.
        try:
            repaired = book.RepairMode
        except AttributeError:
            repaired = None
        if repaired:
            raise AssertionError(f"{path.name}: Excel RepairMode is true")
        logs = changed_repair_logs(before)
        if logs:
            # TEMP is shared with user Excel processes. Activity is not proof
            # this workbook was repaired, and another workbook's log is private.
            raise RuntimeError(f"{path.name}: verification indeterminate: {len(logs)} unattributed repair log(s) "
                               "changed in shared TEMP during open")
        yield book
    finally:
        if book is not None:
            book.Close(SaveChanges=False)


def evaluate_function(sheet: Any) -> Callable[[str], Any]:
    """Resolve one owned worksheet's English Evaluate dispatcher once."""
    import pythoncom
    dispatch = sheet._oleobj_
    identifier = dispatch.GetIDsOfNames("Evaluate")
    result_type = (pythoncom.VT_VARIANT, 0)
    argument_types = ((pythoncom.VT_VARIANT, 1),)
    method = pythoncom.DISPATCH_METHOD

    def evaluate(expression: str) -> Any:
        # Native Excel16/French UI returns #NAME? for ERROR.TYPE at LCID0,
        # but all eight classic error numbers at1033. No workbook text changes.
        return dispatch.InvokeTypes(identifier, 1033, method, result_type, argument_types, expression)

    return evaluate


def text_function(app: Any) -> Callable[[Any, str], str]:
    """Resolve Excel TEXT once, passing original format codes under en-US."""
    import pythoncom
    function = app.WorksheetFunction._oleobj_
    identifier = function.GetIDsOfNames("Text")
    result_type = (pythoncom.VT_VARIANT, 0)
    argument_types = ((pythoncom.VT_VARIANT, 1), (pythoncom.VT_VARIANT, 1))
    method = pythoncom.DISPATCH_METHOD

    def render(value: Any, code: str) -> str:
        # Explicit LCID1033 is observed on French Excel16 build20430; LCID0
        # interprets the same date code as local literals. No code is rewritten.
        return function.InvokeTypes(identifier, 1033, method, result_type, argument_types, value, code)

    return render


def run_formats(app: Any, args: argparse.Namespace, result: dict) -> None:
    if args.cases is None:
        raise ValueError("formats: --cases must name the Rust-exported cases JSON")
    cases = read_cases(args.cases, "formats")
    text = text_function(app)
    result["oracle_method"] = "Application.WorksheetFunction.Text(value, code); LCID1033; original code; en-US calibration; decimal='.'; thousands=','"
    result["comparison_contract"] = "TEXT results are observations; bool/null coercion and General cell display may differ from FormatCode::render"
    separators = (app.UseSystemSeparators, app.DecimalSeparator, app.ThousandsSeparator)
    result["application_separators_before"] = dict(zip(
        ("UseSystemSeparators", "DecimalSeparator", "ThousandsSeparator"), separators))
    # Persist recovery values before changing application-wide preferences.
    write_json(args.output_dir / "results.json", result)
    book = None
    try:
        app.UseSystemSeparators = False
        app.DecimalSeparator = "."
        app.ThousandsSeparator = ","
        book = app.Workbooks.Add()
        # Locale calibration prevents French OS settings becoming format defects.
        calibration = {"decimal": text(1234.5, "#,##0.00"),
                       "date": text(45292, "dddd, mmmm d, yyyy")}
        result["locale_calibration"] = calibration
        if calibration != {"decimal": "1,234.50", "date": "Monday, January 1, 2024"}:
            raise ValueError(f"formats: Excel is not rendering en-US: {calibration!r}")
        result["cases"] = []
        for case in cases:
            answer = dict(case)
            try:
                book.Date1904 = case.get("date_system", "1900") == "1904"
                actual = text(format_value(case), case["code"])
                answer["actual"] = actual
                answer["passed"] = actual == case["expected"]
                if not answer["passed"]:
                    result["failures"].append(f"formats/{case['id']}: expected {case['expected']!r}, Excel {actual!r}")
            except Exception as error:
                answer.update(passed=False, error=str(error))
                result["failures"].append(f"formats/{case['id']}: {error}")
            result["cases"].append(answer)
    finally:
        try:
            if book is not None:
                book.Close(SaveChanges=False)
        finally:
            # These are application preferences, even in an isolated COM
            # instance. Restore each one before Quit can persist changes.
            try:
                app.DecimalSeparator = separators[1]
            finally:
                try:
                    app.ThousandsSeparator = separators[2]
                finally:
                    app.UseSystemSeparators = separators[0]


def inspect_workbooks(app: Any, paths: list[Path], result: dict, *, checkpoint: Path | None = None) -> None:
    result["workbooks"] = {}
    for path in paths:
        print(f"workbook/start {path.name[:160]}", flush=True)
        try:
            with open_checked(app, path) as book:
                result["workbooks"][str(path.resolve())] = snapshot(book)
            print(f"workbook/ok {path.name[:160]}", flush=True)
        except Exception as error:
            result["failures"].append(f"workbook/{path.name}: {error}")
            print(f"workbook/failed {path.name[:160]}", flush=True)
        if checkpoint is not None:
            # Diagnostic observations remain unverified until worker cleanup.
            # Only the supervisor can publish the staged Excel fixture.
            write_json(checkpoint, result)


def require_shift_coverage(observed: dict, sheet_name: str) -> None:
    sheet = next((sheet for sheet in observed["sheets"] if sheet["name"] == sheet_name), None)
    if sheet is None:
        raise ValueError(f"fidelity/shift coverage: missing sheet {sheet_name!r}")
    required = {"formulas": "cell formulas", "conditional_formats": "conditional formats",
                "validations": "data validation", "hyperlinks": "hyperlinks", "autofilter": "autofilter",
                "tables": "tables", "comments": "comments/VML anchors", "merges": "merges"}
    missing = [label for key, label in required.items() if not sheet.get(key)]
    if not any(chart.get("series") for chart in sheet.get("charts", [])):
        missing.append("chart formulas/drawing anchors")
    if not observed.get("names"):
        missing.append("defined names")
    if missing:
        raise ValueError(f"fidelity/shift coverage: {sheet_name!r} lacks {', '.join(missing)}; full structural validation refused")


def run_shift(app: Any, args: argparse.Namespace, result: dict) -> None:
    """Compare native and Rust edits before and after equal Excel save cycles."""
    checkpoint = args.output_dir / "results.json"
    original = args.shift_original or args.exchange / "from-openpyxl-anchored-fidelity.xlsx"
    shifted = args.shifted or args.exchange / "from-rust-shifted.xlsx"
    print(f"fidelity/shift/start {original.name[:160]} -> {shifted.name[:160]}", flush=True)
    shift_failures = len(result["failures"])
    try:
        native = args.output_dir / "shift_excel.xlsx"
        rust_native = args.output_dir / "shift_rust_excel.xlsx"
        for target in (native, rust_native):
            if target.exists():
                raise ValueError(f"shift fixture already exists; select a fresh --output-dir: {target}")
        with open_checked(app, original) as book:
            require_shift_coverage(snapshot(book), args.sheet)
            sheet = book.Worksheets.Item(args.sheet)
            sheet.Rows(args.insert_row).Insert()
            sheet.Columns(args.insert_column).Insert()
            before_save = snapshot(book)
            book.SaveAs(str(native.resolve()), FileFormat=51, AddToMru=False, Local=False)
        # Excel may quantize all four note offsets again at each save. Keep
        # the first-open comparison and give both sides the same save cycle;
        # a tolerance here would hide a real shift or serialization defect.
        with open_checked(app, native) as book:
            expected = snapshot(book)
        with open_checked(app, shifted) as book:
            actual_before_save = snapshot(book)
            book.SaveAs(str(rust_native.resolve()), FileFormat=51, AddToMru=False, Local=False)
        with open_checked(app, rust_native) as book:
            actual = snapshot(book)
        result["shift"] = {"original": str(original), "shifted": str(shifted), "native": str(native),
                           "rust_native": str(rust_native), "expected_before_save": before_save,
                           "actual_before_save": actual_before_save, "expected": expected, "actual": actual}
        if difference := first_difference(before_save, actual_before_save):
            result["failures"].append(f"fidelity/shift/before-save: {difference}")
        if difference := first_difference(expected, actual):
            result["failures"].append(f"fidelity/shift/after-save: {difference}")
    except Exception as error:
        result["failures"].append(f"fidelity/shift: {error}")
    print(f"fidelity/shift/{'failed' if len(result['failures']) != shift_failures else 'ok'}", flush=True)
    write_json(checkpoint, result)


def native_save_normalization(before: dict, after: dict) -> list[dict]:
    """Record native VML coordinate quantization; every other fact stays exact."""
    coordinates = ("Left", "Top", "Width", "Height")

    def without_comment_coordinates(value: dict) -> dict:
        # Borrow the large cell/style payloads. Native SaveAs may quantize only
        # these coordinates: Placement and any future geometry facts stay exact.
        result = dict(value)
        if "sheets" in value:
            result["sheets"] = []
            for sheet in value["sheets"]:
                held = dict(sheet)
                if "comments" in sheet:
                    held["comments"] = []
                    for comment in sheet["comments"]:
                        comparison = dict(comment)
                        geometry = comment.get("geometry")
                        if isinstance(geometry, dict):
                            comparison["geometry"] = {key: item for key, item in geometry.items()
                                                      if key not in coordinates}
                        held["comments"].append(comparison)
                result["sheets"].append(held)
        return result

    if difference := first_difference(without_comment_coordinates(before), without_comment_coordinates(after)):
        raise AssertionError(f"fidelity/Excel-save/source-retention: {difference}")
    changes = []
    for sheet_index, (left, right) in enumerate(zip(before.get("sheets", []), after.get("sheets", []))):
        for comment_index, (source, saved) in enumerate(zip(left.get("comments", []), right.get("comments", []))):
            path = f"$.sheets[{sheet_index}].comments[{comment_index}].geometry"
            source, saved = source.get("geometry"), saved.get("geometry")
            if (not isinstance(source, dict) or not isinstance(saved, dict)
                    or any(name not in source or name not in saved for name in coordinates)):
                raise AssertionError(f"fidelity/Excel-save/source-retention: {path}: expected four native geometry coordinates")
            for name in coordinates:
                first, last = source[name], saved[name]
                if any(isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value)
                       for value in (first, last)):
                    raise AssertionError(f"fidelity/Excel-save/source-retention: {path}.{name}: expected finite native geometry")
                if first_difference(first, last):
                    changes.append({"path": f"{path}.{name}", "before": first, "after": last})
    return changes


def run_fidelity(app: Any, args: argparse.Namespace, result: dict) -> None:
    paths = fidelity_inputs(args.exchange, args.rich)
    checkpoint = args.output_dir / "results.json"
    inspect_workbooks(app, paths + args.workbook, result, checkpoint=checkpoint)
    original_key = str((args.exchange / "from-openpyxl-fidelity.xlsx").resolve())
    edited_key = str((args.exchange / "from-rust-fidelity.xlsx").resolve())
    if original_key in result["workbooks"] and edited_key in result["workbooks"]:
        print(f"fidelity/edited/start {Path(original_key).name[:160]}", flush=True)
        # Let Excel extend styles and geometry when an edit expands UsedRange.
        with open_checked(app, Path(original_key)) as book:
            for name, edits in {"Data": {"A5": "Kiwi, edited", "E2": 50}, "Report": {"A3": "written again"}}.items():
                for address, value in edits.items():
                    book.Worksheets.Item(name).Range(address).Value2 = value
            expected = snapshot(book)
        if difference := first_difference(expected, result["workbooks"][edited_key]):
            result["failures"].append(f"fidelity/edited-feature-retention: {difference}")
        print(f"fidelity/edited/{'failed' if difference else 'ok'}", flush=True)
        write_json(checkpoint, result)
    run_shift(app, args, result)
    if result["failures"]:
        return
    target = args.output_dir / "rich_excel.xlsx"
    if target.exists() or target.with_suffix(".json").exists():
        raise ValueError(f"fidelity: fixture already exists; select a fresh --output-dir: {target}")
    staging = target.with_name("rich_excel.pending.xlsx")
    note = staging.with_suffix(".json")
    if staging.exists() or note.exists():
        raise ValueError(f"fidelity: unverified fixture exists; select a fresh --output-dir: {staging}")
    control = target.with_name("rich_excel.control.xlsx")
    if control.exists():
        raise ValueError(f"fidelity: native control already exists; select a fresh --output-dir: {control}")
    observed = {"source": str(args.rich.resolve()), "control_file": str(control),
                "candidate_file": str(staging), "native_save_cycles": 1}
    result["rich_save"] = observed
    print(f"fidelity/Excel-save/start {args.rich.name[:160]}", flush=True)
    try:
        # Two independent opens and one native SaveAs per leg. Excel itself
        # quantizes VML offsets on each save; a one-sided extra save would test
        # native idempotence rather than the fixture's retained information.
        print("fidelity/Excel-save/control/start", flush=True)
        with open_checked(app, args.rich) as book:
            observed["control_before_save"] = snapshot(book)
            write_json(checkpoint, result)
            source = result["workbooks"].get(observed["source"])
            if source is not None and (difference := first_difference(source, observed["control_before_save"])):
                raise AssertionError(f"fidelity/Excel-save/source-open: {difference}")
            book.SaveAs(str(control.resolve()), FileFormat=51, AddToMru=False, Local=False)
        with open_checked(app, control) as book:
            observed["control"] = snapshot(book)
        write_json(checkpoint, result)
        print("fidelity/Excel-save/control/ok", flush=True)
        print("fidelity/Excel-save/candidate/start", flush=True)
        with open_checked(app, args.rich) as book:
            observed["candidate_before_save"] = snapshot(book)
            write_json(checkpoint, result)
            if difference := first_difference(observed["control_before_save"], observed["candidate_before_save"]):
                raise AssertionError(f"fidelity/Excel-save/before-save: {difference}")
            book.SaveAs(str(staging.resolve()), FileFormat=51, AddToMru=False, Local=False)
        with open_checked(app, staging) as book:
            observed["candidate"] = snapshot(book)
        write_json(checkpoint, result)
        if difference := first_difference(observed["control"], observed["candidate"]):
            raise AssertionError(f"fidelity/Excel-save/after-save: {difference}")
        observed["normalization"] = native_save_normalization(observed["control_before_save"], observed["control"])
        write_json(note, {"schema_version": 1, "excel": result["excel"],
                          "source": str(args.rich), "verified": True, "snapshot": observed["candidate"],
                          "native_save": {"cycles_per_leg": 1, "normalization": observed["normalization"],
                                          "observations": "results.json#/rich_save"}})
        # The supervisor publishes only after the worker and its cleanup succeed.
        result["fixture_candidate"] = str(staging)
        write_json(checkpoint, result)
    except Exception:
        print("fidelity/Excel-save/failed", flush=True)
        result.pop("fixture_candidate", None)
        staging.unlink(missing_ok=True)
        note.unlink(missing_ok=True)
        raise
    print("fidelity/Excel-save/ok", flush=True)


def finish_fixture(result: dict, output: Path, passed: bool) -> None:
    staging = output / "rich_excel.pending.xlsx"
    note = staging.with_suffix(".json")
    if not passed:
        staging.unlink(missing_ok=True)
        note.unlink(missing_ok=True)
        result.pop("fixture_candidate", None)
        return
    if "fixture_candidate" not in result:
        return
    if Path(result["fixture_candidate"]).resolve() != staging.resolve():
        raise ValueError("fidelity: fixture candidate is outside the output directory")
    target = output / "rich_excel.xlsx"
    target_note = target.with_suffix(".json")
    if target.exists() or target_note.exists():
        raise ValueError("fidelity: refusing to replace an existing fixture or version note")
    staging.rename(target)
    try:
        note.rename(target_note)
    except Exception:
        target.unlink()
        raise
    result.pop("fixture_candidate")
    result["fixture"] = str(target)


def run_styles(app: Any, args: argparse.Namespace, result: dict) -> None:
    if args.cases:
        cases = read_cases(args.cases, "styles")
        base = args.cases.parent
    else:
        cases = [{"id": "header", "file": "from-rust-styled.xlsx", "sheet": "Styled", "cell": "A1",
                  "expected": {"bold": True, "fill": "#FFFF00", "borders": dict.fromkeys(EDGES, "#000000")}},
                 {"id": "custom", "file": "from-rust-styled.xlsx", "sheet": "Styled", "cell": "B1",
                  "expected": {"number_format": '#,##0.000 "kg"'}},
                 {"id": "percent", "file": "from-rust-styled.xlsx", "sheet": "Styled", "cell": "C1",
                  "expected": {"number_format": "0.00%"}}]
        base = args.exchange
    # Intake is complete before any workbook opens. A file is owned once, and
    # every observed case is checkpointed before moving to the next one.
    groups: dict[Path, list[tuple[int, dict[str, Any]]]] = {}
    for position, case in enumerate(cases):
        groups.setdefault((base / case["file"]).resolve(), []).append((position, case))
    answers: list[dict[str, Any] | None] = [None] * len(cases)
    result["cases"] = []
    completed = 0
    checkpoint = args.output_dir / "results.json"

    def record(position: int, answer: dict[str, Any]) -> None:
        nonlocal completed
        answers[position] = answer
        completed += 1
        result["cases"] = [item for item in answers if item is not None]
        result["styles_progress"]["completed"] = completed
        write_json(checkpoint, result)

    for path, group in groups.items():
        result["styles_progress"] = {"file": str(path), "completed": completed, "total": len(cases)}
        write_json(checkpoint, result)
        processed = 0
        try:
            with open_checked(app, path) as book:
                for position, case in group:
                    answer = dict(case)
                    try:
                        cell = book.Worksheets.Item(case["sheet"]).Range(case["cell"])
                        actual, expected = {}, {}
                        for key, value in case["expected"].items():
                            if key in ("fill", "font"):
                                actual[key] = int((cell.Interior if key == "fill" else cell.Font).Color)
                                expected[key] = rgb_to_excel(value)
                            elif key == "borders":
                                for edge in value:
                                    if cell.Borders.Item(EDGES[edge]).LineStyle == -4142:
                                        raise AssertionError(f"{case['cell']}: {edge} border is absent")
                                actual[key] = {edge: int(cell.Borders.Item(EDGES[edge]).Color) for edge in value}
                                expected[key] = {edge: rgb_to_excel(color) for edge, color in value.items()}
                            else:
                                actual[key] = {"bold": lambda: bool(cell.Font.Bold), "italic": lambda: bool(cell.Font.Italic),
                                               "number_format": lambda: cell.NumberFormat}[key]()
                                expected[key] = value
                        difference = first_difference(expected, actual)
                        answer.update(actual=actual, passed=difference is None)
                        if difference:
                            result["failures"].append(f"styles/{case['id']}: {difference}")
                    except Exception as error:
                        answer.update(passed=False, error=str(error))
                        result["failures"].append(f"styles/{case['id']}: {error}")
                    processed += 1
                    record(position, answer)
        except Exception as error:
            # Opening and close verification belong to the file. If opening
            # failed, keep every case accounted for without opening it again.
            if processed == len(group):
                result["failures"].append(f"styles/{path.name}: {error}")
                write_json(checkpoint, result)
            else:
                for position, case in group[processed:]:
                    answer = dict(case)
                    answer.update(passed=False, error=str(error))
                    result["failures"].append(f"styles/{case['id']}: {error}")
                    record(position, answer)
    result["styles_progress"] = {"completed": completed, "total": len(cases)}
    write_json(checkpoint, result)


def pivot_snapshot(book: Any, cases: list[dict[str, Any]]) -> list[dict[str, Any]]:
    observed = []
    for case in cases:
        sheet = book.Worksheets.Item(case["sheet"])
        pivots = collection(sheet.PivotTables())
        if len(pivots) != 1:
            raise AssertionError(f"pivots/{case['id']}: expected one owned PivotTable")
        pivot = pivots[0]
        area = pivot.TableRange2
        if int(area.Rows.Count) * int(area.Columns.Count) > 4096:
            raise ValueError(f"pivots/{case['id']}: table exceeds bounded observation")
        grid = [[function_value(area.Cells(row, column).Value2)
                 for column in range(1, int(area.Columns.Count) + 1)]
                for row in range(1, int(area.Rows.Count) + 1)]
        def cell_is_error(cell: Any) -> bool:
            answer = book.Application.WorksheetFunction.IsError(cell)
            if type(answer) is tuple and len(answer) == 1 and type(answer[0]) is bool:
                answer = answer[0]
            if type(answer) is not bool:
                raise TypeError("pivots: ISERROR must return a Boolean")
            return answer
        errors = [[cell_is_error(area.Cells(row, column))
                   for column in range(1, int(area.Columns.Count) + 1)]
                  for row in range(1, int(area.Rows.Count) + 1)]
        def fields(axis: Any) -> list[dict[str, Any]]:
            return [{"name": str(field.Name), "caption": str(field.Caption),
                     "position": int(field.Position)} for field in collection(axis)]
        observed.append({"id": case["id"], "name": str(pivot.Name),
                         "table_range2": str(area.Address), "value2": grid, "is_error": errors,
                         "row_fields": fields(pivot.RowFields),
                         "column_fields": fields(pivot.ColumnFields),
                         "data_fields": [{**entry, "function": int(field.Function),
                                         "number_format": str(field.NumberFormat)}
                                         for entry, field in zip(fields(pivot.DataFields),
                                                                 collection(pivot.DataFields))],
                         "row_grand": bool(pivot.RowGrand), "column_grand": bool(pivot.ColumnGrand)})
    return observed


def assert_pivot_cached(source: Path, sheet: str, refreshed: dict[str, Any],
                        native_saved: Path | None = None) -> list[list[dict[str, Any]]]:
    """Compare typed saved cells to Value2 before Excel can refresh on open."""
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    left, top, right, bottom = range_boundaries(refreshed["table_range2"].replace("$", ""))
    if (right - left + 1) * (bottom - top + 1) > 4096:
        raise ValueError("pivots: raw candidate cache exceeds bounded observation")
    cases = [{"sheet": sheet, "cell": f"{get_column_letter(column)}{row}"}
             for row in range(top, bottom + 1) for column in range(left, right + 1)]
    cells = saved_function_cache(source, cases, formulas_only=False, allow_absent=True)
    native = (saved_function_cache(native_saved, cases, formulas_only=False, allow_absent=True)
              if native_saved is not None else None)
    grid: list[list[dict[str, Any]]] = []
    for row in range(top, bottom + 1):
        observed = []
        for column in range(left, right + 1):
            at = f"{get_column_letter(column)}{row}"
            cached = cells.get((sheet, at))
            kind = None if cached is None else cached["type"]
            if refreshed["is_error"][row - top][column - left] != (kind == "e"):
                raise AssertionError(f"pivots/{sheet}!{at}: cached worksheet error kind differs from native RefreshTable")
            if cached is None:
                value = function_value(None)
            else:
                text = cached["value_text"]
                if kind == "n" and text is not None:
                    numeric = float(text)
                    if not math.isfinite(numeric):
                        raise AssertionError(f"pivots/{sheet}!{at}: unproved nonfinite numeric cache")
                    value = function_value(numeric)
                elif kind == "b" and text in ("0", "1"):
                    value = function_value(text == "1")
                elif kind == "str":
                    value = function_value(decode_ooxml_cached_text(text or ""))
                elif kind == "e":
                    if native is not None:
                        agreed = native.get((sheet, at))
                        if agreed is None or agreed["type"] != "e" or agreed["value_text"] != text:
                            raise AssertionError(f"pivots/{sheet}!{at}: cached error literal differs from native saved error")
                        value = {"variant": "error", "value": text}
                    else:
                        codes = {"#NULL!": 2000, "#DIV/0!": 2007, "#VALUE!": 2015,
                                 "#REF!": 2023, "#NAME?": 2029, "#NUM!": 2036,
                                 "#N/A": 2042, "#GETTING_DATA": 2043}
                        if text not in codes:
                            raise AssertionError(f"pivots/{sheet}!{at}: unproved cached error {text!r}")
                        value = function_value(-2146828288 + codes[text])
                elif kind == "n" and text is None:
                    value = function_value(None)
                else:
                    raise AssertionError(f"pivots/{sheet}!{at}: unproved cached cell type {kind!r}")
            expected = refreshed["value2"][row - top][column - left]
            if value["variant"] == "error" and native is not None:
                same = True  # ISERROR and the native saved literal were checked above.
            elif value["variant"] == "float" and expected["variant"] == "int":
                same = value["value"] == expected["value"]
            elif value["variant"] == "int" and expected["variant"] == "float":
                same = value["value"] == expected["value"]
            else:
                same = value == expected
            if not same:
                raise AssertionError(f"pivots/{sheet}!{at}: cached worksheet cell differs from native RefreshTable")
            observed.append(value)
        grid.append(observed)
    return grid


def pivot_xml(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        return {name: archive.read(name).decode("utf-8")
                for name in sorted(archive.namelist())
                if name.startswith(("xl/pivotTables/", "xl/pivotCache/", "xl/worksheets/"))
                and name.endswith((".xml", ".rels"))}


def assert_pivot_snapshots(before: list[dict[str, Any]], after: list[dict[str, Any]], phase: str) -> None:
    if before != after:
        for left, right in zip(before, after):
            if left != right:
                raise AssertionError(f"pivots/{left['id']}: {phase} changed TableRange2, fields or Value2")
        raise AssertionError(f"pivots: {phase} changed PivotTable count")


def run_existing_pivots(app: Any, args: argparse.Namespace,
                        result: dict, manifest: dict[str, Any]) -> None:
    pair = manifest.get("check_existing", manifest.get("compare_existing"))
    result["pivot_existing"] = []
    result["pivot_existing_candidates"] = [
        str(args.output_dir / f"pivot-existing-{item['id']}.pending.xlsx") for item in pair]
    result["pivot_existing_manifest_sha256"] = hashlib.sha256(args.cases.read_bytes()).hexdigest()
    write_json(args.output_dir / "results.json", result)
    for item, pending_name in zip(pair, result["pivot_existing_candidates"]):
        source = args.cases.parent / item["file"]
        selected = {"id": "selected", "sheet": item["sheet"]}
        pending = Path(pending_name)
        with open_checked(app, source) as book:
            pivots = collection(book.Worksheets.Item(item["sheet"]).PivotTables())
            if len(pivots) != 1 or str(pivots[0].Name) != item["name"]:
                raise AssertionError(f"pivots/{item['id']}: selected pivot identity disagrees")
            before = pivot_snapshot(book, [selected])
            if not bool(pivots[0].RefreshTable()):
                raise AssertionError(f"pivots/{item['id']}: RefreshTable returned false")
            refreshed = pivot_snapshot(book, [selected])
            book.SaveAs(str(pending.resolve()), FileFormat=51, AddToMru=False, Local=False)
            saved = pivot_snapshot(book, [selected])
            assert_pivot_snapshots(refreshed, saved, "SaveAs")
        with open_checked(app, pending) as reopened:
            after_reopen = pivot_snapshot(reopened, [selected])
        assert_pivot_snapshots(refreshed, after_reopen, "reopen")
        result["pivot_existing"].append({
            "id": item["id"], "input_sha256": item["sha256"],
            "before_refresh": before, "after_refresh": refreshed,
            "after_save": saved, "after_reopen": after_reopen,
            "input_xml": {name: text for name, text in pivot_xml(source).items()
                          if name in manifest.get("modified_members", ())},
            "saved_xml": {name: text for name, text in pivot_xml(pending).items()
                          if name in manifest.get("modified_members", ())},
            "saved_sha256": hashlib.sha256(pending.read_bytes()).hexdigest(),
        })
        write_json(args.output_dir / "results.json", result)
    if not manifest.get("observe_labels", False):
        candidate = pair[1]
        refreshed = result["pivot_existing"][1]["after_refresh"][0]
        differences = []
        try:
            result["candidate_input_cache"] = assert_pivot_cached(
                args.cases.parent / candidate["file"], candidate["sheet"], refreshed,
                Path(result["pivot_existing_candidates"][1]))
            result["candidate_raw_cache_equal"] = True
        except AssertionError as error:
            result["candidate_raw_cache_equal"] = False
            differences.append(str(error))
        try:
            assert_pivot_snapshots(result["pivot_existing"][0]["after_refresh"],
                                   result["pivot_existing"][1]["after_refresh"],
                                   "control/candidate RefreshTable")
            result["control_candidate_equal"] = True
        except AssertionError as error:
            result["control_candidate_equal"] = False
            differences.append(str(error))
        write_json(args.output_dir / "results.json", result)
        if differences:
            raise AssertionError("; ".join(differences))


def run_authored_pivots(app: Any, args: argparse.Namespace, result: dict, manifest: dict) -> None:
    """Prove authored raw caches against Excel refresh, then SaveAs and reopen."""
    authored = manifest["verify_existing"]
    source = args.cases.parent / authored["file"]
    cases = authored["pivots"]
    staged = args.output_dir / "pivot-native.pending.xlsx"
    result["pivot_fixture_candidate"] = str(staged)
    evidence = result["pivots"] = {
        "manifest_sha256": hashlib.sha256(args.cases.read_bytes()).hexdigest(),
        "source_sha256": authored["sha256"], "source_xml": pivot_xml(source),
        "rust_equivalence_checked": False,
    }
    with open_checked(app, source) as book:
        evidence["before_refresh"] = pivot_snapshot(book, cases)
        for case, observed in zip(cases, evidence["before_refresh"]):
            if observed["name"] != case["name"]:
                raise AssertionError(f"pivots/{case['id']}: unexpected pivot name")
            pivot = book.Worksheets.Item(case["sheet"]).PivotTables().Item(1)
            if not bool(pivot.RefreshTable()):
                raise AssertionError(f"pivots/{case['id']}: RefreshTable returned false")
        evidence["after_refresh"] = pivot_snapshot(book, cases)
        assert_pivot_snapshots(evidence["before_refresh"], evidence["after_refresh"], "RefreshTable")
        book.SaveAs(str(staged.resolve()), FileFormat=51, AddToMru=False, Local=False)
        evidence["after_save"] = pivot_snapshot(book, cases)
        assert_pivot_snapshots(evidence["after_refresh"], evidence["after_save"], "SaveAs")
    with open_checked(app, staged) as reopened:
        evidence["after_reopen"] = pivot_snapshot(reopened, cases)
    assert_pivot_snapshots(evidence["after_save"], evidence["after_reopen"], "reopen")
    evidence["saved_xml"] = pivot_xml(staged)
    evidence["saved_sha256"] = hashlib.sha256(staged.read_bytes()).hexdigest()
    evidence["raw_cache_checks"] = []
    failures = []
    for case, refreshed in zip(cases, evidence["after_refresh"]):
        check = {"id": case["id"], "sheet": case["sheet"]}
        try:
            check["raw_grid"] = assert_pivot_cached(source, case["sheet"], refreshed, staged)
            check["equal"] = True
        except AssertionError as error:
            check.update(equal=False, error=str(error))
            failures.append(str(error))
        evidence["raw_cache_checks"].append(check)
    evidence["rust_equivalence_checked"] = True
    write_json(args.output_dir / "results.json", result)
    if failures:
        raise AssertionError("; ".join(failures))


def apply_pivot_captions(pivot: Any, case: dict[str, Any]) -> dict[str, str]:
    """Set only the owned pivot's explicitly authored caption properties."""
    observed = {}
    if "data_caption" in case:
        field = pivot.DataPivotField
        field.Caption = case["data_caption"]
        observed["data_caption"] = str(field.Caption)
        if observed["data_caption"] != case["data_caption"]:
            raise AssertionError(f"pivots/{case['id']}: data caption readback differs")
    if "grand_total_caption" in case:
        pivot.GrandTotalName = case["grand_total_caption"]
        observed["grand_total_caption"] = str(pivot.GrandTotalName)
        if observed["grand_total_caption"] != case["grand_total_caption"]:
            raise AssertionError(f"pivots/{case['id']}: grand total caption readback differs")
    return observed


def apply_pivot_order(pivot: Any, case: dict[str, Any]) -> list[dict[str, Any]]:
    """Resolve an explicitly requested automatic sort on each owned axis."""
    if "item_order" not in case:
        return []
    order = {"ascending": 1, "descending": 2}[case["item_order"]]
    observed = []
    for name in case["rows"] + case["columns"]:
        field = pivot.PivotFields(name)
        # AutoSort requires SourceName, not the displayed field caption:
        # https://learn.microsoft.com/en-us/office/vba/api/excel.pivotfield.autosort
        source = str(field.SourceName)
        field.AutoSort(order, source)
        actual = int(field.AutoSortOrder)
        if actual != order:
            raise AssertionError(f"pivots/{case['id']}/{name}: AutoSort readback differs")
        observed.append({"field": source, "order": actual})
    return observed


def run_pivots(app: Any, args: argparse.Namespace, result: dict) -> None:
    if not args.active:
        raise ValueError("pivots: active attachment required")
    if args.cases is None:
        raise ValueError("pivots: --cases must name the authored pivot manifest")
    manifest = read_pivot_manifest(args.cases)
    if "verify_existing" in manifest:
        return run_authored_pivots(app, args, result, manifest)
    if "check_existing" in manifest or "compare_existing" in manifest:
        return run_existing_pivots(app, args, result, manifest)
    cases = manifest["cases"]
    staged = args.output_dir / "pivot-native.pending.xlsx"
    result["pivot_fixture_candidate"] = str(staged)
    result["pivots"] = {"manifest_sha256": hashlib.sha256(args.cases.read_bytes()).hexdigest(),
                        "source_sha256": manifest["workbook"]["sha256"],
                        "source_xml": pivot_xml(args.cases.parent / manifest["workbook"]["file"]),
                        "execution_only": True, "rust_equivalence_checked": False,
                        "source_order_grand_is_observed_not_assumed": True}
    write_json(args.output_dir / "results.json", result)
    with open_checked(app, args.cases.parent / manifest["workbook"]["file"]) as book:
        for case in cases:
            source = case["source"]
            # PivotCaches.Create documents SourceData as a string reference.
            # The source and destination are both in this owned workbook.
            source_data = "'" + source["sheet"].replace("'", "''") + "'!" + source["range"]
            cache = book.PivotCaches().Create(SourceType=1, SourceData=source_data)
            pivot = cache.CreatePivotTable(
                TableDestination=book.Worksheets.Item(case["sheet"]).Range(case["anchor"]),
                TableName="P6_" + case["id"])
            pivot.ManualUpdate = True
            for position, name in enumerate(case["rows"], 1):
                field = pivot.PivotFields(name)
                field.Orientation = 1
                field.Position = position
                if not case["subtotals"]:
                    field.Subtotals = (False,) * 12
            for position, name in enumerate(case["columns"], 1):
                field = pivot.PivotFields(name)
                field.Orientation = 2
                field.Position = position
            for value in case["values"]:
                pivot.AddDataField(pivot.PivotFields(value["field"]), value["caption"],
                                   PIVOT_AGGREGATES[value["aggregate"]])
            pivot.RowGrand = case["row_grand_totals"]
            pivot.ColumnGrand = case["column_grand_totals"]
            captions = apply_pivot_captions(pivot, case)
            if captions:
                result["pivots"].setdefault("caption_readback", {})[case["id"]] = captions
                write_json(args.output_dir / "results.json", result)
            pivot.RowAxisLayout(1)  # xlTabularRow, explicit for nested row axes.
            pivot.ManualUpdate = False
            ordering = apply_pivot_order(pivot, case)
            if ordering:
                result["pivots"].setdefault("order_readback", {})[case["id"]] = ordering
            before_refresh = pivot_snapshot(book, [case])
            if not bool(pivot.RefreshTable()):
                raise AssertionError(f"pivots/{case['id']}: RefreshTable returned false")
            after_refresh = pivot_snapshot(book, [case])
            assert_pivot_snapshots(before_refresh, after_refresh, "RefreshTable")
            result["pivots"].setdefault("created", []).append(case["id"])
            write_json(args.output_dir / "results.json", result)
        result["pivots"]["before_save"] = pivot_snapshot(book, cases)
        write_json(args.output_dir / "results.json", result)
        book.SaveAs(str(staged.resolve()), FileFormat=51, AddToMru=False, Local=False)
        result["pivots"]["after_save"] = pivot_snapshot(book, cases)
        assert_pivot_snapshots(result["pivots"]["before_save"], result["pivots"]["after_save"], "SaveAs")
        write_json(args.output_dir / "results.json", result)
    with open_checked(app, staged) as reopened:
        result["pivots"]["after_reopen"] = pivot_snapshot(reopened, cases)
    assert_pivot_snapshots(result["pivots"]["before_save"], result["pivots"]["after_reopen"], "reopen")
    result["pivots"]["saved_xml"] = pivot_xml(staged)
    result["pivots"]["saved_sha256"] = hashlib.sha256(staged.read_bytes()).hexdigest()
    write_json(args.output_dir / "results.json", result)


def finish_pivot_fixture(result: dict, output: Path, passed: bool) -> None:
    if "pivot_existing_candidates" in result:
        expected = [output / f"pivot-existing-{role}.pending.xlsx"
                    for role in ("control", "candidate")]
        actual = result.pop("pivot_existing_candidates")
        if actual != [str(path) for path in expected]:
            raise ValueError("pivots: expected two owned existing-pivot candidates")
        if not passed:
            for path in expected:
                path.unlink(missing_ok=True)
            return
        published = [output / f"pivot-existing-{role}.xlsx"
                     for role in ("control", "candidate")]
        if any(not path.is_file() for path in expected) or any(path.exists() for path in published):
            raise ValueError("pivots: missing candidate or existing published file")
        renamed = []
        try:
            for source, target in zip(expected, published):
                source.rename(target)
                renamed.append((source, target))
        except Exception:
            for source, target in reversed(renamed):
                target.rename(source)
            raise
        result["pivot_existing_published"] = [path.name for path in published]
        return
    expected = output / "pivot-native.pending.xlsx"
    candidate = result.pop("pivot_fixture_candidate", None)
    if not passed and candidate is None:
        expected.unlink(missing_ok=True)
        return
    if candidate is None or Path(candidate).resolve() != expected.resolve():
        raise ValueError("pivots: expected one output-directory candidate")
    if not passed:
        expected.unlink(missing_ok=True)
        return
    target = output / "pivot-native.xlsx"
    if not expected.is_file() or target.exists():
        raise ValueError("pivots: missing candidate or existing published fixture")
    expected.rename(target)
    result["pivots"]["published_file"] = target.name


def function_text_formula(expression: str, code: str) -> str:
    return 'TEXT(' + expression + ',"' + code.replace('"', '""') + '")'


def local_number_format(cell: Any, code: str) -> dict[str, Any]:
    """Ask Excel to translate one invariant format, without a token translator."""
    import pythoncom
    dispatch = cell._oleobj_
    identifier = dispatch.GetIDsOfNames("NumberFormat")
    dispatch.InvokeTypes(identifier, 1033, pythoncom.DISPATCH_PROPERTYPUT,
                         (pythoncom.VT_EMPTY, 0), ((pythoncom.VT_VARIANT, 1),), code)
    invariant = dispatch.InvokeTypes(identifier, 1033, pythoncom.DISPATCH_PROPERTYGET,
                                     (pythoncom.VT_VARIANT, 0), ())
    localized = dispatch.InvokeTypes(dispatch.GetIDsOfNames("NumberFormatLocal"), 1033,
                                     pythoncom.DISPATCH_PROPERTYGET, (pythoncom.VT_VARIANT, 0), ())
    if not isinstance(localized, str) or not localized:
        raise ValueError("functions: NumberFormatLocal must return a nonempty string")
    return {"input_code": code, "input_lcid": 1033, "invariant_readback": invariant,
            "number_format_local": localized}


def prepare_function_text_oracle(sheet: Any, case: dict, evaluate: Callable[[str], Any]) -> dict:
    """Separate owned companion formula; the authored case cell stays intact."""
    plan = case["text_format_oracle"]
    if "cell" not in plan:
        return {"sheet": case["sheet"], "cell": case["cell"],
                "kind": "original_native_cache", "wire_formula": case["wire_formula"],
                "en_us_original": function_value(evaluate(case["wire_formula"])),
                "scope": "unchanged authored cell cache equals original expression evaluated at LCID1033; no format translation"}
    scratch, target = sheet.Range(plan["format_cell"]), sheet.Range(plan["cell"])
    for cell in (scratch, target):
        if bool(cell.HasFormula) or cell.Value2 is not None or bool(cell.MergeCells):
            raise ValueError("functions: TEXT companion and format cells must be empty and unmerged")
    converted = local_number_format(scratch, "[$-409]" + plan["code"])
    formula = function_text_formula(plan["value_expression"], converted["number_format_local"])
    target.Formula = "=" + formula
    target.Calculate()
    return {"sheet": case["sheet"], "cell": plan["cell"], "format_cell": plan["format_cell"],
            "kind": "native_number_format_local_companion", "format_conversion": converted, "wire_formula": formula,
            "en_us_original": function_value(evaluate(case["wire_formula"])),
            "scope": "native-local companion cache equals original expression evaluated at LCID1033; original cache remains separately observed"}


def run_functions(app: Any, args: argparse.Namespace, result: dict) -> None:
    if args.cases is None:
        raise ValueError("functions: --cases must name the authored formula manifest")
    manifest = read_function_manifest(args.cases)
    workbook_specs = manifest["workbooks"]
    cases = manifest["cases"]
    staging = {item["file"]: args.output_dir / f"functions-cached-{item['date_system']}.pending.xlsx"
               for item in workbook_specs}
    for target in staging.values():
        if (target.exists() or target.with_name(target.name.replace(".pending", "")).exists()
                or target.with_name(target.name.replace(".pending.xlsx", ".json")).exists()):
            raise ValueError(f"functions: existing fixture output {target.name}; use a fresh --output-dir")
    result["function_fixture_candidates"] = [str(path) for path in staging.values()]
    result["functions"] = {"manifest": args.cases.name,
                           "manifest_sha256": hashlib.sha256(args.cases.read_bytes()).hexdigest(),
                           "calculation_scope": "owned workbook worksheets, Cases last; never Application.Calculate",
                           "execution_only": True, "rust_equivalence_checked": False,
                           "workbooks": [], "cases": [], "calibrations": {},
                           "text_oracle_case_ids": [case["id"] for case in cases if "text_format_oracle" in case],
                           "english_cache_scope": "original caches retain native locale; declared original-cell or translated-companion TEXT oracles are checked against explicit en-US evaluation"}
    checkpoint = args.output_dir / "results.json"
    answers: list[dict[str, Any] | None] = [None] * len(cases)

    def record(position: int, answer: dict[str, Any]) -> None:
        answers[position] = answer
        result["functions"]["cases"] = [value for value in answers if value is not None]
        write_json(checkpoint, result)

    for item in workbook_specs:
        name = item["file"]
        selected = [(index, case) for index, case in enumerate(cases) if case["file"] == name]
        processed = 0
        observed = {"file": name, "date_system": item["date_system"], "cases": len(selected),
                    "calculated_sheets": [], "saved_candidate": False}
        result["functions"]["workbooks"].append(observed)
        try:
            with open_checked(app, args.cases.parent / name) as book:
                if bool(book.Date1904) != (item["date_system"] == "1904"):
                    raise ValueError(f"functions/{name}: Excel Date1904 disagrees with manifest")
                sheets = collection(book.Worksheets)
                cases_sheet = [sheet for sheet in sheets if sheet.Name == "Cases"]
                if len(cases_sheet) != 1:
                    raise ValueError(f"functions/{name}: expected one Cases worksheet")
                # These are only sheets of the newly opened oracle workbook.
                # https://learn.microsoft.com/en-us/office/vba/api/excel.worksheet.calculate%28method%29
                # Calculate dependencies before the Cases result sheet; the
                # manifest's Scoped formula is then read after its own pass.
                for sheet in sheets:
                    if sheet.Name != "Cases":
                        sheet.Calculate()
                        observed["calculated_sheets"].append(sheet.Name)
                cases_sheet[0].Calculate()
                observed["calculated_sheets"].append("Cases")
                evaluators = {name: evaluate_function(book.Worksheets.Item(name))
                              for name in dict.fromkeys(case["sheet"] for _, case in selected)}
                write_json(checkpoint, result)
                for position, case in selected:
                    answer = dict(case)
                    try:
                        cell = book.Worksheets.Item(case["sheet"]).Range(case["cell"])
                        if not bool(cell.HasFormula):
                            raise AssertionError("expected a physical formula cell")
                        answer["actual"] = function_cell(app, cell, case["cell"], evaluators[case["sheet"]])
                        actual = answer["actual"]["value2"]["value"]
                        if "text_format_oracle" in case:
                            text = prepare_function_text_oracle(book.Worksheets.Item(case["sheet"]), case, evaluators[case["sheet"]])
                            answer["text_oracle"] = text
                            text["actual"] = function_cell(app, book.Worksheets.Item(case["sheet"]).Range(text["cell"]), text["cell"], evaluators[case["sheet"]])
                            if text["actual"]["value2"] != text["en_us_original"]:
                                raise AssertionError("TEXT oracle cache differs from original en-US evaluation")
                            if case["mode"] == "calibration":
                                answer["original_control_passed"] = function_control(case["control"], actual)
                            actual = text["actual"]["value2"]["value"]
                        answer["status"] = "observed"
                        if case["mode"] in ("calibration", "behavior_only"):
                            answer["control_passed"] = function_control(case["control"], actual)
                            if not answer["control_passed"]:
                                result["failures"].append(
                                    f"functions/{case['id']}: {case['mode']} control failed: {answer['actual']!r}")
                    except Exception as error:
                        answer.update(status="error", error=f"{type(error).__name__}: {error}")
                        result["failures"].append(f"functions/{case['id']}: {error}")
                    record(position, answer)
                    processed += 1
                book.SaveAs(str(staging[name].resolve()), FileFormat=51, AddToMru=False, Local=False)
                observed["saved_candidate"] = True
                write_json(checkpoint, result)
                for position, case in selected:
                    answer = answers[position]
                    assert answer is not None
                    try:
                        cell = book.Worksheets.Item(case["sheet"]).Range(case["cell"])
                        answer["after_save"] = function_cell(app, cell, case["cell"], evaluators[case["sheet"]])
                        if text := answer.get("text_oracle"):
                            text["after_save"] = function_cell(app, book.Worksheets.Item(case["sheet"]).Range(text["cell"]), text["cell"], evaluators[case["sheet"]])
                            if text["after_save"]["value2"] != text["en_us_original"]:
                                raise AssertionError("TEXT oracle cache changed from original en-US evaluation after SaveAs")
                        if case["mode"] == "behavior_only":
                            answer["after_save_control_passed"] = function_control(case["control"], answer["after_save"]["value2"]["value"])
                            if not answer["after_save_control_passed"]:
                                result["failures"].append(f"functions/{case['id']}: behavior control failed after SaveAs")
                    except Exception as error:
                        answer["after_save_error"] = f"{type(error).__name__}: {error}"
                        result["failures"].append(f"functions/{case['id']}: after SaveAs observation: {error}")
                    record(position, answer)
            cache_cases = [case for _, case in selected]
            cache_cases += [{**case, "cell": case["text_format_oracle"]["cell"]}
                            for _, case in selected if "cell" in case.get("text_format_oracle", {})]
            persisted = saved_function_cache(staging[name], cache_cases)
            for position, case in selected:
                answer = answers[position]
                assert answer is not None
                answer["saved_cache"] = persisted[(case["sheet"], case["cell"])]
                comparisons = {}
                # Volatile controls may change legitimately on SaveAs. Their
                # after-save value must match the persisted cache; deterministic
                # cases and calibration controls must match from both sides.
                for phase in (("actual", "after_save") if case["mode"] != "behavior_only" else ("after_save",)):
                    if phase not in answer:
                        continue
                    proof = compare_function_cache(answer[phase], answer["saved_cache"])
                    comparisons[phase] = proof
                    if proof["status"] == "different":
                        result["failures"].append(f"functions/{case['id']}: {phase}/saved cache: {proof['reason']}")
                answer["native_cache_comparisons"] = comparisons
                if text := answer.get("text_oracle"):
                    text["saved_cache"] = persisted[(case["sheet"], text["cell"])]
                    text["native_cache_comparisons"] = {}
                    for phase in ("actual", "after_save"):
                        if phase not in text:
                            continue
                        proof = compare_function_cache(text[phase], text["saved_cache"])
                        text["native_cache_comparisons"][phase] = proof
                        if proof["status"] != "equal":
                            result["failures"].append(f"functions/{case['id']}: TEXT companion/{phase}/saved cache: {proof['reason']}")
                record(position, answer)
        except Exception as error:
            result["failures"].append(f"functions/{name}: {type(error).__name__}: {error}")
            for position, case in selected[processed:]:
                record(position, {**case, "status": "error", "error": f"workbook: {error}"})

    for answer in answers:
        assert answer is not None
        if answer["mode"] == "calibration":
            gate = answer["control"]["gate"]
            result["functions"]["calibrations"][gate] = {
                "passed": answer.get("control_passed") is True,
                "actual": answer.get("actual"), "control": answer["control"],
                "scope": answer.get("text_oracle", {}).get("kind", "original_native_cache"),
                "text_oracle": answer.get("text_oracle"),
                "original_control_passed": answer.get("original_control_passed"),
            }
    for answer in answers:
        answer["gate_status"] = {gate: result["functions"]["calibrations"][gate]["passed"]
                                  for gate in answer["gates"]}
    result["functions"]["cases"] = answers
    result["functions"]["cache_authoritative_cases"] = [answer["id"] for answer in answers
        if any(proof["status"] == "cache-authoritative" for proof in answer.get("native_cache_comparisons", {}).values())]
    write_json(checkpoint, result)


def finish_function_fixtures(result: dict, output: Path, passed: bool) -> None:
    candidates = [Path(path) for path in result.get("function_fixture_candidates", [])]
    if not passed:
        for path in candidates:
            if path.parent.resolve() == output.resolve() and path.name in (
                    "functions-cached-1900.pending.xlsx", "functions-cached-1904.pending.xlsx"):
                path.unlink(missing_ok=True)
        result.pop("function_fixture_candidates", None)
        return
    if not candidates:
        raise ValueError("functions: expected cached fixture candidates")
    expected = {output / f"functions-cached-{year}.pending.xlsx" for year in ("1900", "1904")}
    if {path.resolve() for path in candidates} != {path.resolve() for path in expected}:
        raise ValueError("functions: fixture candidates must be the two output-directory workbooks")
    targets = [(source, source.with_name(source.name.replace(".pending", "")))
               for source in sorted(expected)]
    if any(not source.is_file() or target.exists() or target.with_suffix(".json").exists()
           for source, target in targets):
        raise ValueError("functions: missing candidate or existing published fixture")
    published = []
    try:
        for source, target in targets:
            source.rename(target)
            published.append(target)
            functions = result["functions"]
            year = target.stem.rsplit("-", 1)[1]
            cases = [case for case in functions["cases"] if case["date_system"] == year]
            # A focused manifest can save a whole workbook containing other
            # formulas. The file's contents must never imply broader coverage.
            write_json(target.with_suffix(".json"), {
                "schema_version": 1, "kind": "p5_native_function_cache",
                "excel": result["excel"], "manifest": functions["manifest"],
                "manifest_sha256": functions["manifest_sha256"],
                "file": target.name, "sha256": hashlib.sha256(target.read_bytes()).hexdigest(),
                "observations": "results.json#/functions",
                "oracle": "native worksheet calculation and Excel SaveAs; application locale recorded in excel",
                "execution_only": functions["execution_only"],
                "rust_equivalence_checked": functions["rust_equivalence_checked"],
                "coverage": {
                    "scope": "selected manifest cases; other workbook formulas are not validated",
                    "manifest_case_count": len(functions["cases"]), "checked_case_count": len(cases),
                    "text_oracle_case_ids": [case["id"] for case in cases if "text_oracle" in case],
                    "english_cache_scope": functions.get("english_cache_scope", "original native caches; no language adaptation"),
                    "checked_case_ids": [case["id"] for case in cases],
                    "case_gate_status": {case["id"]: case["gate_status"] for case in cases},
                    "calibration_status": {name: gate["passed"] for name, gate in functions["calibrations"].items()},
                    "cache_authoritative_case_ids": [case["id"] for case in cases
                        if case["id"] in functions["cache_authoritative_cases"]],
                },
            })
    except Exception:
        for target in published:
            target.unlink(missing_ok=True)
            target.with_suffix(".json").unlink(missing_ok=True)
        raise
    result.pop("function_fixture_candidates", None)
    result["function_fixtures"] = [str(target) for _, target in targets]


def identity_from_handle(handle: Any, pid: int) -> dict:
    import ctypes
    from ctypes import wintypes
    import win32process
    query = ctypes.windll.kernel32.QueryFullProcessImageNameW
    query.argtypes = [wintypes.HANDLE, wintypes.DWORD, wintypes.LPWSTR, ctypes.POINTER(wintypes.DWORD)]
    query.restype = wintypes.BOOL
    size = wintypes.DWORD(32768)
    buffer = ctypes.create_unicode_buffer(size.value)
    if not query(int(handle), 0, buffer, ctypes.byref(size)):
        raise ctypes.WinError()
    return {"pid": pid, "created": str(win32process.GetProcessTimes(handle)["CreationTime"]),
            "image": buffer.value}


def process_identity(pid: int) -> dict:
    import win32api
    handle = win32api.OpenProcess(0x1000, False, pid)
    try:
        return identity_from_handle(handle, pid)
    finally:
        handle.Close()


def excel_instance(state: Path, *, owned_only: bool = False) -> dict | None:
    if not state.exists():
        return None
    saved = json.loads(state.read_text(encoding="utf-8"))
    if type(saved.get("owned")) is not bool or (owned_only and not saved["owned"]):
        return None
    try:
        identity = {key: saved[key] for key in ("pid", "created", "image")}
        actual = process_identity(identity["pid"])
    except Exception:
        return None
    if actual != identity or Path(actual["image"]).name.lower() != "excel.exe":
        return None
    return identity


def stop_owned_excel(state: Path) -> None:
    import win32api
    if identity := excel_instance(state, owned_only=True):
        handle = win32api.OpenProcess(0x1001, False, identity["pid"])
        try:
            # Validate and terminate through one handle: PID reuse between two
            # OpenProcess calls must never target a later user Excel instance.
            if identity_from_handle(handle, identity["pid"]) == identity:
                win32api.TerminateProcess(handle, 1)
        finally:
            handle.Close()


def visible_modal(state: Path) -> str | None:
    import win32gui
    import win32process
    identity = excel_instance(state)
    if not identity:
        return None
    found = []
    def visit(hwnd: int, _: Any) -> None:
        if win32gui.IsWindowVisible(hwnd) and win32process.GetWindowThreadProcessId(hwnd)[1] == identity["pid"]:
            classname = win32gui.GetClassName(hwnd)
            if classname in ("#32770", "bosa_sdm_XL9", "NUIDialog"):
                found.append(win32gui.GetWindowText(hwnd) or classname)
    win32gui.EnumWindows(visit, None)
    return "; ".join(found) if found else None


def worker(args: argparse.Namespace) -> int:
    import pythoncom
    import win32com.client
    import win32process
    pythoncom.CoInitialize()
    app = None
    window = None
    preferences = []
    result = {"schema_version": 1, "mode": args.mode, "utc": dt.datetime.now(dt.timezone.utc).isoformat(),
              "session": "active" if args.active else "isolated", "cleanup_completed": False, "failures": []}
    try:
        # Publish ownership before either COM call can block. Attached state
        # can never authorize process termination, even without an HWND yet.
        write_json(args.state, {"owned": not args.active})
        app = (win32com.client.GetActiveObject("Excel.Application") if args.active
               else win32com.client.DispatchEx("Excel.Application"))
        write_json(args.state, {**process_identity(win32process.GetWindowThreadProcessId(app.Hwnd)[1]),
                                "owned": not args.active})
        wait_ready(app, active=args.active)
        if not args.active:
            app.Visible = False
        if args.mode != "probe":
            if args.active:
                # Activating this retained window restores its existing sheet
                # and selection; do not overwrite a user's selection explicitly.
                window = app.ActiveWindow
            wanted = (("DisplayAlerts", False), ("EnableEvents", False),
                      ("AskToUpdateLinks", False), ("AutomationSecurity", 3))
            original = {name: getattr(app, name) for name, _ in wanted}
            # A blocked worker may be stopped before finally can restore these.
            # Persist only the four settings we change, for explicit recovery.
            result["application_preferences_before"] = original
            write_json(args.output_dir / "results.json", result)
            for name, value in wanted:
                preferences.append((name, original[name]))
                setattr(app, name, value)
        result["excel"] = {"version": app.Version, "build": str(app.Build), "operating_system": app.OperatingSystem,
                           "country_setting": indexed_property(app, "International", 2),
                           "ui_language": indexed_property(app.LanguageSettings, "LanguageID", 2)}
        write_json(args.output_dir / "results.json", result)
        if not args.active and app.Workbooks.Count != 0:
            raise RuntimeError("isolated Excel unexpectedly contains workbooks")
        if args.mode == "formats":
            run_formats(app, args, result)
        elif args.mode == "fidelity":
            run_fidelity(app, args, result)
        elif args.mode == "shift":
            run_shift(app, args, result)
        elif args.mode == "styles":
            run_styles(app, args, result)
        elif args.mode == "functions":
            run_functions(app, args, result)
        elif args.mode == "pivots":
            run_pivots(app, args, result)
        elif args.mode == "inspect":
            if not args.workbook:
                raise ValueError("inspect: at least one --workbook is required")
            inspect_workbooks(app, args.workbook, result)
    except Exception as error:
        traceback.print_exc()
        result["failures"].append(f"{args.mode}: {type(error).__name__}: {error}")
    finally:
        cleanup_failures = len(result["failures"])
        if app is not None:
            if not args.active:
                try:
                    for book in collection(app.Workbooks):
                        book.Close(SaveChanges=False)
                except Exception as error:
                    result["failures"].append(f"cleanup/workbooks: {error}")
            elif window is not None:
                try:
                    # Restore focus while events are still disabled. User
                    # workbooks opened during the check are never enumerated.
                    window.Activate()
                except Exception as error:
                    result["failures"].append(f"cleanup/window: {error}")
            for name, value in reversed(preferences):
                try:
                    setattr(app, name, value)
                except Exception as error:
                    result["failures"].append(f"cleanup/{name}: {error}")
            if not args.active:
                try:
                    app.Quit()
                except Exception as error:
                    result["failures"].append(f"cleanup/quit: {error}")
        result["cleanup_completed"] = len(result["failures"]) == cleanup_failures
        window = app = None
        gc.collect()
        pythoncom.CoUninitialize()
        result["passed"] = not result["failures"]
        write_json(args.output_dir / "results.json", result)
    return 0 if result["passed"] else 1


def parser() -> argparse.ArgumentParser:
    result = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    result.add_argument("mode", choices=("probe", "inspect", "fidelity", "shift", "formats", "styles", "functions", "pivots"))
    result.add_argument("--output-dir", type=Path, default=OUTPUT)
    result.add_argument("--exchange", type=Path, default=EXCHANGE)
    result.add_argument("--rich", type=Path)
    result.add_argument("--workbook", type=Path, action="append", default=[])
    result.add_argument("--cases", type=Path)
    result.add_argument("--shift-original", type=Path)
    result.add_argument("--shifted", type=Path)
    result.add_argument("--sheet", default="Data")
    result.add_argument("--insert-row", type=int, default=3)
    result.add_argument("--insert-column", type=int, default=1)
    result.add_argument("--timeout", type=float, default=180)
    result.add_argument("--active", action="store_true",
                        help="attach to existing Excel; never quit or terminate it")
    result.add_argument("--worker", action="store_true", help=argparse.SUPPRESS)
    result.add_argument("--state", type=Path, help=argparse.SUPPRESS)
    return result


def main() -> int:
    # Console encodings may be narrower than workbook text. Keep JSON exact
    # and render unencodable diagnostic characters visibly instead of crashing.
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="backslashreplace")
    args = parser().parse_args()
    if sys.platform != "win32":
        raise SystemExit("desktop oracle requires Windows Microsoft Excel and pywin32")
    if not math.isfinite(args.timeout) or args.timeout <= 0 or args.insert_row < 1 or args.insert_column < 1:
        raise SystemExit("timeout must be finite and positive; insert coordinates must be positive")
    args.output_dir = args.output_dir.resolve()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    if args.worker:
        return worker(args)
    if args.mode == "fidelity":
        for name in ("rich_excel.xlsx", "rich_excel.json", "rich_excel.pending.xlsx", "rich_excel.pending.json", "rich_excel.control.xlsx"):
            if (args.output_dir / name).exists():
                raise SystemExit(f"fidelity: {name} already exists; select a fresh --output-dir")
    if args.mode == "pivots":
        if not args.active or args.cases is None:
            raise SystemExit("pivots: --active and --cases are required")
        manifest = read_pivot_manifest(args.cases)
        names = ((f"pivot-existing-{role}{suffix}.xlsx"
                  for role in ("control", "candidate") for suffix in (".pending", ""))
                 if "check_existing" in manifest or "compare_existing" in manifest else
                 iter(("pivot-native.pending.xlsx", "pivot-native.xlsx")))
        if any((args.output_dir / name).exists() for name in names):
            raise SystemExit("pivots: fixture output exists; select a fresh --output-dir")
    if args.mode == "functions":
        if args.cases is None:
            raise SystemExit("functions: --cases must name the authored formula manifest")
        read_function_manifest(args.cases)
        for year in ("1900", "1904"):
            for name in (f"functions-cached-{year}.pending.xlsx", f"functions-cached-{year}.xlsx",
                         f"functions-cached-{year}.json"):
                if (args.output_dir / name).exists():
                    raise SystemExit(f"functions: {name} already exists; select a fresh --output-dir")
    write_json(args.output_dir / "results.json", {"schema_version": 1, "mode": args.mode,
                                                "utc": dt.datetime.now(dt.timezone.utc).isoformat(),
                                                "session": "active" if args.active else "isolated",
                                                "cleanup_completed": False, "failures": [], "passed": False})
    with tempfile.TemporaryDirectory(prefix="yggdryl-excel-oracle-") as directory:
        state = Path(directory) / "instance.json"
        command = [sys.executable, str(Path(__file__).resolve()), *sys.argv[1:], "--worker", "--state", str(state)]
        failure = None
        with (args.output_dir / "worker.log").open("w", encoding="utf-8") as log:
            child = subprocess.Popen(command, stdout=log, stderr=subprocess.STDOUT, creationflags=subprocess.CREATE_NO_WINDOW)
            deadline = time.monotonic() + args.timeout
            try:
                while child.poll() is None:
                    if time.monotonic() >= deadline:
                        failure = f"{args.mode}: timeout after {args.timeout:g}s"
                        break
                    if modal := visible_modal(state):
                        failure = f"{args.mode}: Excel modal dialog {modal!r}"
                        break
                    time.sleep(0.2)
            finally:
                if not failure and (modal := visible_modal(state)):
                    failure = f"{args.mode}: Excel modal dialog {modal!r}"
                if child.poll() is None:
                    child.terminate()
                    child.wait(timeout=10)
                if not args.active:
                    stop_owned_excel(state)
        result_path = args.output_dir / "results.json"
        if failure:
            result = json.loads(result_path.read_text(encoding="utf-8")) if result_path.exists() else {"schema_version": 1, "mode": args.mode, "failures": []}
            result["failures"].append(failure)
            result["passed"] = False
            write_json(result_path, result)
        elif child.returncode:
            result = json.loads(result_path.read_text(encoding="utf-8"))
            if not result["failures"]:
                result["failures"].append(f"{args.mode}: worker exited {child.returncode}; see worker.log")
            result["passed"] = False
            write_json(result_path, result)
        if not result_path.exists():
            print(f"FAIL {args.mode}: worker failed; see {args.output_dir / 'worker.log'}", file=sys.stderr)
            return 1
        result = json.loads(result_path.read_text(encoding="utf-8"))
        if args.active and not result.get("cleanup_completed", False):
            result["cleanup_limitation"] = (
                "The attached Excel process was not terminated. Worker cleanup could not complete: "
                "application preferences/focus may be unrestored and oracle workbooks may remain open; "
                "user workbooks were not closed by supervisor cleanup")
            result["failures"].append(result["cleanup_limitation"])
        elif not state.exists() or "pid" not in json.loads(state.read_text(encoding="utf-8")):
            result["cleanup_limitation"] = (
                "Excel startup identity unavailable: attachment, DispatchEx or Hwnd may have blocked before identity was recorded; "
                "any unidentified Excel process was not terminated")
            result["failures"].append(result["cleanup_limitation"])
        passed = result.get("passed", False) and child.returncode == 0 and not failure and not result["failures"]
        try:
            if args.mode == "fidelity":
                finish_fixture(result, args.output_dir, passed)
            elif args.mode == "functions":
                finish_function_fixtures(result, args.output_dir, passed)
            elif args.mode == "pivots":
                finish_pivot_fixture(result, args.output_dir, passed)
        except Exception as error:
            result["failures"].append(f"fixture finalization: {error}")
            passed = False
        result["passed"] = passed
        write_json(result_path, result)
        for message in result["failures"]:
            print(f"FAIL {message}")
        print(f"{'PASS' if passed else 'FAIL'} {args.mode}: Excel {result.get('excel', {})}; {result_path}")
        return 0 if passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
